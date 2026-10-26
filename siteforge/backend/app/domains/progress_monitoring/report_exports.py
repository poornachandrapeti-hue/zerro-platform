"""Frozen snapshot based exports; no current/live query or privilege widening."""
from io import BytesIO
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

SECTIONS = ('summary','activities','targets','exceptions')
def check_sections(sections):
    selected = list(dict.fromkeys(sections or SECTIONS))
    if not selected or any(s not in SECTIONS for s in selected):
        raise ValueError('Select one or more valid export sections')
    return selected

def summary(row, snapshot):
    pub=snapshot.get('publication') or {}
    construction=(snapshot.get('construction_measures') or {})
    return [('Report',row.report_number),('Revision',row.revision),('Status',row.status),
        ('Report type',row.report_type),('Period',f'{row.period_start} to {row.period_end}'),
        ('Data date',str(row.data_date)),('Source cutoff',str(row.source_cutoff or 'Unavailable')),
        ('Measure',snapshot.get('measure') or 'Unavailable'),('Baseline revision',pub.get('revision','Unavailable')),
        ('Reported %',construction.get('reported_percent')),('Verified %',construction.get('verified_percent')),
        ('QA-accepted %',construction.get('accepted_percent')),
        ('Overall EPCC %',snapshot.get('overall_percent')),
        ('Calculation',snapshot.get('calculation_method') or 'Unavailable'),
        ('Commentary',row.commentary or 'No commentary')]

def section_rows(row,snapshot,section):
    if section=='summary':return [('Field','Value')]+summary(row,snapshot)
    if section=='activities':
        keys=('code','name','unit','scope_quantity','reported_quantity','verified_quantity','accepted_quantity','reported_percent','verified_percent','accepted_percent')
        return [tuple(keys)]+[tuple(a.get(k) for k in keys) for a in snapshot.get('actuals',[]) if isinstance(a,dict)]
    if section=='targets':return [('Period end','Planned %')]+[(t.get('period_end'),t.get('planned_percent')) for t in snapshot.get('planned',[]) if isinstance(t,dict)]
    if section=='exceptions':return [('Data quality exceptions',)]+[(str(e),) for e in snapshot.get('exceptions',[])]
    raise ValueError('Invalid section')

def xlsx_bytes(row, snapshot, sections):
    wb=Workbook();wb.remove(wb.active)
    for section in check_sections(sections):
        sheet=wb.create_sheet(section.title());sheet.freeze_panes='A2'
        for record in section_rows(row,snapshot,section):
            sheet.append([v if v is None or isinstance(v,(int,float,str,bool)) else str(v) for v in record])
        for cell in sheet[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='0B2140')
        sheet.auto_filter.ref=sheet.dimensions
        for cells in sheet.columns:
            letter=get_column_letter(cells[0].column)
            max_len=max((len(str(c.value or '')) for c in cells),default=5)
            sheet.column_dimensions[letter].width=min(55,max(14,max_len+3))
            for cell in cells:cell.alignment=Alignment(vertical='top',wrap_text=True)
    stream=BytesIO();wb.save(stream);return stream.getvalue()

def pdf_bytes(row,snapshot,sections):
    selected=check_sections(sections)
    stream=BytesIO();doc=SimpleDocTemplate(stream,pagesize=A4,rightMargin=34,leftMargin=34,topMargin=36,bottomMargin=36)
    styles=getSampleStyleSheet();normal=styles['BodyText'];normal.fontSize=8;normal.leading=11
    story=[Paragraph('ZerroBuild | Progress report',styles['Title']),
        Paragraph(escape(f'{row.report_number}  •  Revision {row.revision}  •  {row.status}'),normal),Spacer(1,14)]
    width=A4[0]-68
    for section in selected:
        story.append(Paragraph(escape(section.title()),styles['Heading2']))
        rows=section_rows(row,snapshot,section)
        if len(rows)==1:story.append(Paragraph('No data available in the frozen report snapshot.',normal));continue
        for start in range(1,len(rows),35):
            chunk=[rows[0]]+rows[start:start+35]
            n=len(chunk[0]);cellw=width/n
            data=[[Paragraph(escape(str(v if v is not None else 'N/A'))[:450],normal) for v in record] for record in chunk]
            t=Table(data,colWidths=[cellw]*n,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0B2140')),
                ('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),0.2,colors.lightgrey),
                ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),
                ('BOTTOMPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5)]))
            story.append(t);story.append(Spacer(1,10))
    doc.build(story);return stream.getvalue()
