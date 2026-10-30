from .models import CatalogueSpecificationTemplate
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import update
from sqlalchemy.orm.attributes import flag_modified

from app.domains.coordination.models import ApprovalPolicy, PolicyReviewer, PolicyStage, ReviewInstance
from app.domains.engineering_documents.models import Drawing, DrawingRevision, StoredFile
from app.domains.identity_projects.models import (
    Area,
    Company,
    Discipline,
    Membership,
    Project,
    ProjectCompany,
    ProjectDiscipline,
    Role,
    User,
)
from app.domains.identity_projects.security import holds, permitted

from .models import (
    MaterialRequirement,
    ProcurementEvent,
    ProcurementFile,
    ProcurementItem,
    PurchaseRequisition,
    RequestForQuotation,
    RequirementLine,
    RequisitionLine,
    UnitConversion,
    Vendor,
    Warehouse,
    now,
)
from .schemas import ITEM_TYPES, PRIORITIES, VENDOR_STATES

REQUIREMENT = "procurement.requirement"
REQUISITION = "procurement.requisition"
PURCHASE_ORDER = "procurement.purchase_order"
EDITABLE = {"Draft", "Returned for Correction"}
HOLDING = {"Draft", "Pending Approval", "Approved", "Returned for Correction"}
OPEN_REQUIREMENTS = {"Draft", "Pending Approval", "Returned for Correction"}


def _key(value):
    return value.strip().casefold()


def _amount(value):
    return Decimal(str(value or 0))


def _stamp(row, user):
    row.modified_by = user.id
    row.modified_at = now()


def _event(db, user, project_id, resource, resource_id, kind, detail=""):
    db.add(
        ProcurementEvent(
            project_id=project_id,
            resource=resource,
            resource_id=resource_id,
            kind=kind,
            detail=detail,
            actor_id=user.id,
        )
    )


def _bump(db, model, row_id, version, values):
    result = db.execute(
        update(model).where(model.id == row_id, model.version == version).values(version=version + 1, **values)
    )
    if result.rowcount != 1:
        raise HTTPException(409, "This record changed. Reload it and try again")


def role_allows(db, user, project, action, discipline_id=None):
    if not user or not user.active or project.organization_id != user.organization_id:
        return False
    enabled = {
        row.discipline_id
        for row in db.query(ProjectDiscipline).filter_by(project_id=project.id, active=True)
    }
    if discipline_id is not None and discipline_id not in enabled:
        return False
    memberships = db.query(Membership).filter_by(project_id=project.id, user_id=user.id, active=True)
    for member in memberships:
        role = db.get(Role, member.role_id)
        if not role or action not in (role.permissions or []):
            continue
        if discipline_id is None or member.scope == "ALL" or discipline_id in (member.discipline_ids or []):
            return True
    return False


def reviewer_can_view(db, user, project):
    if not user or not user.active or project.organization_id != user.organization_id:
        return False
    policies = (
        db.query(ApprovalPolicy)
        .filter(
            ApprovalPolicy.project_id == project.id,
            ApprovalPolicy.active.is_(True),
            ApprovalPolicy.record_type.in_((REQUIREMENT, REQUISITION, PURCHASE_ORDER)),
        )
        .all()
    )
    stage_ids = [
        stage.id
        for policy in policies
        for stage in db.query(PolicyStage).filter_by(policy_id=policy.id)
    ]
    if stage_ids and (
        db.query(PolicyReviewer)
        .filter(PolicyReviewer.stage_id.in_(stage_ids), PolicyReviewer.user_id == user.id)
        .first()
    ):
        return True
    from app.domains.coordination.service import current_stage

    instances = (
        db.query(ReviewInstance)
        .filter(
            ReviewInstance.project_id == project.id,
            ReviewInstance.record_type.in_((REQUIREMENT, REQUISITION, PURCHASE_ORDER)),
            ReviewInstance.state.in_(("In Review", "Returned")),
        )
        .all()
    )
    for instance in instances:
        stage = current_stage(db, instance)
        if stage and db.query(PolicyReviewer).filter_by(stage_id=stage.id, user_id=user.id).first():
            return True
    return False


def _project(db, user, pid, action="procurement.view", discipline_id=None):
    project = db.get(Project, pid)
    if not project or project.organization_id != user.organization_id:
        raise HTTPException(404, "Project not found")
    if action:
        try:
            permitted(db, user, project, action, discipline_id)
        except HTTPException:
            if action == "procurement.view" and reviewer_can_view(db, user, project):
                return project
            raise
    return project


def _org_master(db, user):
    if user.organization_admin:
        return
    memberships = (
        db.query(Membership)
        .join(Project, Project.id == Membership.project_id)
        .filter(
            Membership.user_id == user.id,
            Membership.active.is_(True),
            Project.organization_id == user.organization_id,
        )
    )
    for member in memberships:
        role = db.get(Role, member.role_id)
        if role and "procurement.master.edit" in (role.permissions or []):
            return
    raise HTTPException(403, "Procurement master data requires procurement configuration access")


def _commercial(db, user, project):
    return holds(db, user, project, "procurement.commercial.view")


def _company_actor(db, user, project, company_id):
    if user.organization_admin:
        return
    rows = db.query(Membership).filter_by(project_id=project.id, user_id=user.id, active=True)
    if not any(row.company_id == company_id for row in rows):
        raise HTTPException(403, "You can change records only for your assigned company")


def _project_company(db, project, company_id):
    row = db.query(ProjectCompany).filter_by(
        project_id=project.id, company_id=company_id, active=True
    ).first()
    if not row:
        raise HTTPException(422, "Choose a company assigned to this project")


def _discipline(db, project, discipline_id):
    row = db.query(ProjectDiscipline).filter_by(
        project_id=project.id, discipline_id=discipline_id, active=True
    ).first()
    if not row:
        raise HTTPException(422, "Choose a discipline enabled on this project")


def _member(db, project, user_id):
    person = db.get(User, user_id)
    if not person or not person.active or person.organization_id != project.organization_id:
        raise HTTPException(422, "Choose an active person in this organization")
    member = db.query(Membership).filter_by(
        project_id=project.id, user_id=user_id, active=True
    ).first()
    if not member and not person.organization_admin:
        raise HTTPException(422, "Choose a person assigned to this project")
    return person


def _area(db, project, area_id):
    if not area_id:
        return None
    area = db.get(Area, area_id)
    if not area or area.project_id != project.id or not area.active:
        raise HTTPException(422, "Choose an area from this project")
    return area


def _wbs(db, project, wbs_id):
    if not wbs_id:
        return None
    from app.domains.planning_reporting.workspace_models import PlanningWBS

    wbs = db.get(PlanningWBS, wbs_id)
    if not wbs or wbs.project_id != project.id or wbs.organization_id != project.organization_id or not wbs.active:
        raise HTTPException(422, "Choose an active WBS element from this project")
    return wbs


def _editable_package(package):
    if package.status != "Draft" or package.release_state == "In Review":
        raise HTTPException(
            409,
            f"Work package {package.code} is {package.status if package.status != 'Draft' else 'in release review'}; "
            "change its material requirement links from Construction before release",
        )


def _sync_work_packages(db, project, row, package_ids):
    if package_ids is None:
        return
    from app.domains.construction.models import WorkPackage, WorkPackageRequirement

    if len(package_ids) != len(set(package_ids)):
        raise HTTPException(422, "Choose each work package once")
    wanted = set(package_ids)
    existing = {
        link.package_id: link
        for link in db.query(WorkPackageRequirement).filter_by(requirement_id=row.id)
    }
    for package_id in wanted - existing.keys():
        package = db.get(WorkPackage, package_id)
        if not package or package.project_id != project.id or package.organization_id != project.organization_id:
            raise HTTPException(422, "Choose a work package from this project")
        _editable_package(package)
        db.add(WorkPackageRequirement(package_id=package_id, requirement_id=row.id))
    for package_id in existing.keys() - wanted:
        _editable_package(db.get(WorkPackage, package_id))
        db.delete(existing[package_id])


def _requirement_packages(db, row):
    from app.domains.construction.models import WorkPackage, WorkPackageRequirement

    packages = []
    for link in db.query(WorkPackageRequirement).filter_by(requirement_id=row.id):
        package = db.get(WorkPackage, link.package_id)
        if package:
            packages.append({"id": package.id, "code": package.code, "name": package.name, "status": package.status})
    return sorted(packages, key=lambda package: package["code"])


def _drawing(db, project, drawing_id, revision_id):
    if revision_id:
        revision = db.get(DrawingRevision, revision_id)
        drawing = db.get(Drawing, revision.drawing_id) if revision else None
        if (
            not drawing
            or drawing.project_id != project.id
            or drawing.organization_id != project.organization_id
        ):
            raise HTTPException(422, "Choose a drawing revision from this project")
        if drawing_id and drawing_id != drawing.id:
            raise HTTPException(422, "The revision does not belong to the selected drawing")
        return drawing.id, revision.id
    if drawing_id:
        drawing = db.get(Drawing, drawing_id)
        if not drawing or drawing.project_id != project.id:
            raise HTTPException(422, "Choose a drawing from this project")
        return drawing.id, None
    return None, None


def _unit(db, organization_id, base_unit, unit):
    if _key(unit) == _key(base_unit):
        return
    forward = db.query(UnitConversion).filter_by(
        organization_id=organization_id, from_unit=unit, to_unit=base_unit
    ).first()
    backward = db.query(UnitConversion).filter_by(
        organization_id=organization_id, from_unit=base_unit, to_unit=unit
    ).first()
    if not forward and not backward:
        raise HTTPException(422, "Configure a unit conversion before using a different unit")


def _item(db, user, item_id):
    item = db.get(ProcurementItem, item_id)
    if not item or item.organization_id != user.organization_id or not item.active:
        raise HTTPException(422, "Choose an active item from this organization")
    return item


def _name(db, user_id):
    person = db.get(User, user_id) if user_id else None
    return person.name if person else ""


def _files(db, resource, resource_id):
    rows = db.query(ProcurementFile).filter_by(resource=resource, resource_id=resource_id).all()
    result = []
    for row in rows:
        stored = db.get(StoredFile, row.file_id)
        if stored:
            result.append(
                {"id": row.id, "file_id": stored.id, "name": stored.original_name, "size": stored.size}
            )
    return result


def _events(db, resource, resource_id):
    rows = (
        db.query(ProcurementEvent)
        .filter_by(resource=resource, resource_id=resource_id)
        .order_by(ProcurementEvent.id)
        .all()
    )
    return [
        {
            "id": row.id,
            "kind": row.kind,
            "detail": row.detail,
            "actor": _name(db, row.actor_id),
            "created_at": row.created_at.isoformat(),
        }
        for row in rows
    ]


def _review(db, user, project_id, record_type, record_id):
    row = (
        db.query(ReviewInstance)
        .filter_by(project_id=project_id, record_type=record_type, record_id=record_id)
        .order_by(ReviewInstance.id.desc())
        .first()
    )
    if not row:
        return None
    from app.domains.coordination.service import review_payload

    return review_payload(db, user, row)


def _reserved(db, line_id):
    total = Decimal("0")
    approved = Decimal("0")
    rows = (
        db.query(RequisitionLine, PurchaseRequisition)
        .join(PurchaseRequisition, PurchaseRequisition.id == RequisitionLine.requisition_id)
        .filter(RequisitionLine.requirement_line_id == line_id)
        .all()
    )
    for line, requisition in rows:
        if requisition.status in HOLDING:
            total += _amount(line.quantity)
        if requisition.status == "Approved":
            approved += _amount(line.quantity)
    return total, approved


def _line_trace(db, line):
    from app.domains.engineering_documents.models import MaterialTakeoff, TakeoffLine
    from app.domains.planning_reporting.workspace_models import PlanningWBS

    revision = db.get(DrawingRevision, line.drawing_revision_id) if line.drawing_revision_id else None
    drawing = db.get(Drawing, revision.drawing_id) if revision else None
    source = db.get(TakeoffLine, line.takeoff_line_id) if line.takeoff_line_id else None
    takeoff = db.get(MaterialTakeoff, source.takeoff_id) if source else None
    wbs = db.get(PlanningWBS, line.planning_wbs_id) if line.planning_wbs_id else None
    return {
        "planning_wbs_id": line.planning_wbs_id,
        "wbs_code": wbs.code if wbs else "",
        "wbs_name": wbs.name if wbs else "",
        "takeoff_line_id": line.takeoff_line_id,
        "takeoff_id": takeoff.id if takeoff else None,
        "takeoff_number": takeoff.number if takeoff else "",
        "takeoff_line_position": source.position if source else None,
        "drawing_revision_id": line.drawing_revision_id,
        "drawing_number": drawing.number if drawing else "",
        "drawing_revision_code": revision.code if revision else "",
        "source_page": source.page if source else None,
    }


def _line_payload(db, line):
    from .fulfilment import allocated

    reserved, approved = _reserved(db, line.id)
    stock = allocated(db, line.id)
    required = _amount(line.quantity)
    return {
        "id": line.id,
        "position": line.position,
        "item_id": line.item_id,
        "item_code": db.get(ProcurementItem, line.item_id).code,
        "description": db.get(ProcurementItem, line.item_id).description,
        "specification": line.specification,
        "quantity": format(required.quantize(Decimal("0.001")), "f"),
        "unit": line.unit,
        "required_date": line.required_date.isoformat(),
        "tag": line.tag,
        "engineering_reference": line.engineering_reference,
        "reserved_quantity": format(reserved, "f"),
        "approved_requisition_quantity": format(approved, "f"),
        "accepted_stock_quantity": format(stock, "f"),
        "stock_allocated_quantity": format(stock, "f"),
        "outstanding_quantity": format(required - approved - stock, "f"),
        "remaining_for_requisition": format(required - reserved - stock, "f"),
        **_line_trace(db, line),
    }


def _fulfillment(lines):
    if not lines:
        return "No lines"
    approved = sum(_amount(line["approved_requisition_quantity"]) for line in lines)
    stock = sum(_amount(line["stock_allocated_quantity"]) for line in lines)
    required = sum(_amount(line["quantity"]) for line in lines)
    if approved + stock <= 0:
        return "Not allocated"
    if stock and approved + stock >= required:
        return "Fully covered by stock and approved requisitions"
    if approved >= required:
        return "Fully allocated to approved requisitions"
    if stock:
        return "Partially covered by stock and approved requisitions"
    return "Partially allocated to approved requisitions"


def _takeoff_number(db, takeoff_id):
    if not takeoff_id:
        return ""
    from app.domains.engineering_documents.models import MaterialTakeoff

    takeoff = db.get(MaterialTakeoff, takeoff_id)
    return takeoff.number if takeoff else ""


def requirement_payload(db, user, project, row):
    lines = [
        _line_payload(db, line)
        for line in db.query(RequirementLine)
        .filter_by(requirement_id=row.id)
        .order_by(RequirementLine.position)
    ]
    return {
        "id": row.id,
        "number": row.number,
        "project_id": row.project_id,
        "company_id": row.company_id,
        "company": db.get(Company, row.company_id).name,
        "discipline_id": row.discipline_id,
        "discipline": db.get(Discipline, row.discipline_id).name,
        "area_id": row.area_id,
        "requested_by": row.requested_by,
        "requested_by_name": _name(db, row.requested_by),
        "required_date": row.required_date.isoformat(),
        "priority": row.priority,
        "purpose": row.purpose,
        "remarks": row.remarks,
        "drawing_id": row.drawing_id,
        "drawing_revision_id": row.drawing_revision_id,
        "drawing_number": db.get(Drawing, row.drawing_id).number if row.drawing_id else "",
        "drawing_revision_code": db.get(DrawingRevision, row.drawing_revision_id).code
        if row.drawing_revision_id
        else "",
        "source": row.source,
        "takeoff_id": row.takeoff_id,
        "takeoff_number": _takeoff_number(db, row.takeoff_id),
        "work_packages": _requirement_packages(db, row),
        "status": row.status,
        "fulfillment": _fulfillment(lines),
        "cancellation_reason": row.cancellation_reason,
        "impact_state": row.impact_state,
        "impact_detail": row.impact_detail,
        "impact_decision": row.impact_decision,
        "version": row.version,
        "created_by": row.created_by,
        "created_by_name": _name(db, row.created_by),
        "modified_by_name": _name(db, row.modified_by),
        "created_at": row.created_at.isoformat(),
        "modified_at": row.modified_at.isoformat() if row.modified_at else "",
        "lines": lines,
        "files": _files(db, "requirements", row.id),
        "events": _events(db, "requirements", row.id),
        "review": _review(db, user, project.id, REQUIREMENT, row.id),
        "stock_available": False,
    }


def requisition_payload(db, user, project, row):
    from .orders import ordering_summary

    ordering_state, ordering_note = ordering_summary(db, row)
    show_prices = _commercial(db, user, project)
    lines = []
    for line in (
        db.query(RequisitionLine)
        .filter_by(requisition_id=row.id)
        .order_by(RequisitionLine.position)
    ):
        price = _amount(line.estimated_price) if line.estimated_price is not None else None
        source_line = db.get(RequirementLine, line.requirement_line_id)
        source = db.get(MaterialRequirement, source_line.requirement_id) if source_line else None
        lines.append(
            {
                "id": line.id,
                "position": line.position,
                "requirement_line_id": line.requirement_line_id,
                "requirement_id": source.id if source else None,
                "requirement_number": source.number if source else "",
                **(_line_trace(db, source_line) if source_line else {}),
                "item_id": line.item_id,
                "item_code": db.get(ProcurementItem, line.item_id).code,
                "specification": line.specification,
                "quantity": format(_amount(line.quantity), "f"),
                "unit": line.unit,
                "estimated_price": format(price, "f") if show_prices and price is not None else None,
            }
        )
    return {
        "id": row.id,
        "number": row.number,
        "project_id": row.project_id,
        "company_id": row.company_id,
        "company": db.get(Company, row.company_id).name,
        "discipline_id": row.discipline_id,
        "discipline": db.get(Discipline, row.discipline_id).name,
        "requested_by": row.requested_by,
        "requested_by_name": _name(db, row.requested_by),
        "request_date": row.request_date.isoformat(),
        "required_delivery": row.required_delivery.isoformat(),
        "justification": row.justification,
        "currency": row.currency if show_prices else "",
        "prices_visible": show_prices,
        "status": row.status,
        "ordering": ordering_state,
        "ordering_note": ordering_note,
        "cancellation_reason": row.cancellation_reason,
        "version": row.version,
        "created_by_name": _name(db, row.created_by),
        "modified_by_name": _name(db, row.modified_by),
        "created_at": row.created_at.isoformat(),
        "modified_at": row.modified_at.isoformat() if row.modified_at else "",
        "lines": lines,
        "files": _files(db, "requisitions", row.id),
        "events": _events(db, "requisitions", row.id),
        "review": _review(db, user, project.id, REQUISITION, row.id),
    }


def _check_item_use(db, user, project, item, discipline_id, unit):
    if discipline_id not in (item.discipline_ids or []):
        raise HTTPException(422, "This item is not configured for the selected discipline")
    _unit(db, user.organization_id, item.base_unit, unit)


def _sync_requirement_lines(db, user, project, row, incoming):
    existing = {
        line.id: line for line in db.query(RequirementLine).filter_by(requirement_id=row.id)
    }
    seen = set()
    for index, item in enumerate(incoming, start=1):
        stored = _item(db, user, item.item_id)
        _check_item_use(db, user, project, stored, row.discipline_id, item.unit)
        wbs_id = getattr(item, "planning_wbs_id", None)
        _wbs(db, project, wbs_id)
        if item.id:
            line = existing.get(item.id)
            if not line:
                raise HTTPException(422, "A line on this requirement was not found")
            from .fulfilment import allocated

            reserved, _approved = _reserved(db, line.id)
            if _amount(item.quantity) < reserved + allocated(db, line.id):
                raise HTTPException(
                    422,
                    "Quantity cannot be less than the quantity already on purchase requisitions or allocated stock",
                )
            line.position = index
            line.item_id = item.item_id
            line.specification = item.specification
            line.quantity = item.quantity
            line.unit = item.unit
            line.required_date = item.required_date
            line.tag = item.tag
            line.engineering_reference = item.engineering_reference
            if "planning_wbs_id" in getattr(item, "model_fields_set", ()):
                line.planning_wbs_id = wbs_id
            seen.add(line.id)
        else:
            created = RequirementLine(
                requirement_id=row.id,
                position=index,
                item_id=item.item_id,
                specification=item.specification,
                quantity=item.quantity,
                unit=item.unit,
                required_date=item.required_date,
                tag=item.tag,
                engineering_reference=item.engineering_reference,
                planning_wbs_id=wbs_id,
            )
            db.add(created)
    for line in existing.values():
        if line.id in seen:
            continue
        from .fulfilment import allocated

        reserved, _approved = _reserved(db, line.id)
        if reserved > 0 or allocated(db, line.id) > 0:
            raise HTTPException(
                422, "A line that is already on a purchase requisition or has allocated stock cannot be removed"
            )
        db.delete(line)


def _fill_requirement(db, user, project, row, data):
    _project_company(db, project, data.company_id)
    _company_actor(db, user, project, data.company_id)
    _discipline(db, project, data.discipline_id)
    _member(db, project, data.requested_by)
    _area(db, project, data.area_id)
    if data.priority not in PRIORITIES:
        raise HTTPException(422, "Priority must be Normal, High, or Critical")
    drawing_id, revision_id = _drawing(db, project, data.drawing_id, data.drawing_revision_id)
    row.number = data.number.strip()
    row.number_key = _key(data.number)
    row.company_id = data.company_id
    row.discipline_id = data.discipline_id
    row.area_id = data.area_id
    row.requested_by = data.requested_by
    row.required_date = data.required_date
    row.priority = data.priority
    row.purpose = data.purpose
    row.remarks = data.remarks
    row.drawing_id = drawing_id
    row.drawing_revision_id = revision_id
    _stamp(row, user)


def create_requirement(db, user, project, data):
    permitted(db, user, project, "procurement.requirement.edit", data.discipline_id)
    row = MaterialRequirement(
        organization_id=user.organization_id,
        project_id=project.id,
        number=data.number.strip(),
        number_key=_key(data.number),
        company_id=data.company_id,
        discipline_id=data.discipline_id,
        requested_by=data.requested_by,
        required_date=data.required_date,
        purpose=data.purpose,
        status="Draft",
        created_by=user.id,
    )
    _fill_requirement(db, user, project, row, data)
    db.add(row)
    db.flush()
    _sync_requirement_lines(db, user, project, row, data.lines)
    _sync_work_packages(db, project, row, getattr(data, "work_package_ids", None))
    _event(db, user, project.id, "requirements", row.id, "created")
    return row


def update_requirement(db, user, project, row, data):
    if row.status not in EDITABLE:
        raise HTTPException(409, "An approved or in-review requirement is changed by amendment or review")
    permitted(db, user, project, "procurement.requirement.edit", data.discipline_id)
    _bump(db, MaterialRequirement, row.id, data.version, {})
    db.refresh(row)
    _fill_requirement(db, user, project, row, data)
    _sync_requirement_lines(db, user, project, row, data.lines)
    _sync_work_packages(db, project, row, getattr(data, "work_package_ids", None))
    _event(db, user, project.id, "requirements", row.id, "updated")
    return row


def _policy(db, project, record_type, company_id, discipline_id):
    rows = db.query(ApprovalPolicy).filter_by(
        project_id=project.id, record_type=record_type, active=True
    )
    matches = []
    for policy in rows:
        if policy.company_id and policy.company_id != company_id:
            continue
        if policy.discipline_id and policy.discipline_id != discipline_id:
            continue
        matches.append(policy)
    if not matches:
        return None
    return max(
        matches,
        key=lambda policy: (
            int(policy.company_id is not None) + int(policy.discipline_id is not None),
            policy.policy_version,
        ),
    )


def _project_member(db, project, person):
    if not person or not person.active or person.organization_id != project.organization_id:
        return False
    return (
        db.query(Membership)
        .filter_by(project_id=project.id, user_id=person.id, active=True)
        .first()
        is not None
    )


def _saved_reviewer(db, policy, user_id):
    if not policy:
        return False
    stages = db.query(PolicyStage).filter_by(policy_id=policy.id).all()
    return any(
        db.query(PolicyReviewer).filter_by(stage_id=stage.id, user_id=user_id).first()
        for stage in stages
    )


def _authority(db, project, policy):
    stages = db.query(PolicyStage).filter_by(policy_id=policy.id).all()
    if not stages:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    for stage in stages:
        reviewers = db.query(PolicyReviewer).filter_by(stage_id=stage.id).all()
        if not reviewers:
            raise HTTPException(422, "Mandatory reviewer configuration is missing")
        for reviewer in reviewers:
            person = db.get(User, reviewer.user_id)
            if not _project_member(db, project, person):
                name = person.name if person else "Reviewer"
                raise HTTPException(422, f"{name} is not an active member of this project")


def validate_procurement_policy(db, project, data):
    for stage in data.stages:
        ids = stage.reviewer_ids or [row.user_id for row in stage.reviewers]
        if not ids:
            raise HTTPException(422, "Mandatory reviewer configuration is missing")
        for user_id in ids:
            person = db.get(User, user_id)
            if not _project_member(db, project, person):
                name = person.name if person else "Reviewer"
                raise HTTPException(422, f"{name} is not an active member of this project")


def submit_requirement(db, user, project, row, version):
    permitted(db, user, project, "procurement.requirement.edit", row.discipline_id)
    if row.status != "Draft":
        raise HTTPException(409, "Only a draft can be submitted. A returned requirement is resubmitted.")
    if not db.query(RequirementLine).filter_by(requirement_id=row.id).first():
        raise HTTPException(422, "Add at least one item line")
    policy = _policy(db, project, REQUIREMENT, row.company_id, row.discipline_id)
    if not policy:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    _authority(db, project, policy)
    _bump(db, MaterialRequirement, row.id, version, {"status": "Pending Approval", "modified_by": user.id})
    db.refresh(row)
    from app.domains.engineering_documents.workflow import submit_record

    instance = submit_record(db, user, project, REQUIREMENT, row.id, row.version, policy)
    _event(db, user, project.id, "requirements", row.id, "submitted")
    return instance


def resubmit_requirement(db, user, project, row, data):
    permitted(db, user, project, "procurement.requirement.edit", row.discipline_id)
    if row.status != "Returned for Correction":
        raise HTTPException(409, "Only a returned requirement can be resubmitted")
    instance = (
        db.query(ReviewInstance)
        .filter_by(project_id=project.id, record_type=REQUIREMENT, record_id=row.id, state="Returned")
        .order_by(ReviewInstance.id.desc())
        .first()
    )
    if not instance:
        raise HTTPException(409, "The returned review was not found")
    _resubmit(db, user, project, instance, data.review_version)
    _bump(
        db,
        MaterialRequirement,
        row.id,
        data.version,
        {"status": "Pending Approval", "modified_by": user.id},
    )
    db.refresh(row)
    instance.record_version = row.version
    _event(db, user, project.id, "requirements", row.id, "resubmitted")


def _resubmit(db, user, project, instance, review_version):
    from app.domains.coordination.service import _close_tasks, _lock, _open_stage, current_stage

    _lock(db, instance, review_version)
    _close_tasks(db, project.id, instance.record_type, instance.id, "Completed", "Corrected and resubmitted")
    instance.attempt += 1
    instance.state = "In Review"
    _open_stage(db, user, project, instance, current_stage(db, instance))


def cancel_requirement(db, user, project, row, data):
    permitted(db, user, project, "procurement.requirement.edit", row.discipline_id)
    if row.status == "Pending Approval":
        raise HTTPException(409, "This requirement is in review. The reviewer returns or rejects it first.")
    if row.status not in {"Draft", "Returned for Correction", "Approved"}:
        raise HTTPException(409, "This requirement cannot be cancelled")
    for line in db.query(RequirementLine).filter_by(requirement_id=row.id):
        reserved, _approved = _reserved(db, line.id)
        if reserved > 0:
            raise HTTPException(409, "Cancel the purchase requisitions that use this requirement first")
    _bump(
        db,
        MaterialRequirement,
        row.id,
        data.version,
        {"status": "Cancelled", "cancellation_reason": data.reason, "modified_by": user.id},
    )
    from .models import FulfilmentAllocation

    for line in db.query(RequirementLine).filter_by(requirement_id=row.id):
        for allocation in db.query(FulfilmentAllocation).filter_by(requirement_line_id=line.id, status="Active"):
            allocation.status = "Released"
            allocation.release_reason = "Requirement cancelled"
            allocation.released_by = user.id
            allocation.released_at = now()
    _stop_review(db, project, REQUIREMENT, row.id)
    _event(db, user, project.id, "requirements", row.id, "cancelled", data.reason)


def amend_requirement(db, user, project, row, data):
    permitted(db, user, project, "procurement.requirement.edit", row.discipline_id)
    if row.status != "Approved":
        raise HTTPException(409, "Only an approved requirement can be amended")
    if row.impact_state == "Review required":
        raise HTTPException(409, "Record the drawing impact decision before amending this requirement")
    _bump(db, MaterialRequirement, row.id, data.version, {"status": "Draft", "modified_by": user.id})
    _event(db, user, project.id, "requirements", row.id, "amended", data.reason)


def decide_impact(db, user, project, row, data):
    policy = _policy(db, project, REQUIREMENT, row.company_id, row.discipline_id)
    if not _saved_reviewer(db, policy, user.id) and not role_allows(
        db, user, project, "procurement.requirement.review", row.discipline_id
    ):
        raise HTTPException(403, "Drawing impact is decided by the saved requirement reviewer")
    if row.impact_state != "Review required":
        raise HTTPException(409, "This requirement has no open drawing impact")
    detail = f"{data.decision} Affected quantity {data.affected_quantity}."
    _bump(
        db,
        MaterialRequirement,
        row.id,
        data.version,
        {"impact_state": "Reviewed", "impact_decision": detail, "modified_by": user.id},
    )
    from app.domains.coordination.service import _close_tasks

    _close_tasks(db, project.id, REQUIREMENT, row.id, "Completed", detail)
    _event(db, user, project.id, "requirements", row.id, "impact-reviewed", detail)


def allocate_stock(db, user, project, row):
    permitted(db, user, project, "procurement.requirement.edit", row.discipline_id)
    raise HTTPException(
        409,
        "Accepted stock is not filled from this screen. Use Stock to accept, reserve, or issue it, or send this requirement to a purchase requisition.",
    )


def create_requisition(db, user, project, data):
    if any(line.estimated_price is not None for line in data.lines) and not _commercial(db, user, project):
        raise HTTPException(403, "Estimated prices require commercial access")
    if any(line.estimated_price is not None for line in data.lines) and not data.currency.strip():
        raise HTTPException(422, "Enter a currency when estimated prices are used")
    if not any(line.estimated_price is not None for line in data.lines) and data.currency.strip():
        raise HTTPException(422, "Currency is entered only when estimated prices are used")
    first = db.get(RequirementLine, data.lines[0].requirement_line_id)
    if not first:
        raise HTTPException(422, "Choose a requirement line")
    source = db.get(MaterialRequirement, first.requirement_id)
    if not source or source.project_id != project.id:
        raise HTTPException(422, "Choose a requirement line from this project")
    permitted(db, user, project, "procurement.requisition.edit", source.discipline_id)
    _company_actor(db, user, project, source.company_id)
    row = PurchaseRequisition(
        organization_id=user.organization_id,
        project_id=project.id,
        number=data.number.strip(),
        number_key=_key(data.number),
        company_id=source.company_id,
        discipline_id=source.discipline_id,
        requested_by=user.id,
        request_date=data.request_date,
        required_delivery=data.required_delivery,
        justification=data.justification,
        currency=data.currency.strip().upper(),
        status="Draft",
        created_by=user.id,
    )
    db.add(row)
    db.flush()
    _sync_requisition_lines(db, user, project, row, data.lines)
    _event(db, user, project.id, "requisitions", row.id, "created")
    return row


def update_requisition(db, user, project, row, data):
    if row.status not in EDITABLE:
        raise HTTPException(409, "An approved or in-review requisition stays as submitted")
    commercial = _commercial(db, user, project)
    if any(line.estimated_price is not None for line in data.lines) and not commercial:
        raise HTTPException(403, "Estimated prices require commercial access")
    permitted(db, user, project, "procurement.requisition.edit", row.discipline_id)
    _company_actor(db, user, project, row.company_id)
    if commercial:
        if any(line.estimated_price is not None for line in data.lines) and not data.currency.strip():
            raise HTTPException(422, "Enter a currency when estimated prices are used")
        if not any(line.estimated_price is not None for line in data.lines) and data.currency.strip():
            raise HTTPException(422, "Currency is entered only when estimated prices are used")
    preserved_prices = {
        line.requirement_line_id: line.estimated_price
        for line in db.query(RequisitionLine).filter_by(requisition_id=row.id)
    }
    preserved_currency = row.currency or ""
    _bump(db, PurchaseRequisition, row.id, data.version, {"modified_by": user.id})
    db.refresh(row)
    row.request_date = data.request_date
    row.required_delivery = data.required_delivery
    row.justification = data.justification
    row.currency = data.currency.strip().upper() if commercial else preserved_currency
    row.number = data.number.strip()
    row.number_key = _key(data.number)
    _stamp(row, user)
    for line in db.query(RequisitionLine).filter_by(requisition_id=row.id):
        db.delete(line)
    db.flush()
    _sync_requisition_lines(
        db,
        user,
        project,
        row,
        data.lines,
        preserve_prices=None if commercial else preserved_prices,
    )
    _event(db, user, project.id, "requisitions", row.id, "updated")
    return row


def _sync_requisition_lines(db, user, project, row, incoming, preserve_prices=None):
    seen = set()
    for index, item in enumerate(incoming, start=1):
        source_line = db.get(RequirementLine, item.requirement_line_id)
        source = db.get(MaterialRequirement, source_line.requirement_id) if source_line else None
        if not source or source.project_id != project.id:
            raise HTTPException(422, "Choose a requirement line from this project")
        if source.status != "Approved":
            raise HTTPException(422, "A purchase requisition uses approved requirement quantities")
        if source.impact_state == "Review required":
            raise HTTPException(
                422, "Resolve the drawing impact review before purchasing more of this requirement"
            )
        if source.company_id != row.company_id or source.discipline_id != row.discipline_id:
            raise HTTPException(422, "All lines on a requisition use the same company and discipline")
        if item.requirement_line_id in seen:
            raise HTTPException(422, "A requirement line can appear only once on a requisition")
        seen.add(item.requirement_line_id)
        from .fulfilment import allocated

        reserved, _approved = _reserved(db, source_line.id)
        if _amount(item.quantity) > _amount(source_line.quantity) - reserved - allocated(db, source_line.id):
            raise HTTPException(422, "Quantity is above the remaining approved requirement")
        stored = db.get(ProcurementItem, source_line.item_id)
        price = item.estimated_price
        if preserve_prices is not None:
            price = preserve_prices.get(item.requirement_line_id)
        db.add(
            RequisitionLine(
                requisition_id=row.id,
                requirement_line_id=source_line.id,
                position=index,
                item_id=source_line.item_id,
                specification=source_line.specification,
                quantity=item.quantity,
                unit=source_line.unit,
                estimated_price=price,
            )
        )
        _ = stored


def submit_requisition(db, user, project, row, version):
    permitted(db, user, project, "procurement.requisition.edit", row.discipline_id)
    if row.status != "Draft":
        raise HTTPException(409, "Only a draft can be submitted. A returned requisition is resubmitted.")
    if not db.query(RequisitionLine).filter_by(requisition_id=row.id).first():
        raise HTTPException(422, "Add at least one item line")
    policy = _policy(db, project, REQUISITION, row.company_id, row.discipline_id)
    if not policy:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    _authority(db, project, policy)
    _bump(db, PurchaseRequisition, row.id, version, {"status": "Pending Approval", "modified_by": user.id})
    db.refresh(row)
    from app.domains.engineering_documents.workflow import submit_record

    instance = submit_record(db, user, project, REQUISITION, row.id, row.version, policy)
    _event(db, user, project.id, "requisitions", row.id, "submitted")
    return instance


def resubmit_requisition(db, user, project, row, data):
    permitted(db, user, project, "procurement.requisition.edit", row.discipline_id)
    _company_actor(db, user, project, row.company_id)
    if row.status != "Returned for Correction":
        raise HTTPException(409, "Only a returned requisition can be resubmitted")
    instance = (
        db.query(ReviewInstance)
        .filter_by(project_id=project.id, record_type=REQUISITION, record_id=row.id, state="Returned")
        .order_by(ReviewInstance.id.desc())
        .first()
    )
    if not instance:
        raise HTTPException(409, "The returned review was not found")
    policy = _policy(db, project, REQUISITION, row.company_id, row.discipline_id)
    if not policy:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    _authority(db, project, policy)
    _resubmit(db, user, project, instance, data.review_version)
    # Pick up the latest active route and treat the resubmitter as the submitter
    # so the named reviewer is not blocked by self-approval rules.
    instance.policy_id = policy.id
    instance.submitted_by = user.id
    _bump(
        db,
        PurchaseRequisition,
        row.id,
        data.version,
        {"status": "Pending Approval", "modified_by": user.id},
    )
    db.refresh(row)
    instance.record_version = row.version
    _event(db, user, project.id, "requisitions", row.id, "resubmitted")


def cancel_requisition(db, user, project, row, data):
    permitted(db, user, project, "procurement.requisition.edit", row.discipline_id)
    if row.status == "Pending Approval":
        raise HTTPException(409, "This requisition is in review. The reviewer returns or rejects it first.")
    if row.status not in {"Draft", "Returned for Correction", "Approved"}:
        raise HTTPException(409, "This requisition cannot be cancelled")
    _company_actor(db, user, project, row.company_id)
    _bump(
        db,
        PurchaseRequisition,
        row.id,
        data.version,
        {"status": "Cancelled", "cancellation_reason": data.reason, "modified_by": user.id},
    )
    _stop_review(db, project, REQUISITION, row.id)
    _event(db, user, project.id, "requisitions", row.id, "cancelled", data.reason)


def _stop_review(db, project, record_type, record_id):
    from app.domains.coordination.service import _close_tasks

    instance = (
        db.query(ReviewInstance)
        .filter_by(project_id=project.id, record_type=record_type, record_id=record_id)
        .order_by(ReviewInstance.id.desc())
        .first()
    )
    if instance and instance.state in {"In Review", "Returned"}:
        instance.state = "Cancelled"
        _close_tasks(db, project.id, record_type, instance.id, "Cancelled", "Record cancelled")


def on_review_event(db, user, instance, project, action, stage=None):
    if instance.record_type == REQUIREMENT:
        row = db.get(MaterialRequirement, instance.record_id)
        resource = "requirements"
    elif instance.record_type == REQUISITION:
        row = db.get(PurchaseRequisition, instance.record_id)
        resource = "requisitions"
    elif instance.record_type == PURCHASE_ORDER:
        from .models import PurchaseOrder

        row = db.get(PurchaseOrder, instance.record_id)
        resource = "orders"
    else:
        return
    if not row or row.project_id != project.id:
        return
    if action == "advance" and stage is not None:
        from app.domains.coordination.jobs import notify

        submitter = db.get(User, instance.submitted_by)
        notify(
            db,
            submitter,
            "Review moved to the next person",
            f"{row.number} moved to {stage.name}.",
            f"review:{instance.id}:attempt:{instance.attempt}:advanced:{stage.id}:buyer",
            project_id=project.id,
            link=f"/projects/{project.id}",
        )
        _event(db, user, project.id, resource, row.id, "advanced", stage.name)
        return
    if action == "return":
        status = "Returned for Correction"
    elif action == "reject":
        status = "Rejected"
    elif instance.state == "Approved":
        status = "Approved"
    else:
        return
    row.status = status
    row.version += 1
    _stamp(row, user)
    _event(db, user, project.id, resource, row.id, status.casefold(), "")


def review_guard(db, user, instance, stage, action):
    if instance.record_type not in {REQUIREMENT, REQUISITION, PURCHASE_ORDER} or action == "comment":
        return
    project = db.get(Project, instance.project_id)
    if instance.record_type == PURCHASE_ORDER:
        from .models import PurchaseOrder

        row = db.get(PurchaseOrder, instance.record_id)
    else:
        model = MaterialRequirement if instance.record_type == REQUIREMENT else PurchaseRequisition
        row = db.get(model, instance.record_id)
    if not row or row.project_id != project.id:
        raise HTTPException(403, "Approval requires the saved reviewer for this document")
    if instance.record_type == PURCHASE_ORDER:
        if action == "approve":
            from .orders import guard_limit

            guard_limit(db, user, stage, row)
        return
    if action != "approve" or instance.record_type != REQUISITION:
        return
    reviewer = db.query(PolicyReviewer).filter_by(stage_id=stage.id, user_id=user.id).first()
    if not reviewer or reviewer.approval_limit is None:
        return
    lines = db.query(RequisitionLine).filter_by(requisition_id=row.id).all()
    if any(line.estimated_price is None for line in lines):
        raise HTTPException(
            422, "Enter an estimated price on every line before this monetary limit can be applied"
        )
    total = sum(_amount(line.quantity) * _amount(line.estimated_price) for line in lines)
    if total > Decimal(reviewer.approval_limit):
        raise HTTPException(
            403, "This requisition total is above your approval limit in the requisition currency"
        )


def flag_revision(db, user, revision_id, detail):
    revision = db.get(DrawingRevision, revision_id)
    drawing = db.get(Drawing, revision.drawing_id) if revision else None
    if not drawing:
        return
    project = db.get(Project, drawing.project_id)
    rows = db.query(MaterialRequirement).filter(
        MaterialRequirement.project_id == drawing.project_id,
        MaterialRequirement.status.notin_(["Cancelled", "Rejected"]),
    )
    from app.domains.coordination.service import _task

    for row in rows:
        if row.drawing_revision_id != revision.id and row.drawing_id != drawing.id:
            continue
        if row.impact_state == "Review required":
            continue
        row.impact_state = "Review required"
        row.impact_detail = detail
        _stamp(row, user)
        _event(db, user, row.project_id, "requirements", row.id, "impact", detail)
        _task(
            db,
            user,
            project,
            row.requested_by,
            f"Review drawing impact on {row.number}",
            REQUIREMENT,
            row.id,
        )


def _item_duplicate_matches(db, user, data, exclude_id=None):
    from .catalogue_duplicates import (
        classify_match,
        normalize_mpn,
        normalize_text,
        summarize_duplicates,
        technical_fingerprint,
    )

    mpn = normalize_mpn(getattr(data, "manufacturer_part_number", ""))
    fingerprint = technical_fingerprint(
        category=data.category,
        subcategory=getattr(data, "subcategory", "") or "",
        description=data.description,
        material_grade=getattr(data, "material_grade", "") or "",
        manufacturer_part_number=getattr(data, "manufacturer_part_number", "") or "",
        size_rating=getattr(data, "size_rating", "") or "",
        specification_attributes=getattr(data, "specification_attributes", None) or {},
    )
    description_key = normalize_text(data.description)
    query = db.query(ProcurementItem).filter_by(organization_id=user.organization_id)
    if exclude_id:
        query = query.filter(ProcurementItem.id != exclude_id)
    matches = []
    for candidate in query.limit(500).all():
        match_type = classify_match(
            candidate,
            mpn=mpn,
            fingerprint=fingerprint,
            description_key=description_key,
            category=data.category,
        )
        if not match_type:
            continue
        matches.append(
            {
                "id": candidate.id,
                "code": candidate.code,
                "description": candidate.description,
                "category": candidate.category,
                "subcategory": getattr(candidate, "subcategory", "") or "",
                "manufacturer_part_number": getattr(candidate, "manufacturer_part_number", "") or "",
                "approval_status": candidate.approval_status,
                "match_type": match_type,
            }
        )
    return summarize_duplicates(matches)


def save_item(db, user, data, row=None):
    _org_master(db, user)
    if row is not None and row.approval_status not in ("Draft", "Rejected"):
        raise HTTPException(409, "Submitted or approved item cannot be edited; use a controlled revision")
    if data.item_type not in ITEM_TYPES:
        raise HTTPException(422, "Type must be Material, Equipment, Consumable, Spare, or Tool")
    if getattr(data, "shelf_life_applicable", False) and not getattr(data, "shelf_life_days", None):
        raise HTTPException(422, "Enter shelf life days when shelf life applies")
    if getattr(data, "hsn_code", "") and not str(data.hsn_code).replace(" ", "").isalnum():
        raise HTTPException(422, "HSN code must be alphanumeric")
    preferred_vendor_id = getattr(data, "preferred_vendor_id", None)
    if preferred_vendor_id:
        vendor = db.get(Vendor, preferred_vendor_id)
        if not vendor or vendor.organization_id != user.organization_id or not vendor.active:
            raise HTTPException(422, "Choose an active preferred vendor from this organization")
    disciplines = []
    for discipline_id in data.discipline_ids:
        discipline = db.get(Discipline, discipline_id)
        if not discipline:
            raise HTTPException(422, "Choose a discipline from the organization list")
        disciplines.append(discipline.id)
    discipline_name = getattr(data, "specification_discipline", "") or ""
    if not discipline_name and disciplines:
        first = db.get(Discipline, disciplines[0])
        discipline_name = first.name if first else ""
    from .catalogue_masters import enforce_item_masters
    if discipline_name:
        enforce_item_masters(db, user, data, discipline_name)
    duplicates = _item_duplicate_matches(db, user, data, exclude_id=row.id if row is not None else None)
    if duplicates["blocking"] and not getattr(data, "allow_duplicate", False):
        codes = ", ".join(item["code"] for item in duplicates["exact"][:5])
        raise HTTPException(409, f"Exact catalogue duplicate detected ({codes}). Confirm allow_duplicate to proceed or revise identity fields.")
    # Preserve structured attributes on unrelated legacy item edits. Category changes
    # invalidate old attributes unless a matching new set is provided.
    existing_category = row.category if row is not None else None
    existing_attributes = dict(row.specification_attributes or {}) if row is not None else {}
    if row is None:
        row = ProcurementItem(
            organization_id=user.organization_id,
            code=data.code.strip(),
            code_key=_key(data.code),
            description=data.description,
            item_type=data.item_type,
            category=data.category,
            base_unit=data.base_unit.strip(),
            created_by=user.id,
        )
        db.add(row)
    else:
        _bump(db, ProcurementItem, row.id, data.version, {})
        db.refresh(row)
        row.code = data.code.strip()
        row.code_key = _key(data.code)
        row.description = data.description
        row.item_type = data.item_type
        row.category = data.category
        row.base_unit = data.base_unit.strip()
    purchase_unit = (getattr(data, "purchase_unit", "") or "").strip() or row.base_unit
    issue_unit = (getattr(data, "issue_unit", "") or "").strip() or row.base_unit
    row.purchase_unit = purchase_unit
    row.issue_unit = issue_unit
    row.subcategory = (getattr(data, "subcategory", "") or "").strip()
    row.discipline_ids = disciplines
    flag_modified(row, "discipline_ids")
    from .specification_templates import validate_configured_attributes
    template_row = db.query(CatalogueSpecificationTemplate).filter_by(
        organization_id=user.organization_id,
        discipline=data.specification_discipline or discipline_name,
        category=data.category,
        active=True,
    ).first() if (data.specification_discipline or discipline_name) else None
    if data.specification_attributes:
        # Use the discipline name from the authenticated organization's master, never a client-supplied ID alone.
        if not data.specification_discipline:
            raise HTTPException(422, "Select a specification discipline")
        names = [db.get(Discipline, discipline_id).name for discipline_id in disciplines]
        if data.specification_discipline not in names and not template_row:
            raise HTTPException(422, "Specification discipline must match the item's assigned discipline")
        row.specification_attributes = validate_configured_attributes(db, user.organization_id, data.specification_discipline, data.category, data.specification_attributes)
        row.specification_template_version = template_row.version if template_row else None
    elif existing_category == data.category and existing_attributes:
        row.specification_attributes = existing_attributes
    else:
        row.specification_attributes = {}
        row.specification_template_version = None
    row.specification = data.specification
    row.size_rating = data.size_rating
    row.material_grade = (getattr(data, "material_grade", "") or "").strip()
    row.applicable_standard = (getattr(data, "applicable_standard", "") or "").strip()
    row.approved_make = data.approved_make
    row.model_number = (getattr(data, "model_number", "") or "").strip()
    row.manufacturer_part_number = (getattr(data, "manufacturer_part_number", "") or "").strip()
    row.hsn_code = (getattr(data, "hsn_code", "") or "").strip()
    row.inspection_required = data.inspection_required
    row.tracking_required = data.tracking_required
    row.mtc_required = bool(getattr(data, "mtc_required", False))
    row.calibration_certificate_required = bool(getattr(data, "calibration_certificate_required", False))
    row.shelf_life_applicable = bool(getattr(data, "shelf_life_applicable", False))
    row.shelf_life_days = getattr(data, "shelf_life_days", None) if row.shelf_life_applicable else None
    row.preferred_vendor_id = preferred_vendor_id
    row.active = False if row.approval_status == "Draft" else data.active
    _stamp(row, user)
    db.flush()
    return row


def item_payload(row, db):
    from .catalogue_masters import qa_documents_for_item

    disciplines = [
        db.get(Discipline, discipline_id).name
        for discipline_id in (row.discipline_ids or [])
        if db.get(Discipline, discipline_id)
    ]
    return {
        "id": row.id,
        "code": row.code,
        "description": row.description,
        "item_type": row.item_type,
        "category": row.category,
        "subcategory": getattr(row, "subcategory", "") or "",
        "discipline_ids": row.discipline_ids or [],
        "disciplines": disciplines,
        "specification": row.specification,
        "specification_attributes": row.specification_attributes or {},
        "approval_status": row.approval_status,
        "size_rating": row.size_rating,
        "material_grade": getattr(row, "material_grade", "") or "",
        "applicable_standard": getattr(row, "applicable_standard", "") or "",
        "base_unit": row.base_unit,
        "purchase_unit": getattr(row, "purchase_unit", "") or row.base_unit,
        "issue_unit": getattr(row, "issue_unit", "") or row.base_unit,
        "approved_make": row.approved_make,
        "model_number": getattr(row, "model_number", "") or "",
        "manufacturer_part_number": getattr(row, "manufacturer_part_number", "") or "",
        "hsn_code": getattr(row, "hsn_code", "") or "",
        "inspection_required": row.inspection_required,
        "tracking_required": row.tracking_required,
        "mtc_required": bool(getattr(row, "mtc_required", False)),
        "calibration_certificate_required": bool(getattr(row, "calibration_certificate_required", False)),
        "shelf_life_applicable": bool(getattr(row, "shelf_life_applicable", False)),
        "shelf_life_days": getattr(row, "shelf_life_days", None),
        "preferred_vendor_id": getattr(row, "preferred_vendor_id", None),
        "specification_template_version": getattr(row, "specification_template_version", None),
        "qa_document_types": qa_documents_for_item(db, row.organization_id, disciplines[0] if disciplines else "", row.category),
        "active": row.active,
        "version": row.version,
        "created_by": row.created_by,
        "created_at": row.created_at.isoformat() if getattr(row, "created_at", None) else None,
        "modified_by": getattr(row, "modified_by", None),
        "modified_at": row.modified_at.isoformat() if getattr(row, "modified_at", None) else None,
        "files": _files(db, "catalogue_items", row.id),
    }


def save_vendor(db, user, data, row=None):
    _org_master(db, user)
    if data.approval_status not in VENDOR_STATES:
        raise HTTPException(422, "Vendor approval must be Pending, Approved, or Suspended")
    company = db.get(Company, data.company_id)
    if not company or company.organization_id != user.organization_id or not company.active:
        raise HTTPException(422, "Choose a company from this organization")
    if data.email and "@" not in data.email:
        raise HTTPException(422, "Enter a valid vendor email")
    if row is None:
        row = Vendor(
            organization_id=user.organization_id,
            code=data.code.strip(),
            code_key=_key(data.code),
            company_id=data.company_id,
            created_by=user.id,
        )
        db.add(row)
    else:
        _bump(db, Vendor, row.id, data.version, {})
        db.refresh(row)
        row.code = data.code.strip()
        row.code_key = _key(data.code)
        row.company_id = data.company_id
    row.categories = data.categories
    row.contact_name = data.contact_name
    row.email = data.email
    row.phone = data.phone
    row.address = data.address
    row.tax_details = data.tax_details
    row.approval_status = data.approval_status
    row.active = data.active
    _stamp(row, user)
    db.flush()
    return row


def vendor_payload(db, user, row, project=None):
    show = user.organization_admin or (project and _commercial(db, user, project))
    if project is None:
        show = user.organization_admin or _any_commercial(db, user)
    return {
        "id": row.id,
        "code": row.code,
        "company_id": row.company_id,
        "company": db.get(Company, row.company_id).name,
        "categories": row.categories,
        "contact_name": row.contact_name,
        "email": row.email,
        "phone": row.phone,
        "address": row.address,
        "tax_details": row.tax_details if show else "",
        "prices_visible": show,
        "approval_status": row.approval_status,
        "active": row.active,
        "version": row.version,
        "files": _files(db, "vendors", row.id),
    }


def _any_commercial(db, user):
    if user.organization_admin:
        return True
    memberships = db.query(Membership).filter_by(user_id=user.id, active=True)
    for member in memberships:
        project = db.get(Project, member.project_id)
        if project and project.organization_id == user.organization_id and _commercial(db, user, project):
            return True
    return False


def save_warehouse(db, user, project, data, row=None):
    permitted(db, user, project, "procurement.master.edit")
    _member(db, project, data.officer_id)
    _area(db, project, data.area_id)
    if row is None:
        row = Warehouse(
            organization_id=user.organization_id,
            project_id=project.id,
            code=data.code.strip(),
            code_key=_key(data.code),
            name=data.name,
            officer_id=data.officer_id,
            created_by=user.id,
        )
        db.add(row)
    else:
        _bump(db, Warehouse, row.id, data.version, {})
        db.refresh(row)
        row.code = data.code.strip()
        row.code_key = _key(data.code)
        row.name = data.name
        row.officer_id = data.officer_id
    row.location = data.location
    row.area_id = data.area_id
    row.active = data.active
    _stamp(row, user)
    db.flush()
    return row


def warehouse_payload(db, row):
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "location": row.location,
        "area_id": row.area_id,
        "officer_id": row.officer_id,
        "officer": _name(db, row.officer_id),
        "active": row.active,
        "version": row.version,
    }


def save_conversion(db, user, data):
    _org_master(db, user)
    if _key(data.from_unit) == _key(data.to_unit):
        raise HTTPException(422, "The two units must be different")
    row = UnitConversion(
        organization_id=user.organization_id,
        from_unit=data.from_unit.strip(),
        to_unit=data.to_unit.strip(),
        factor=data.factor,
        created_by=user.id,
    )
    db.add(row)
    db.flush()
    return row


def dashboard(db, project):
    requirements = db.query(MaterialRequirement).filter_by(project_id=project.id)
    requisitions = db.query(PurchaseRequisition).filter_by(project_id=project.id)
    open_requirements = requirements.filter(MaterialRequirement.status.in_(OPEN_REQUIREMENTS)).count()
    pending = requisitions.filter(PurchaseRequisition.status == "Pending Approval").count()
    impact = requirements.filter(MaterialRequirement.impact_state == "Review required").count()
    awaiting = (
        db.query(RequestForQuotation)
        .filter_by(project_id=project.id, status="Awaiting quotations")
        .count()
    )
    from .invoices import recorded_count
    from .models import PurchaseOrder
    from .orders import unordered_line_count
    from .stock import on_hand_count

    orders_pending = (
        db.query(PurchaseOrder).filter_by(project_id=project.id, status="Pending Approval").count()
    )
    awaiting_receipt = (
        db.query(PurchaseOrder)
        .filter_by(project_id=project.id)
        .filter(PurchaseOrder.status.in_(("Issued", "Acknowledged", "Partially Received")))
        .count()
    )
    return {
        "open_requirements": open_requirements,
        "requisitions_pending": pending,
        "approved_lines_not_ordered": unordered_line_count(db, project.id),
        "drawing_impacts": impact,
        "rfqs_awaiting": awaiting,
        "orders_pending": orders_pending,
        "orders_awaiting_receipt": awaiting_receipt,
        "stock_on_hand": on_hand_count(db, project.id),
        "invoices_recorded": recorded_count(db, project.id),
        "ordering_note": "An issued purchase order is the commitment. A goods receipt records delivery. Stock changes only when that receipt is accepted into a warehouse. A vendor invoice records the supplier bill and does not match or pay it.",
        "unavailable": [
            {"name": "Inspection", "reason": "QA/QC is not open. Warehouse acceptance is a stores posting, not a quality decision."},
            {"name": "Invoice matching and payment", "reason": "Financials owns matching and payment. This screen only records the supplier bill."},
            {"name": "Planning work packages", "reason": "Planning is not open. Required dates stay on the requirement."},
            {"name": "BOQ and quantities from geometry", "reason": "MTO lines come from explicit drawing tables, imports or manual entry. Quantities are never calculated from drawing geometry or scale."},
            {"name": "Site execution and billing", "reason": "Installed quantities and payment execution are not open."},
        ],
    }


def context(db, user, project):
    from app.domains.construction.models import WorkPackage
    from app.domains.planning_reporting.workspace_models import PlanningWBS

    companies = [
        {"id": company.id, "name": company.name}
        for link in db.query(ProjectCompany).filter_by(project_id=project.id, active=True)
        if (company := db.get(Company, link.company_id)) and company.active
    ]
    people = []
    seen = set()
    for member in db.query(Membership).filter_by(project_id=project.id, active=True):
        if member.user_id in seen:
            continue
        person = db.get(User, member.user_id)
        if person and person.active:
            seen.add(person.id)
            people.append({"id": person.id, "name": person.name, "company_id": member.company_id})
    def eligible(action):
        chosen = []
        for row in people:
            person = db.get(User, row["id"])
            if role_allows(db, person, project, action):
                chosen.append({"id": row["id"], "name": row["name"]})
        return chosen
    drawings = []
    for drawing in (
        db.query(Drawing)
        .filter_by(project_id=project.id, record_status="Active")
        .order_by(Drawing.number)
        .limit(200)
    ):
        revisions = db.query(DrawingRevision).filter_by(drawing_id=drawing.id).order_by(DrawingRevision.id)
        drawings.append(
            {
                "id": drawing.id,
                "number": drawing.number,
                "revisions": [{"id": revision.id, "code": revision.code} for revision in revisions],
            }
        )
    return {
        "disciplines": [
            {"id": row.discipline_id, "name": db.get(Discipline, row.discipline_id).name}
            for row in db.query(ProjectDiscipline).filter_by(project_id=project.id, active=True)
            if db.get(Discipline, row.discipline_id)
        ],
        "companies": companies,
        "people": people,
        "requirement_reviewers": eligible("procurement.requirement.review"),
        "requisition_reviewers": eligible("procurement.requisition.review"),
        "areas": [
            {"id": area.id, "code": area.code, "name": area.name}
            for area in db.query(Area).filter_by(project_id=project.id, active=True)
        ],
        "warehouses": [
            {"id": row.id, "code": row.code, "name": row.name}
            for row in db.query(Warehouse).filter_by(project_id=project.id, active=True).order_by(Warehouse.code)
        ],
        "drawings": drawings,
        "wbs": [
            {"id": row.id, "code": row.code, "name": row.name}
            for row in db.query(PlanningWBS)
            .filter_by(project_id=project.id, organization_id=project.organization_id, active=True)
            .order_by(PlanningWBS.code)
        ],
        "work_packages": [
            {"id": row.id, "code": row.code, "name": row.name, "status": row.status,
             "editable": row.status == "Draft" and row.release_state != "In Review"}
            for row in db.query(WorkPackage)
            .filter(WorkPackage.project_id == project.id, WorkPackage.status != "Cancelled")
            .order_by(WorkPackage.code)
        ],
        "items": [
            item_payload(row, db)
            for row in db.query(ProcurementItem)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(ProcurementItem.code)
        ],
        "access": {
            "requirement_edit": holds(db, user, project, "procurement.requirement.edit"),
            "requirement_review": role_allows(db, user, project, "procurement.requirement.review"),
            "requisition_edit": holds(db, user, project, "procurement.requisition.edit"),
            "requisition_review": role_allows(db, user, project, "procurement.requisition.review"),
            "rfq_edit": holds(db, user, project, "procurement.rfq.edit"),
            "po_edit": holds(db, user, project, "procurement.po.edit"),
            "receipt_edit": holds(db, user, project, "procurement.receipt.edit"),
            "stock_edit": holds(db, user, project, "procurement.stock.edit"),
            "invoice_edit": holds(db, user, project, "procurement.invoice.edit"),
            "commercial": _commercial(db, user, project),
            "master": user.organization_admin or role_allows(db, user, project, "procurement.master.edit"),
            "admin": user.organization_admin,
        },
    }


def _workbook(headers):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# Map common alternate headings onto the import template names.
_IMPORT_HEADER_ALIASES = {
    "item_code": "code",
    "itemcode": "code",
    "item_type": "type",
    "itemtype": "type",
    "item_description": "description",
    "desc": "description",
    "category_name": "category",
    "item_category": "category",
    "material_category": "category",
    "cat": "category",
    "sub_category": "subcategory",
    "sub_cat": "subcategory",
    "base_unit": "unit",
    "uom": "unit",
    "unit_of_measure": "unit",
    "purchase_uom": "purchase_unit",
    "issue_uom": "issue_unit",
    "size_rating": "size",
    "applicable_standard": "standard",
    "approved_make": "make",
    "manufacturer": "make",
    "brand": "make",
    "model_number": "model",
    "manufacturer_part_number": "mpn",
    "mpn_number": "mpn",
    "hsn_code": "hsn",
    "discipline": "disciplines",
    "discipline_name": "disciplines",
    "discipline_names": "disciplines",
}


def _normalize_import_header(raw):
    text = str(raw or "").strip().casefold()
    text = "_".join(text.replace("-", " ").split())
    return _IMPORT_HEADER_ALIASES.get(text, text)


def _row_value(header, row, index):
    if index < 0 or index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value).strip()


def _resolve_import_value(header, row, canonical, extras=()):
    """Read a field by canonical name, aliases, or a close header label."""
    for name in (canonical, *extras):
        if name in header:
            text = _row_value(header, row, header.index(name))
            if text:
                return text
    for index, label in enumerate(header):
        if not label or label == canonical:
            continue
        # Avoid matching subcategory when asking for category.
        if canonical == "category" and "sub" in label.split("_"):
            continue
        tokens = set(label.split("_"))
        if canonical in tokens or label.endswith(f"_{canonical}") or label.startswith(f"{canonical}_"):
            text = _row_value(header, row, index)
            if text:
                return text
    return ""


def item_template(db=None, user=None):
    """Blank import workbook plus an Allowed_Masters sheet for clients.

    When db/user are provided, Allowed_Masters is filled from this organization's
    active catalogue masters and discipline list so uploads can copy exact values.
    """
    from .catalogue_exchange import IMPORT_COLUMNS
    from .models import (
        CatalogueCategory,
        CatalogueManufacturer,
        CatalogueMaterialGrade,
        CatalogueStandard,
        CatalogueSubcategory,
        CatalogueUom,
    )

    book = Workbook()
    sheet = book.active
    sheet.title = "Catalogue"
    headers = list(IMPORT_COLUMNS)
    sheet.append(headers)

    disciplines = []
    categories = []
    uoms = []
    grades = []
    standards = []
    manufacturers = []
    if db is not None and user is not None:
        disciplines = [row.name for row in db.query(Discipline).order_by(Discipline.name).all()]
        categories = (
            db.query(CatalogueCategory)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(CatalogueCategory.discipline, CatalogueCategory.name)
            .all()
        )
        uoms = (
            db.query(CatalogueUom)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(CatalogueUom.code)
            .all()
        )
        grades = (
            db.query(CatalogueMaterialGrade)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(CatalogueMaterialGrade.code)
            .all()
        )
        standards = (
            db.query(CatalogueStandard)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(CatalogueStandard.code)
            .all()
        )
        manufacturers = (
            db.query(CatalogueManufacturer)
            .filter_by(organization_id=user.organization_id, active=True)
            .order_by(CatalogueManufacturer.code)
            .all()
        )

    sample_discipline = disciplines[0] if disciplines else "Civil"
    sample_category = next((row.name for row in categories if _key(row.discipline) == _key(sample_discipline)), None)
    if sample_category is None and categories:
        sample_category = categories[0].name
        sample_discipline = categories[0].discipline
    if sample_category is None:
        sample_category = "Concrete"
    sample_unit = uoms[0].code if uoms else "KG"

    sample = {name: "" for name in headers}
    sample.update(
        {
            "code": "SAMPLE-001",
            "description": "Sample item — delete this row before Confirm",
            "type": "Material",
            "category": sample_category,
            "subcategory": "",
            "disciplines": sample_discipline,
            "unit": sample_unit,
            "purchase_unit": sample_unit,
            "issue_unit": sample_unit,
            "inspection": "no",
            "tracking": "no",
            "mtc": "no",
            "calibration": "no",
            "shelf_life": "no",
            "specification_attributes": "{}",
        }
    )
    sheet.append([sample[name] for name in headers])

    masters = book.create_sheet("Allowed_Masters")
    masters.append(["Copy values from this sheet into the Catalogue sheet. Do not rename Catalogue headers."])
    masters.append([])
    masters.append(["disciplines (use in column disciplines)"])
    masters.append(["discipline"])
    for name in disciplines:
        masters.append([name])
    if not disciplines:
        masters.append(["(none — configure organization disciplines first)"])

    masters.append([])
    masters.append(["categories (use in column category / subcategory)"])
    masters.append(["discipline", "category", "subcategory"])
    if categories:
        for category in categories:
            subs = (
                db.query(CatalogueSubcategory)
                .filter_by(organization_id=user.organization_id, category_id=category.id, active=True)
                .order_by(CatalogueSubcategory.name)
                .all()
            )
            if subs:
                for sub in subs:
                    masters.append([category.discipline, category.name, sub.name])
            else:
                masters.append([category.discipline, category.name, ""])
    else:
        masters.append(["(none — seed or create Catalogue reference masters before client upload)"])

    masters.append([])
    masters.append(["uoms (use in unit / purchase_unit / issue_unit)"])
    masters.append(["code", "name"])
    if uoms:
        for row in uoms:
            masters.append([row.code, row.name])
    else:
        masters.append(["(none — any unit text allowed until UOM masters are added)"])

    masters.append([])
    masters.append(["material_grades (optional column material_grade)"])
    masters.append(["code", "name"])
    if grades:
        for row in grades:
            masters.append([row.code, row.name])
    else:
        masters.append(["(none)"])

    masters.append([])
    masters.append(["standards (optional column standard)"])
    masters.append(["code", "title"])
    if standards:
        for row in standards:
            masters.append([row.code, row.title])
    else:
        masters.append(["(none)"])

    masters.append([])
    masters.append(["manufacturers (optional column make)"])
    masters.append(["code", "name"])
    if manufacturers:
        for row in manufacturers:
            masters.append([row.code, row.name])
    else:
        masters.append(["(none)"])

    guide = book.create_sheet("Instructions")
    guide.append(["Client catalogue import pack"])
    guide.append(["1. Open Allowed_Masters and copy exact discipline / category / unit values."])
    guide.append(["2. Enter items only on the Catalogue sheet. Keep row-1 headers unchanged."])
    guide.append(["3. Required columns: code, description, type, category, disciplines, unit."])
    guide.append(["4. type must be Material, Equipment, Consumable, Spare, or Tool."])
    guide.append(["5. When masters exist, category and unit must match Allowed_Masters exactly."])
    guide.append(["6. Delete SAMPLE-001 before upload Confirm (or change its code)."])
    guide.append(["7. After import, items are Draft — Submit and Approve before purchase use."])
    guide.append(["8. yes/no columns: inspection, tracking, mtc, calibration, shelf_life."])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def requirement_template():
    return _workbook(
        [
            "number",
            "company",
            "discipline",
            "requested_by",
            "required_date",
            "priority",
            "purpose",
            "remarks",
            "item_code",
            "quantity",
            "unit",
            "line_date",
            "tag",
            "engineering_reference",
        ]
    )


def _sheet(data):
    try:
        # data_only=False keeps typed cell values (safer for freshly filled imports).
        book = load_workbook(BytesIO(data), data_only=False)
    except Exception as exc:
        raise HTTPException(422, "The file is not a readable Excel workbook") from exc
    sheet = None
    for candidate in book.worksheets:
        title = (candidate.title or "").casefold()
        if title in {"catalogue", "items", "import", "item catalogue"}:
            sheet = candidate
            break
    if sheet is None:
        sheet = book.active
    if (sheet.title or "").casefold() == "instructions" and len(book.worksheets) > 1:
        sheet = next((item for item in book.worksheets if (item.title or "").casefold() != "instructions"), sheet)
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(422, "The workbook is empty")
    header = [_normalize_import_header(cell) for cell in rows[0]]
    # Drop trailing blank header placeholders from Excel's used range.
    while header and header[-1] == "":
        header.pop()
    if not header:
        raise HTTPException(422, "Row 1 must contain the import column headers")
    body = []
    width = len(header)
    for row in rows[1:]:
        values = list(row or ())
        if len(values) < width:
            values.extend([None] * (width - len(values)))
        body.append(tuple(values[:width]))
    return header, body


def _cell(header, row, name):
    if name not in header:
        return ""
    index = header.index(name)
    if index >= len(row):
        return ""
    value = row[index]
    if value is None:
        return ""
    return str(value).strip()


def import_items(db, user, data, mode):
    _org_master(db, user)
    from .catalogue_import_validation import duplicate_codes, MAX_IMPORT_ROWS
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(413, "Catalogue import exceeds 5 MB")
    if mode not in {"preview", "confirm", "update"}:
        raise HTTPException(422, "Import mode must be preview, confirm, or update")
    header, rows = _sheet(data)
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(422, f"Catalogue import supports at most {MAX_IMPORT_ROWS} rows")
    required = ["code", "description", "type", "category", "disciplines", "unit"]
    missing = [name for name in required if name not in header]
    if missing:
        detected = ", ".join(name for name in header if name) or "(none)"
        raise HTTPException(
            422,
            "Missing columns: "
            + ", ".join(missing)
            + f". Detected headers: {detected}. Download a fresh blank template and keep row 1 unchanged.",
        )
    disciplines = {_key(row.name): row for row in db.query(Discipline)}
    duplicate_rows = duplicate_codes(header, rows)
    preview = []
    errors = []
    for index, row in enumerate(rows, start=2):
        if not any(cell not in (None, "") for cell in row):
            continue
        code = _cell(header, row, "code")
        names = [part.strip() for part in _cell(header, row, "disciplines").split(",") if part.strip()]
        discipline_ids = []
        row_errors = []
        if not code:
            row_errors.append("code is required")
        if index in duplicate_rows:
            row_errors.append(f"duplicate item code in workbook (rows {', '.join(map(str, duplicate_rows[index]))})")
        item_type = _resolve_import_value(header, row, "type", ("item_type",))
        description = _resolve_import_value(header, row, "description", ("item_description", "desc"))
        category = _resolve_import_value(
            header, row, "category", ("category_name", "item_category", "material_category", "cat")
        )
        subcategory = _resolve_import_value(header, row, "subcategory", ("sub_category", "sub_cat"))
        # Common mistake: value typed only in subcategory — use it as category.
        if not category and subcategory:
            category = subcategory
            subcategory = ""
        unit = _resolve_import_value(header, row, "unit", ("base_unit", "uom", "unit_of_measure"))
        if item_type not in ITEM_TYPES:
            row_errors.append("type must be Material, Equipment, Consumable, Spare, or Tool")
        blank_required = [label for label, value in (
            ("description", description),
            ("category", category),
            ("unit", unit),
        ) if not value]
        if blank_required:
            row_errors.append(
                "blank required field(s): "
                + ", ".join(blank_required)
                + " — fill column header category (not only subcategory). Example: category=Concrete"
            )
        if not names:
            row_errors.append("at least one discipline is required (column disciplines, e.g. Civil)")
        for name in names:
            found = disciplines.get(_key(name))
            if not found:
                row_errors.append(f"discipline {name} was not found")
            else:
                discipline_ids.append(found.id)
        inspection = _cell(header, row, "inspection").casefold()
        tracking = _cell(header, row, "tracking").casefold()
        mtc = _cell(header, row, "mtc").casefold()
        calibration = _cell(header, row, "calibration").casefold()
        shelf_life = _cell(header, row, "shelf_life").casefold()
        for label, value in (
            ("inspection", inspection),
            ("tracking", tracking),
            ("mtc", mtc),
            ("calibration", calibration),
            ("shelf_life", shelf_life),
        ):
            if value not in {"", "yes", "no"}:
                row_errors.append(f"{label} must be yes or no")
        shelf_life_days_raw = _cell(header, row, "shelf_life_days")
        shelf_life_days = None
        if shelf_life == "yes":
            if not shelf_life_days_raw.isdigit() or int(shelf_life_days_raw) < 1:
                row_errors.append("shelf_life_days must be a positive whole number when shelf_life is yes")
            else:
                shelf_life_days = int(shelf_life_days_raw)
        elif shelf_life_days_raw:
            row_errors.append("shelf_life_days is only allowed when shelf_life is yes")
        spec_attributes, spec_error = _import_specification_attributes(header, row)
        if spec_error:
            row_errors.append(spec_error)
        existing = db.query(ProcurementItem).filter_by(
            organization_id=user.organization_id, code_key=_key(code)
        ).first() if code else None
        action = "update" if existing else "create"
        if existing and mode != "update":
            row_errors.append("this item code already exists")
        if existing and mode == "update" and existing.approval_status not in ("Draft", "Rejected"):
            row_errors.append("approved or submitted item cannot be updated by import; use controlled amendment")
        candidate = {
            "row": index,
            "action": action,
            "code": code,
            "discipline_ids": discipline_ids,
            "description": description,
            "item_type": item_type,
            "category": category,
            "subcategory": subcategory,
            "specification": _cell(header, row, "specification"),
            "size_rating": _resolve_import_value(header, row, "size", ("size_rating",)),
            "material_grade": _cell(header, row, "material_grade"),
            "applicable_standard": _resolve_import_value(header, row, "standard", ("applicable_standard",)),
            "base_unit": unit,
            "purchase_unit": _resolve_import_value(header, row, "purchase_unit", ("purchase_uom",)) or unit,
            "issue_unit": _resolve_import_value(header, row, "issue_unit", ("issue_uom",)) or unit,
            "approved_make": _cell(header, row, "make"),
            "model_number": _cell(header, row, "model"),
            "manufacturer_part_number": _cell(header, row, "mpn"),
            "hsn_code": _cell(header, row, "hsn"),
            "inspection_required": inspection == "yes",
            "tracking_required": tracking == "yes",
            "mtc_required": mtc == "yes",
            "calibration_certificate_required": calibration == "yes",
            "shelf_life_applicable": shelf_life == "yes",
            "shelf_life_days": shelf_life_days,
            "specification_discipline": names[0] if names else "",
            "specification_attributes": spec_attributes,
            "version": existing.version if existing else None,
        }
        if not row_errors and names:
            from .catalogue_masters import enforce_item_masters

            try:
                enforce_item_masters(db, user, ItemInPayload(candidate), names[0])
            except HTTPException as exc:
                detail = exc.detail if isinstance(exc.detail, str) else "Catalogue master validation failed"
                row_errors.append(detail)
        if row_errors:
            errors.append({"row": index, "messages": row_errors})
        preview.append(candidate)
    if errors:
        if mode == "preview":
            return {"errors": errors, "rows": preview, "saved": 0}
        raise HTTPException(422, _import_message(errors))
    if mode == "preview":
        return {"errors": [], "rows": preview, "saved": 0}
    saved = 0
    for item in preview:
        payload = ItemInPayload(item)
        existing = db.query(ProcurementItem).filter_by(
            organization_id=user.organization_id, code_key=_key(item["code"])
        ).first()
        if existing:
            payload.version = existing.version
            save_item(db, user, payload, existing)
        else:
            save_item(db, user, payload, None)
        saved += 1
    return {"errors": [], "rows": preview, "saved": saved}


def _import_specification_attributes(header, row):
    import json

    attributes = {}
    raw = _cell(header, row, "specification_attributes")
    if raw:
        try:
            parsed = json.loads(raw)
        except Exception:
            return {}, "specification_attributes must be valid JSON object"
        if not isinstance(parsed, dict):
            return {}, "specification_attributes must be a JSON object"
        attributes.update(parsed)
    for name in header:
        if not name.startswith("spec."):
            continue
        key = name[5:].strip()
        if not key:
            continue
        value = _cell(header, row, name)
        if value == "":
            continue
        if value.casefold() in {"true", "false"}:
            attributes[key] = value.casefold() == "true"
        else:
            try:
                attributes[key] = int(value) if "." not in value else float(value)
            except ValueError:
                attributes[key] = value
    return attributes, ""


class ItemInPayload:
    def __init__(self, item):
        self.code = item["code"]
        self.description = item["description"]
        self.item_type = item["item_type"]
        self.category = item["category"]
        self.subcategory = item.get("subcategory", "")
        self.discipline_ids = item["discipline_ids"]
        self.specification = item["specification"]
        self.size_rating = item["size_rating"]
        self.material_grade = item.get("material_grade", "")
        self.applicable_standard = item.get("applicable_standard", "")
        self.base_unit = item["base_unit"]
        self.purchase_unit = item.get("purchase_unit", "") or item["base_unit"]
        self.issue_unit = item.get("issue_unit", "") or item["base_unit"]
        self.approved_make = item["approved_make"]
        self.model_number = item.get("model_number", "")
        self.manufacturer_part_number = item.get("manufacturer_part_number", "")
        self.hsn_code = item.get("hsn_code", "")
        self.inspection_required = item["inspection_required"]
        self.tracking_required = item["tracking_required"]
        self.mtc_required = item.get("mtc_required", False)
        self.calibration_certificate_required = item.get("calibration_certificate_required", False)
        self.shelf_life_applicable = item.get("shelf_life_applicable", False)
        self.shelf_life_days = item.get("shelf_life_days")
        self.preferred_vendor_id = item.get("preferred_vendor_id")
        self.specification_discipline = item.get("specification_discipline", "")
        self.specification_attributes = item.get("specification_attributes") or {}
        self.allow_duplicate = item.get("allow_duplicate", False)
        self.active = True
        self.version = item.get("version")


def import_requirements(db, user, project, data, mode):
    permitted(db, user, project, "procurement.requirement.edit")
    if mode not in {"preview", "confirm"}:
        raise HTTPException(422, "Requirement import creates new records. Use preview or confirm.")
    header, rows = _sheet(data)
    required = ["number", "company", "discipline", "requested_by", "required_date", "priority", "purpose", "item_code", "quantity", "unit"]
    missing = [name for name in required if name not in header]
    if missing:
        raise HTTPException(422, "Missing columns: " + ", ".join(missing))
    companies = {
        _key(link_company.name): link_company
        for link in db.query(ProjectCompany).filter_by(project_id=project.id, active=True)
        if (link_company := db.get(Company, link.company_id))
    }
    disciplines = {
        _key(db.get(Discipline, row.discipline_id).name): row.discipline_id
        for row in db.query(ProjectDiscipline).filter_by(project_id=project.id, active=True)
        if db.get(Discipline, row.discipline_id)
    }
    people = {}
    for member in db.query(Membership).filter_by(project_id=project.id, active=True):
        person = db.get(User, member.user_id)
        if person:
            people[_key(person.email)] = person
    items = {
        _key(row.code): row
        for row in db.query(ProcurementItem).filter_by(organization_id=user.organization_id, active=True, approval_status="Approved")
    }
    grouped = {}
    errors = []
    for index, row in enumerate(rows, start=2):
        if not any(cell not in (None, "") for cell in row):
            continue
        number = _cell(header, row, "number")
        messages = []
        company = companies.get(_key(_cell(header, row, "company")))
        discipline_id = disciplines.get(_key(_cell(header, row, "discipline")))
        person = people.get(_key(_cell(header, row, "requested_by")))
        item = items.get(_key(_cell(header, row, "item_code")))
        if not number:
            messages.append("number is required")
        if not company:
            messages.append("company is not assigned to this project")
        if not discipline_id:
            messages.append("discipline is not enabled on this project")
        if not person:
            messages.append("requested_by must be a project member email")
        if _cell(header, row, "priority") not in PRIORITIES:
            messages.append("priority must be Normal, High, or Critical")
        if not _cell(header, row, "purpose"):
            messages.append("purpose is required")
        if not item:
            messages.append("item_code was not found")
        try:
            quantity = Decimal(_cell(header, row, "quantity") or "0")
            if quantity <= 0:
                messages.append("quantity must be greater than zero")
        except Exception:
            quantity = Decimal("0")
            messages.append("quantity must be a number")
        try:
            required_date = _parse_date(_cell(header, row, "required_date"))
            line_date = _parse_date(_cell(header, row, "line_date") or _cell(header, row, "required_date"))
        except ValueError:
            required_date = None
            line_date = None
            messages.append("dates must be YYYY-MM-DD")
        if item and discipline_id and discipline_id not in (item.discipline_ids or []):
            messages.append("item is not configured for this discipline")
        if item and _cell(header, row, "unit") and _key(_cell(header, row, "unit")) != _key(item.base_unit):
            try:
                _unit(db, user.organization_id, item.base_unit, _cell(header, row, "unit"))
            except HTTPException:
                messages.append("configure a unit conversion for this unit")
        existing = (
            db.query(MaterialRequirement)
            .filter_by(project_id=project.id, number_key=_key(number))
            .first()
            if number
            else None
        )
        if existing:
            messages.append("this requirement number already exists")
        if messages:
            errors.append({"row": index, "messages": messages})
            continue
        bucket = grouped.setdefault(
            _key(number),
            {
                "number": number,
                "company_id": company.id,
                "discipline_id": discipline_id,
                "requested_by": person.id,
                "required_date": required_date,
                "priority": _cell(header, row, "priority"),
                "purpose": _cell(header, row, "purpose"),
                "remarks": _cell(header, row, "remarks"),
                "lines": [],
            },
        )
        if (
            bucket["company_id"] != company.id
            or bucket["discipline_id"] != discipline_id
            or bucket["purpose"] != _cell(header, row, "purpose")
        ):
            errors.append({"row": index, "messages": ["rows for one requirement must share the same header"]})
            continue
        bucket["lines"].append(
            {
                "item_id": item.id,
                "quantity": quantity,
                "unit": _cell(header, row, "unit"),
                "required_date": line_date,
                "tag": _cell(header, row, "tag"),
                "engineering_reference": _cell(header, row, "engineering_reference"),
                "specification": "",
            }
        )
    if errors:
        if mode == "preview":
            return {"errors": errors, "requirements": [], "saved": 0}
        raise HTTPException(422, _import_message(errors))
    if mode == "preview":
        return {"errors": [], "requirements": [{"number": item["number"], "lines": len(item["lines"])} for item in grouped.values()], "saved": 0}
    saved = 0
    for item in grouped.values():
        payload = _requirement_payload(item)
        created = create_requirement(db, user, project, payload)
        created.source = "Import"
        saved += 1
    return {"errors": [], "requirements": [], "saved": saved}


def _import_message(errors):
    parts = [
        f"row {item['row']}: {', '.join(item['messages'])}" for item in errors[:8]
    ]
    return "Import was not saved. " + "; ".join(parts)


def _parse_date(value):
    from datetime import date, datetime as clock

    if isinstance(value, date) and not isinstance(value, clock):
        return value
    text = str(value).strip()
    if " " in text:
        text = text.split(" ")[0]
    return date.fromisoformat(text[:10])


class _RequirementPayload:
    def __init__(self, item):
        self.number = item["number"]
        self.company_id = item["company_id"]
        self.discipline_id = item["discipline_id"]
        self.area_id = None
        self.requested_by = item["requested_by"]
        self.required_date = item["required_date"]
        self.priority = item["priority"]
        self.purpose = item["purpose"]
        self.remarks = item["remarks"]
        self.drawing_id = None
        self.drawing_revision_id = None
        self.lines = [_LinePayload(line) for line in item["lines"]]
        self.version = None


class _LinePayload:
    def __init__(self, line):
        self.id = None
        self.item_id = line["item_id"]
        self.specification = line["specification"]
        self.quantity = line["quantity"]
        self.unit = line["unit"]
        self.required_date = line["required_date"]
        self.tag = line["tag"]
        self.engineering_reference = line["engineering_reference"]


def _requirement_payload(item):
    return _RequirementPayload(item)


from app.domains.coordination.service import REVIEW_GUARDS

REVIEW_GUARDS.append(review_guard)
