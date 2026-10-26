"""Draft-only Excel activity imports. Preview and confirm revalidate identical bytes and DB references."""
import base64
import hashlib
import io
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from openpyxl import load_workbook, Workbook
from app.db import get_db
from app.domains.identity_projects.models import User, Company
from app.domains.identity_projects.security import authenticated
from .workspace_router import access
from .workspace_models import PlanningSchedule, PlanningActivity, PlanningWBS, PlanningCalendar

router=APIRouter(prefix='/api/projects/{pid}/planning', tags=['Planning Excel Import'])
COLUMNS=['activity_code','activity_name','wbs_code','activity_type','duration_hours','company_id','calendar_code','source_type','source_id']
MAX_BYTES=2_000_000
MAX_ROWS=5000

class ExcelImport(BaseModel):
    schedule_id:int
    workbook_base64:str=Field(min_length=1,max_length=3_000_000)
    preview_sha256:str|None=None


def unpack(raw):
    try: data=base64.b64decode(raw, validate=True)
    except Exception: raise HTTPException(422,'Invalid workbook encoding')
    if len(data)>MAX_BYTES: raise HTTPException(413,'Workbook exceeds 2 MB limit')
    if data[:2]!=b'PK': raise HTTPException(422,'Only .xlsx files are supported')
    try:
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        if 'Activities' not in wb.sheetnames: raise ValueError('Workbook needs Activities sheet')
        ws=wb['Activities']
        records=ws.iter_rows(values_only=True)
        header=[str(v).strip().lower() if v is not None else '' for v in next(records,())]
        if len(header)!=len(set(header)): raise ValueError('Duplicate column headings')
        if set(COLUMNS)-set(header): raise ValueError('Missing headings: '+', '.join(sorted(set(COLUMNS)-set(header))))
        rows=[]
        for idx, values in enumerate(records,start=2):
            if all(v is None or str(v).strip()=='' for v in values):continue
            if len(rows)>=MAX_ROWS:raise ValueError('Maximum 5,000 nonblank activity rows')
            rows.append((idx,{key:(values[header.index(key)] if header.index(key)<len(values) else None) for key in COLUMNS}))
        wb.close()
        return data,rows
    except HTTPException:raise
    except Exception as exc:raise HTTPException(422,f'Invalid workbook: {exc}')


def validate(db,pid,org,schedule_id,rows):
    schedule=db.query(PlanningSchedule).filter_by(id=schedule_id,project_id=pid,organization_id=org).one_or_none()
    if not schedule:raise HTTPException(404,'Schedule not found')
    if schedule.status!='Draft':raise HTTPException(409,'Imports are allowed only into draft schedules')
    wbs={x.code:x for x in db.query(PlanningWBS).filter_by(project_id=pid,organization_id=org,active=True).all()}
    calendars={x.code:x for x in db.query(PlanningCalendar).filter_by(project_id=pid,organization_id=org).all()}
    existing={x.code for x in db.query(PlanningActivity).filter_by(project_id=pid,schedule_id=schedule_id).all()}
    companies={x.id for x in db.query(Company).filter_by(organization_id=org).all()}
    errors=[]; ready=[]; seen=set()
    for index,r in rows:
        code=str(r['activity_code'] or '').strip();name=str(r['activity_name'] or '').strip()
        wbs_code=str(r['wbs_code'] or '').strip();kind=str(r['activity_type'] or '').strip()
        calendar_code=str(r['calendar_code'] or '').strip()
        source_type=str(r['source_type'] or '').strip();source_id=str(r['source_id'] or '').strip()
        issues=[]
        if not code or len(code)>80:issues.append('Invalid activity_code')
        if not name or len(name)>240:issues.append('Invalid activity_name')
        if code in existing or code in seen:issues.append('Duplicate activity_code')
        seen.add(code)
        if wbs_code not in wbs:issues.append('Unknown or inactive WBS')
        if calendar_code not in calendars:issues.append('Unknown project calendar')
        if kind not in ('Task','Milestone'):issues.append('Invalid activity type')
        try:
            d=Decimal(str(r['duration_hours']))
            if not d.is_finite() or d!=int(d) or d<0 or d>100000 or (kind=='Task' and d==0) or (kind=='Milestone' and d!=0):raise ValueError()
            hours=int(d)
        except (ValueError,TypeError,InvalidOperation):issues.append('Invalid duration_hours');hours=None
        try:
            company_id=int(r['company_id'])
            if company_id not in companies:raise ValueError()
        except (ValueError,TypeError,OverflowError):issues.append('Unknown company_id');company_id=None
        if bool(source_type)!=bool(source_id):issues.append('Source type and ID must both be populated or both blank')
        if len(source_type)>60 or len(source_id)>100:issues.append('Source reference too long')
        if issues:errors.append({'row':index,'activity_code':code,'errors':issues})
        else:ready.append(dict(organization_id=org,project_id=pid,schedule_id=schedule_id,
                               wbs_id=wbs[wbs_code].id,calendar_id=calendars[calendar_code].id,
                               company_id=company_id,code=code,name=name,activity_type=kind,
                               duration_hours=hours if kind=='Task' else None,
                               source_type=source_type or None,source_id=source_id or None,remarks='Imported from Excel'))
    if not rows:errors.append({'row':1,'activity_code':'','errors':['No activity rows']})
    return errors,ready


def prepare(db,pid,org,body):
    data,rows=unpack(body.workbook_base64)
    errors,ready=validate(db,pid,org,body.schedule_id,rows)
    digest=hashlib.sha256(data).hexdigest()
    return {'total_rows':len(rows),'valid_rows':len(ready),'error_count':len(errors),'errors':errors[:300],
            'sha256':digest,'ready':ready}

@router.post('/imports/activities/excel-preview')
def preview_excel(pid:int,body:ExcelImport,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid,True)
    report=prepare(db,pid,user.organization_id,body)
    return {k:v for k,v in report.items() if k!='ready'} | {'written':0,'preview_only':True}

@router.post('/imports/activities/excel-confirm')
def confirm_excel(pid:int,body:ExcelImport,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid,True)
    report=prepare(db,pid,user.organization_id,body)
    if not body.preview_sha256 or body.preview_sha256!=report['sha256']:
        raise HTTPException(409,'Workbook has changed since preview; preview again')
    if report['error_count']:raise HTTPException(422,detail={'message':'Import validation failed','errors':report['errors']})
    # Lock schedule to serialize imports and prevent concurrent publication while inserting.
    schedule=db.query(PlanningSchedule).filter_by(id=body.schedule_id,project_id=pid,organization_id=user.organization_id).with_for_update().one_or_none()
    if not schedule or schedule.status!='Draft':raise HTTPException(409,'Schedule is no longer editable')
    # Revalidate after locking to reject concurrent duplicated codes.
    _,rows=unpack(body.workbook_base64)
    errors,ready=validate(db,pid,user.organization_id,body.schedule_id,rows)
    if errors:raise HTTPException(409,detail={'message':'Schedule changed; preview again','errors':errors[:300]})
    try:
        db.add_all(PlanningActivity(**data) for data in ready)
        db.commit()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'Concurrent update or duplicate code; preview again')
    except Exception:
        db.rollback();raise
    return {'written':len(ready),'schedule_id':body.schedule_id,'sha256':report['sha256']}

@router.get('/imports/activities/excel-template')
def excel_template(pid:int,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid)
    wb=Workbook();ws=wb.active;ws.title='Activities';ws.append(COLUMNS)
    ws.append(['CV-100','Foundation excavation','CIV','Task',40,1,'STD','',''])
    ws.freeze_panes='A2';ws.auto_filter.ref='A1:I2'
    for col in 'ABCDEFGHI':ws.column_dimensions[col].width=22
    out=io.BytesIO();wb.save(out)
    return {'filename':'planning-activity-template.xlsx','workbook_base64':base64.b64encode(out.getvalue()).decode('ascii'),
            'note':'Example company ID and calendar/WBS codes must be replaced with real project references.'}
