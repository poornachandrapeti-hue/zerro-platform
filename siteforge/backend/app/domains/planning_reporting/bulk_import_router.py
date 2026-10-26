"""Non-destructive Excel imports for dependency networks and new calendars.
Preview returns SHA-256 of input; confirm verifies identical bytes and all DB references.
"""
import base64
import hashlib
import io
from datetime import date
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook, Workbook
from app.db import get_db
from app.domains.identity_projects.models import User
from app.domains.identity_projects.security import authenticated
from .workspace_router import access, draft_schedule
from .workspace_models import PlanningSchedule, PlanningActivity, PlanningDependency, PlanningCalendar
from .scheduling_engine import calculate, ScheduleError

router=APIRouter(prefix='/api/projects/{pid}/planning/imports',tags=['Planning bulk imports'])
MAX_BYTES=2_000_000
MAX_ROWS=5000
COLUMNS={'dependencies':['predecessor_code','successor_code','relation','lag_hours'], 'calendars':['calendar_code','calendar_name','timezone','working_days','shift_start','shift_end','holidays']}

class WorkbookInput(BaseModel):
    workbook_base64:str=Field(min_length=1,max_length=3_000_000)
    schedule_id:int|None=None
    preview_sha256:str|None=None


def parse_workbook(encoded,kind):
    try: data=base64.b64decode(encoded,validate=True)
    except Exception: raise HTTPException(422,'Invalid Base64 workbook')
    if len(data)>MAX_BYTES:raise HTTPException(413,'Workbook exceeds 2 MB')
    if data[:2]!=b'PK':raise HTTPException(422,'Only .xlsx workbooks supported')
    try:
        wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        sheet='Dependencies' if kind=='dependencies' else 'Calendars'
        if sheet not in wb.sheetnames:raise ValueError(f'Missing {sheet} sheet')
        it=wb[sheet].iter_rows(values_only=True)
        header=[str(x).strip().lower() if x is not None else '' for x in next(it,())]
        if len(set(header))!=len(header):raise ValueError('Duplicate headers')
        if set(COLUMNS[kind])-set(header):raise ValueError('Missing columns: '+','.join(sorted(set(COLUMNS[kind])-set(header))))
        rows=[]
        for number,values in enumerate(it,2):
            if all(x is None or str(x).strip()=='' for x in values):continue
            if len(rows)>=MAX_ROWS:raise ValueError('Maximum 5,000 rows')
            rows.append((number,{col:(values[header.index(col)] if header.index(col)<len(values) else None) for col in COLUMNS[kind]}))
        wb.close()
        return data,rows
    except Exception as exc:raise HTTPException(422,f'Invalid workbook: {exc}')


def validate_dependencies(rows,activity_codes,existing_pairs):
    errors=[];prepared=[];pairs=set(existing_pairs)
    for number,r in rows:
        a=str(r['predecessor_code'] or '').strip();b=str(r['successor_code'] or '').strip()
        relation=str(r['relation'] or '').strip().upper();issues=[]
        if a not in activity_codes:issues.append('Unknown predecessor code')
        if b not in activity_codes:issues.append('Unknown successor code')
        if a==b:issues.append('Self dependency')
        if (a,b) in pairs:issues.append('Duplicate relationship')
        if relation not in ('FS','SS','FF','SF'):issues.append('Invalid relation')
        try:
            raw=str(r['lag_hours']);lag=int(raw)
            if str(lag)!=raw.strip() or lag<0 or lag>100000:raise ValueError()
        except (TypeError,ValueError):issues.append('lag_hours must be a nonnegative integer');lag=0
        if issues:errors.append({'row':number,'errors':issues})
        else:prepared.append({'predecessor_id':activity_codes[a], 'successor_id':activity_codes[b],'relation':relation,'lag_hours':lag});pairs.add((a,b))
    if not rows:errors.append({'row':1,'errors':['No dependency rows']})
    return errors,prepared


def validate_calendars(rows,existing_codes):
    errors=[];ready=[];seen=set(existing_codes)
    for number,r in rows:
        code=str(r['calendar_code'] or '').strip();name=str(r['calendar_name'] or '').strip();tz=str(r['timezone'] or '').strip()
        days_string=str(r['working_days'] or '').strip();issues=[]
        if not code or len(code)>80 or code in seen:issues.append('Missing, invalid or duplicate calendar_code')
        if not name or len(name)>200:issues.append('Invalid calendar_name')
        try:ZoneInfo(tz)
        except (ZoneInfoNotFoundError,ValueError,KeyError):issues.append('Invalid timezone')
        try:
            days=[int(x.strip()) for x in days_string.split(',')]
            if not days or len(set(days))!=len(days) or any(x<0 or x>6 for x in days):raise ValueError()
        except (ValueError,TypeError):days=[];issues.append('working_days must be unique weekday numbers 0..6')
        start=str(r['shift_start'] or '').strip();end=str(r['shift_end'] or '').strip()
        import re
        if not all(re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d',x) for x in (start,end)) or start==end:issues.append('Invalid shift times')
        holidays=[]
        try:
            for piece in str(r['holidays'] or '').split(','):
                if piece.strip():holidays.append(date.fromisoformat(piece.strip()).isoformat())
            if len(holidays)!=len(set(holidays)):raise ValueError()
        except ValueError:issues.append('holidays must be unique ISO dates YYYY-MM-DD')
        if issues:errors.append({'row':number,'errors':issues})
        else:ready.append(dict(code=code,name=name,timezone=tz,working_days=days,shift_start=start,shift_end=end,holidays=holidays,version=1));seen.add(code)
    if not rows:errors.append({'row':1,'errors':['No calendar rows']})
    return errors,ready


def prepare(db,pid,org,kind,body):
    data,rows=parse_workbook(body.workbook_base64,kind)
    if kind=='dependencies':
        if body.schedule_id is None:raise HTTPException(422,'schedule_id required for dependencies')
        schedule=db.query(PlanningSchedule).filter_by(id=body.schedule_id,project_id=pid,organization_id=org).one_or_none()
        if not schedule:raise HTTPException(404,'Schedule not found')
        if schedule.status!='Draft':raise HTTPException(409,'Only draft schedules can be changed')
        activities=db.query(PlanningActivity).filter_by(schedule_id=schedule.id,project_id=pid,organization_id=org).all()
        existing=db.query(PlanningDependency).filter_by(schedule_id=schedule.id,project_id=pid).all()
        byid={a.id:a.code for a in activities}
        errors,ready=validate_dependencies(rows,{a.code:a.id for a in activities},{(byid[d.predecessor_id],byid[d.successor_id]) for d in existing})
        if not errors:
            dependencies=[dict(predecessor_id=d.predecessor_id,successor_id=d.successor_id,relation=d.relation,lag_hours=d.lag_hours) for d in existing]+ready
            calendars={c.id:{'working_days':c.working_days,'shift_start':c.shift_start,'shift_end':c.shift_end,'holidays':c.holidays,'timezone':c.timezone} for c in db.query(PlanningCalendar).filter_by(project_id=pid,organization_id=org).all()}
            try:calculate([dict(id=a.id,code=a.code,name=a.name,activity_type=a.activity_type,duration_hours=a.duration_hours,calendar_id=a.calendar_id) for a in activities],dependencies,calendars,schedule.data_date)
            except ScheduleError as exc:errors.append({'row':0,'errors':[f'Network invalid: {exc}']})
    else:
        existing={c.code for c in db.query(PlanningCalendar).filter_by(project_id=pid,organization_id=org).all()}
        errors,ready=validate_calendars(rows,existing)
    return {'total_rows':len(rows),'valid_rows':len(ready),'errors':errors[:300],'error_count':len(errors),'sha256':hashlib.sha256(data).hexdigest()},ready

@router.get('/{kind}/excel-template')
def template(pid:int,kind:str,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid)
    if kind not in COLUMNS:raise HTTPException(404,'Unknown import type')
    wb=Workbook();ws=wb.active;ws.title=kind.title();ws.append(COLUMNS[kind]);ws.freeze_panes='A2'
    if kind=='dependencies':ws.append(['CV-100','CV-110','FS',0])
    else:ws.append(['SITE','Site Working Calendar','Asia/Kolkata','0,1,2,3,4,5','08:00','17:00','2026-12-25'])
    for col in 'ABCDEFG':ws.column_dimensions[col].width=27
    out=io.BytesIO();wb.save(out)
    return {'filename':f'planning-{kind}-template.xlsx','workbook_base64':base64.b64encode(out.getvalue()).decode()}

@router.post('/{kind}/excel-preview')
def preview(pid:int,kind:str,body:WorkbookInput,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid,True)
    if kind not in COLUMNS:raise HTTPException(404,'Unknown import type')
    report,_=prepare(db,pid,user.organization_id,kind,body)
    return report|{'preview_only':True,'written':0}

@router.post('/{kind}/excel-confirm')
def confirm(pid:int,kind:str,body:WorkbookInput,user:User=Depends(authenticated),db:Session=Depends(get_db)):
    access(db,user,pid,True)
    if kind not in COLUMNS:raise HTTPException(404,'Unknown import type')
    if kind=='dependencies':
        locked=db.query(PlanningSchedule).filter_by(id=body.schedule_id,project_id=pid,organization_id=user.organization_id).with_for_update().one_or_none()
        if not locked or locked.status!='Draft':raise HTTPException(409,'Schedule is not an editable draft')
    else:
        # Lock the project to serialize concurrent inserts for this project's calendar codes.
        from app.domains.identity_projects.models import Project
        db.query(Project).filter_by(id=pid,organization_id=user.organization_id).with_for_update().one()
    report,ready=prepare(db,pid,user.organization_id,kind,body)
    if not body.preview_sha256 or body.preview_sha256!=report['sha256']:raise HTTPException(409,'Workbook changed since preview')
    if report['error_count']:raise HTTPException(422,detail=report)
    try:
        for data in ready:
            common=dict(project_id=pid,organization_id=user.organization_id)
            if kind=='dependencies':db.add(PlanningDependency(**common,schedule_id=body.schedule_id,**data))
            else:db.add(PlanningCalendar(**common,**data))
        db.commit()
    except IntegrityError:
        db.rollback();raise HTTPException(409,'Concurrent change or duplicate; preview again')
    except Exception:
        db.rollback();raise
    return {'written':len(ready),'sha256':report['sha256'],'kind':kind}
