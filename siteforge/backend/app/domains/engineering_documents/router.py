import io
import json
from datetime import date
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import quote_sheetname

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.db import get_db
from app.domains.identity_projects.models import Company, Discipline, User
from app.domains.identity_projects.router import audit, commit, project
from app.domains.identity_projects.security import admin, authenticated
from . import dashboard
from .models import (
    Drawing,
    DrawingComment,
    DrawingRevision,
    Transmittal,
    TransmittalItem,
    TransmittalRecipient,
    WrenchLink,
    WrenchSetting,
)
from .schemas import (
    CancelIn,
    CommentClose,
    DisciplineIn,
    CommentIn,
    CommentRespond,
    DispatchIn,
    DrawingIn,
    DrawingLinkIn,
    DrawingUpdate,
    EngineeringSettingIn,
    ExternalDecisionIn,
    PreparationIn,
    ReleaseIn,
    TransmittalIn,
    WrenchIn,
    WrenchLinkIn,
)
from .workflow import (
    acknowledge_transmittal,
    add_comment,
    add_revision,
    cancel_drawing,
    close_comment,
    create_drawing,
    create_transmittal,
    detail,
    drawing_recipients,
    visible_transmittals,
    dispatch_transmittal,
    edit_drawing,
    export_drawings,
    filtered_drawings,
    import_rows,
    link_drawing,
    link_wrench,
    list_categories,
    matches_measure,
    my_tasks,
    public_drawing,
    record_client_decision,
    release_revision,
    respond_comment,
    revision_file,
    save_engineering_setting,
    save_wrench,
    sees_unreleased,
    start_preparation,
    submit_revision,
    summary,
    withdraw_revision,
    _drawing,
    _require,
)

router = APIRouter(prefix="/api")


@router.post("/disciplines", status_code=201)
def add_discipline(
    data: DisciplineIn, u: User = Depends(admin), db: Session = Depends(get_db)
):
    if db.query(Discipline).filter_by(name=data.name).first():
        raise HTTPException(409, "Discipline already exists")
    row = Discipline(name=data.name)
    db.add(row)
    db.flush()
    audit(db, u, "create", "disciplines", row.id, {"name": row.name})
    commit(db)
    return {"id": row.id, "name": row.name}


def _project(db, user, pid):
    found = project(db, user, pid)
    _require(db, user, found, "drawings.view")
    return found


def _visible(db, user, found, pairs, bucket):
    rows = []
    for drawing, latest in pairs:
        try:
            _require(db, user, found, "drawings.view", drawing.discipline_id)
        except HTTPException:
            continue
        if not sees_unreleased(db, user, found):
            from .workflow import _released_revision

            released = _released_revision(db, drawing.id)
            if not released:
                continue
            latest = released
        if bucket and public_drawing(db, user, found, drawing, latest)["bucket"] != bucket:
            continue
        rows.append((drawing, latest))
    return rows


@router.get("/projects/{pid}/drawings")
def drawings(
    pid: int,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
    discipline_id: int | None = None,
    company_id: int | None = None,
    owner_id: int | None = None,
    bucket: str = "",
    measure: str = "",
    planned_from: date | None = None,
    planned_to: date | None = None,
    category_id: int | None = None,
    area_id: int | None = None,
    priority: str = "",
    source: str = "",
    record_status: str = "Active",
    date_kind: str = "planned_preparation",
):
    found = _project(db, u, pid)
    pairs = filtered_drawings(
        db,
        found,
        discipline_id,
        company_id,
        owner_id,
        planned_from,
        planned_to,
        category_id,
        area_id,
        priority,
        source,
        record_status,
        date_kind,
    )
    visible = _visible(db, u, found, pairs, "")
    listed = _visible(db, u, found, pairs, bucket)
    if measure:
        listed = [
            (drawing, latest)
            for drawing, latest in listed
            if matches_measure(db, drawing, latest, measure)
        ]
    return {
        "summary": summary(db, visible),
        "items": [public_drawing(db, u, found, drawing, latest) for drawing, latest in listed],
    }


@router.get("/projects/{pid}/drawings/export")
def export_register(
    pid: int,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = _project(db, u, pid)
    pairs = filtered_drawings(db, found)
    payload = export_drawings(db, u, found, pairs)
    return Response(payload, media_type="text/csv")


def _dashboard_filters(
    discipline_id: int | None = None,
    area_id: int | None = None,
    drawing_type: str = "",
    engineer_id: int | None = None,
    internal_status: str = "",
    client_status: str = "",
    release_status: str = "",
    required_from: date | None = None,
    required_to: date | None = None,
    overdue_only: bool = False,
    q: str = "",
    record_status: str = "Active",
):
    return {
        "discipline_id": discipline_id,
        "area_id": area_id,
        "drawing_type": drawing_type,
        "engineer_id": engineer_id,
        "internal_status": internal_status,
        "client_status": client_status,
        "release_status": release_status,
        "required_from": required_from,
        "required_to": required_to,
        "overdue_only": overdue_only,
        "q": q,
        "record_status": record_status,
    }


@router.get("/projects/{pid}/engineering-dashboard")
def engineering_dashboard(
    pid: int,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
    filters: dict = Depends(_dashboard_filters),
    metric: str = "",
    drill_discipline: int | None = None,
    sort: str = "number",
    direction: str = "asc",
    page: int = 1,
    size: int = 25,
):
    found = _project(db, u, pid)
    data = dashboard.snapshot(db, u, found, filters)
    return dashboard.page(data, metric, drill_discipline, sort, direction, page, size)


@router.get("/projects/{pid}/engineering-dashboard/export")
def engineering_dashboard_export(
    pid: int,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
    filters: dict = Depends(_dashboard_filters),
    metric: str = "",
    drill_discipline: int | None = None,
):
    found = _project(db, u, pid)
    data = dashboard.snapshot(db, u, found, filters)
    return Response(dashboard.export(data, metric, drill_discipline), media_type="text/csv")


@router.get("/drawing-categories")
def categories(u: User = Depends(authenticated), db: Session = Depends(get_db)):
    rows = list_categories(db, u)
    commit(db)
    return rows


@router.get("/projects/{pid}/engineering-tasks")
def tasks(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = _project(db, u, pid)
    return my_tasks(db, u, found)


@router.get("/projects/{pid}/engineering-settings")
def engineering_settings(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = _project(db, u, pid)
    from .models import EngineeringSetting

    row = db.get(EngineeringSetting, found.id)
    return {"allow_release_with_comments": bool(row and row.allow_release_with_comments)}


@router.put("/projects/{pid}/engineering-settings")
def save_settings(
    pid: int,
    data: EngineeringSettingIn,
    u: User = Depends(admin),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = save_engineering_setting(db, found, data.allow_release_with_comments)
    audit(db, u, "update", "engineering_settings", pid, project_id=pid)
    commit(db)
    return {"allow_release_with_comments": row.allow_release_with_comments}


@router.post("/projects/{pid}/drawings", status_code=201)
def create(
    pid: int, data: DrawingIn, u: User = Depends(authenticated), db: Session = Depends(get_db)
):
    found = project(db, u, pid)
    row = create_drawing(db, u, found, data)
    audit(db, u, "create", "drawings", row.id, {"number": row.number}, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.put("/projects/{pid}/drawings/{did}")
def edit(
    pid: int,
    did: int,
    data: DrawingUpdate,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    edit_drawing(db, u, found, row, data)
    audit(db, u, "update", "drawings", did, project_id=pid)
    commit(db)
    db.refresh(row)
    return detail(db, u, found, row)


@router.post("/projects/{pid}/drawings/{did}/preparation")
def prepare(
    pid: int,
    did: int,
    data: PreparationIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    start_preparation(db, u, found, row, data.version)
    commit(db)
    db.refresh(row)
    return detail(db, u, found, row)


@router.get("/projects/{pid}/drawings/{did}")
def one(pid: int, did: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = project(db, u, pid)
    return detail(db, u, found, _drawing(db, u, found, did))


@router.post("/projects/{pid}/drawings/{did}/revisions", status_code=201)
async def revise(
    pid: int,
    did: int,
    code: str = Form(),
    upload: UploadFile = File(),
    description: str = Form(""),
    purpose: str = Form("Review"),
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    payload = type("Upload", (), {})()
    payload.filename = upload.filename
    payload.content_type = upload.content_type
    payload.data = await upload.read()
    revision = add_revision(db, u, found, row, code, payload, description, purpose)
    audit(db, u, "revise", "drawing_revisions", revision.id, {"code": code}, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.get("/projects/{pid}/drawings/{did}/revisions/{rid}/file")
def download(
    pid: int,
    did: int,
    rid: int,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    if not sees_unreleased(db, u, found) and revision.release_state not in {
        "Released",
        "Superseded",
        "Withdrawn",
    }:
        raise HTTPException(404, "Revision not found")
    stored, content = revision_file(db, revision)
    audit(db, u, "download", "drawing_revisions", rid, project_id=pid)
    commit(db)
    filename = stored.original_name.replace('"', "")
    return Response(
        content=content,
        media_type=stored.content_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/projects/{pid}/drawings/{did}/revisions/{rid}/submit", status_code=201)
def submit(
    pid: int, did: int, rid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    instance = submit_revision(db, u, found, row, revision)
    audit(db, u, "submit", "reviews", instance.id, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.post("/projects/{pid}/drawings/{did}/revisions/{rid}/release")
def release(
    pid: int,
    did: int,
    rid: int,
    data: ReleaseIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    release_revision(db, u, found, row, revision, data.version, data.reason)
    audit(db, u, "release", "drawing_revisions", rid, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.post("/projects/{pid}/drawings/{did}/revisions/{rid}/withdraw")
def withdraw(
    pid: int,
    did: int,
    rid: int,
    data: ReleaseIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    withdraw_revision(db, u, found, row, revision, data.version, data.reason)
    audit(db, u, "withdraw", "drawing_revisions", rid, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.post("/projects/{pid}/drawings/{did}/revisions/{rid}/comments", status_code=201)
def comment(
    pid: int,
    did: int,
    rid: int,
    data: CommentIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    row_comment = add_comment(db, u, found, row, revision, data)
    audit(db, u, "comment", "drawing_comments", row_comment.id, project_id=pid)
    commit(db)
    return {"id": row_comment.id}


@router.post("/projects/{pid}/comments/{cid}/respond")
def respond(
    pid: int,
    cid: int,
    data: CommentRespond,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = db.get(DrawingComment, cid)
    revision = db.get(DrawingRevision, row.revision_id) if row else None
    drawing = db.get(Drawing, revision.drawing_id) if revision else None
    if not drawing or drawing.project_id != pid:
        raise HTTPException(404, "Comment not found")
    respond_comment(db, u, found, drawing, row, data.response, data.version)
    commit(db)
    return {"ok": True}


@router.post("/projects/{pid}/comments/{cid}/close")
def close(
    pid: int,
    cid: int,
    data: CommentClose,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = db.get(DrawingComment, cid)
    revision = db.get(DrawingRevision, row.revision_id) if row else None
    drawing = db.get(Drawing, revision.drawing_id) if revision else None
    if not drawing or drawing.project_id != pid:
        raise HTTPException(404, "Comment not found")
    close_comment(db, u, found, drawing, row, data.evidence, data.version)
    commit(db)
    return {"ok": True}


@router.post("/projects/{pid}/drawings/{did}/cancel")
def cancel(
    pid: int,
    did: int,
    data: CancelIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    cancel_drawing(db, u, found, row, data.reason, data.version)
    audit(db, u, "cancel", "drawings", did, {"reason": data.reason}, project_id=pid)
    commit(db)
    db.refresh(row)
    return detail(db, u, found, row)


@router.post("/projects/{pid}/drawings/{did}/links", status_code=201)
def add_link(
    pid: int,
    did: int,
    data: DrawingLinkIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    link = link_drawing(db, u, found, row, data.related_id)
    audit(db, u, "link", "drawing_links", link.id, project_id=pid)
    commit(db)
    return {"id": link.id, "related_id": link.related_id}


@router.post("/projects/{pid}/drawings/{did}/revisions/{rid}/client-decision")
def client_decision(
    pid: int,
    did: int,
    rid: int,
    data: ExternalDecisionIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    revision = db.get(DrawingRevision, rid)
    if not revision or revision.drawing_id != row.id:
        raise HTTPException(404, "Revision not found")
    record_client_decision(db, u, found, row, revision, data)
    audit(db, u, "client_decision", "drawing_revisions", rid, {"outcome": data.outcome}, project_id=pid)
    commit(db)
    return detail(db, u, found, row)


@router.get("/projects/{pid}/drawing-imports/template")
def drawing_import_template(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = project(db, u, pid)
    from .workflow import _require, IMPORT_COLUMNS
    from app.domains.identity_projects.models import Area
    _require(db, u, found, "drawings.edit")
    book = Workbook()
    sheet = book.active
    sheet.title = "Drawings"
    sheet.append(list(IMPORT_COLUMNS))
    lookup = book.create_sheet("Project Locations")
    lookup.append(["Location code", "Location name"])
    locations = db.query(Area).filter_by(project_id=pid, active=True).order_by(Area.code).all()
    for area in locations:
        lookup.append([area.code, area.name])
    if locations:
        location_col = list(IMPORT_COLUMNS).index("location") + 1
        from openpyxl.utils import get_column_letter
        col = get_column_letter(location_col)
        validation = DataValidation(type="list", formula1=f"{quote_sheetname(lookup.title)}!$A$2:$A${len(locations)+1}", allow_blank=True)
        validation.error = "Choose an existing project location"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{col}2:{col}5001")
    for cell in sheet[1]:
        cell.font = __import__('openpyxl').styles.Font(bold=True)
    book.active = 0
    buffer = io.BytesIO()
    book.save(buffer)
    return Response(content=buffer.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="drawing-import-template.xlsx"', "Cache-Control": "no-store"})


@router.post("/projects/{pid}/drawing-imports", status_code=201)
async def importer(
    pid: int,
    upload: UploadFile = File(),
    dry_run: bool = Form(False),
    mode: str = Form("create"),
    location_mappings: str = Form("{}"),
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    try:
        mappings = json.loads(location_mappings)
    except (ValueError, TypeError):
        raise HTTPException(422, "Location mappings must be valid JSON")
    report = import_rows(db, u, found, upload.filename or "", await upload.read(), mode, mappings)
    report["preview"] = dry_run
    if dry_run:
        db.rollback()
        return report
    audit(db, u, "import", "drawings", pid, {"created": len(report["created"])}, project_id=pid)
    commit(db)
    return report


@router.get("/projects/{pid}/drawing-recipients")
def recipients(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = _project(db, u, pid)
    return drawing_recipients(db, u, found)


@router.get("/projects/{pid}/transmittals")
def transmittals(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = _project(db, u, pid)
    return [_transmittal(db, row) for row in visible_transmittals(db, u, found)]


@router.post("/projects/{pid}/transmittals", status_code=201)
def transmit(
    pid: int, data: TransmittalIn, u: User = Depends(authenticated), db: Session = Depends(get_db)
):
    found = project(db, u, pid)
    row = create_transmittal(db, u, found, data)
    audit(db, u, "create", "transmittals", row.id, project_id=pid)
    commit(db)
    return _transmittal(db, row)


@router.post("/projects/{pid}/transmittals/{tid}/dispatch")
def dispatch(
    pid: int,
    tid: int,
    data: DispatchIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = db.get(Transmittal, tid)
    if not row or row.project_id != pid:
        raise HTTPException(404, "Transmittal not found")
    dispatch_transmittal(db, u, found, row, data.version)
    commit(db)
    db.refresh(row)
    return _transmittal(db, row)


@router.post("/projects/{pid}/transmittals/{tid}/acknowledge")
def acknowledge(
    pid: int, tid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)
):
    found = project(db, u, pid)
    row = db.get(Transmittal, tid)
    if not row or row.project_id != pid:
        raise HTTPException(404, "Transmittal not found")
    acknowledge_transmittal(db, u, found, row)
    commit(db)
    db.refresh(row)
    return _transmittal(db, row)


@router.get("/projects/{pid}/transmittals/{tid}")
def transmittal(
    pid: int, tid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)
):
    found = _project(db, u, pid)
    row = db.get(Transmittal, tid)
    if not row or row.project_id != found.id:
        raise HTTPException(404, "Transmittal not found")
    return _transmittal(db, row)


@router.get("/projects/{pid}/wrench")
def wrench(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    found = _project(db, u, pid)
    row = db.get(WrenchSetting, found.id)
    return {
        "enabled": False if not row else row.enabled,
        "specification": "" if not row else row.specification,
        "credential_label": "" if not row else row.credential_label,
        "status": "disabled",
    }


@router.put("/projects/{pid}/wrench")
def wrench_save(
    pid: int, data: WrenchIn, u: User = Depends(admin), db: Session = Depends(get_db)
):
    found = project(db, u, pid)
    row = save_wrench(db, found, data)
    commit(db)
    return {"enabled": row.enabled, "status": "disabled"}


@router.post("/projects/{pid}/wrench/sync")
def wrench_sync(pid: int, u: User = Depends(authenticated), db: Session = Depends(get_db)):
    _project(db, u, pid)
    raise HTTPException(
        409,
        "Wrench synchronization is disabled. A verified API specification and credentials are required before a connector can be enabled.",
    )


@router.post("/projects/{pid}/drawings/{did}/wrench-link")
def wrench_link(
    pid: int,
    did: int,
    data: WrenchLinkIn,
    u: User = Depends(authenticated),
    db: Session = Depends(get_db),
):
    found = project(db, u, pid)
    row = _drawing(db, u, found, did)
    link = link_wrench(db, u, found, row, data)
    commit(db)
    db.refresh(link)
    return {
        "external_document_id": link.external_document_id,
        "external_revision_id": link.external_revision_id,
        "provenance": link.provenance,
        "conflict_status": link.conflict_status,
        "detail": link.detail,
    }


def _transmittal(db, row: Transmittal):
    items = db.query(TransmittalItem).filter_by(transmittal_id=row.id).all()
    recipients = db.query(TransmittalRecipient).filter_by(transmittal_id=row.id).all()
    lines = []
    for item in items:
        revision = db.get(DrawingRevision, item.revision_id)
        drawing = db.get(Drawing, revision.drawing_id) if revision else None
        lines.append(
            {
                "revision_id": item.revision_id,
                "drawing_id": drawing.id if drawing else None,
                "number": drawing.number if drawing else "",
                "title": drawing.title if drawing else "",
                "code": revision.code if revision else "",
            }
        )
    recipient_company = db.get(Company, row.recipient_company_id) if row.recipient_company_id else None
    sender = db.get(User, row.sender_id)
    return {
        "items": lines,
        "recipient_company": recipient_company.name if recipient_company else "",
        "sender": sender.name if sender else "",
        "id": row.id,
        "number": row.number,
        "purpose": row.purpose,
        "issue_date": row.issue_date,
        "status": row.status,
        "dispatched_at": row.dispatched_at,
        "acknowledged_at": row.acknowledged_at,
        "sender_company_id": row.sender_company_id,
        "recipient_company_id": row.recipient_company_id,
        "expected_response": row.expected_response,
        "method": row.method,
        "external_reference": row.external_reference,
        "version": row.version,
        "revision_ids": [item.revision_id for item in items],
        "recipients": [
            {
                "user_id": item.user_id,
                "name": getattr(db.get(User, item.user_id), "name", ""),
                "acknowledged_at": item.acknowledged_at,
            }
            for item in recipients
        ],
    }
