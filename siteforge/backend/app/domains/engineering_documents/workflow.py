"""Drawing register, revisions, release and the document-control slice for transmittals."""

import csv
import io
import re
from datetime import date, datetime, timezone

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.orm import Session

from app.domains.coordination.jobs import notify
from app.domains.coordination.models import ReviewDecision, ReviewInstance
from app.domains.coordination.service import REVIEW_GUARDS, review_payload
from app.domains.identity_projects.models import Area, Company, Discipline, Membership, Project, User
from app.domains.identity_projects.security import permitted
from .files import read_file, store_file
from .models import (
    Drawing,
    DrawingCategory,
    DrawingComment,
    DrawingLink,
    DrawingRevision,
    EngineeringSetting,
    RevisionEvent,
    StoredFile,
    Transmittal,
    TransmittalItem,
    TransmittalRecipient,
    WrenchLink,
    WrenchSetting,
)

DRAWING_REVIEW = "drawing.revision"
PREPARATION = {"Planned", "In Preparation", "Prepared"}
PRIORITIES = {"Low", "Normal", "High", "Critical"}
SEVERITIES = {"Blocking", "Major", "Minor"}
# One drawing contributes to one bucket. The latest revision wins over preparation.
BUCKETS = (
    "to_prepare",
    "in_preparation",
    "under_review",
    "accepted",
    "accepted_with_comments",
    "returned",
    "rejected",
    "released",
)


def _event(db, user, drawing, revision, kind, detail=""):
    db.add(
        RevisionEvent(
            project_id=drawing.project_id,
            drawing_id=drawing.id,
            revision_id=revision.id,
            kind=kind,
            actor_id=user.id,
            detail=detail,
        )
    )


def _latest(db, drawing_id):
    return (
        db.query(DrawingRevision)
        .filter_by(drawing_id=drawing_id)
        .order_by(DrawingRevision.sequence.desc())
        .first()
    )


def bucket_for(drawing, latest):
    if latest and latest.release_state == "Released":
        return "released"
    if latest and latest.review_state in {"Internal Review", "Client Review"}:
        return "under_review"
    if latest and latest.review_state == "Accepted":
        return "accepted"
    if latest and latest.review_state == "Accepted with Comments":
        return "accepted_with_comments"
    if latest and latest.review_state == "Returned":
        return "returned"
    if latest and latest.review_state == "Rejected":
        return "rejected"
    if drawing.preparation == "Planned" and not latest:
        return "to_prepare"
    return "in_preparation"


TECHNICAL = {"drawings.review", "drawings.release", "drawings.client_review"}


def _membership_allows(db, user, project, action, discipline_id=None):
    from app.domains.identity_projects.models import ProjectDiscipline, Role

    if project.organization_id != user.organization_id:
        return False
    enabled = {
        row.discipline_id
        for row in db.query(ProjectDiscipline).filter_by(project_id=project.id, active=True)
    }
    if discipline_id is not None and discipline_id not in enabled:
        return False
    for member in db.query(Membership).filter_by(
        project_id=project.id, user_id=user.id, active=True
    ):
        role = db.get(Role, member.role_id)
        if not role or action not in (role.permissions or []):
            continue
        if discipline_id is None or member.scope == "ALL" or discipline_id in (member.discipline_ids or []):
            return True
    return False


def _can(db, user, project, action, discipline_id=None):
    if action in TECHNICAL:
        return _membership_allows(db, user, project, action, discipline_id)
    try:
        permitted(db, user, project, action, discipline_id)
        return True
    except HTTPException:
        return False


def sees_unreleased(db, user, project):
    if user.organization_admin:
        return True
    return any(
        _can(db, user, project, action)
        for action in (
            "drawings.edit",
            "drawings.review",
            "drawings.client_review",
            "drawings.release",
            "transmittals.issue",
            "drawings.plan",
        )
    )


def _require(db, user, project, action, discipline_id=None):
    if action in TECHNICAL and not _membership_allows(db, user, project, action, discipline_id):
        raise HTTPException(403, "Technical approval is a separate permission")
    if action in TECHNICAL:
        return
    permitted(db, user, project, action, discipline_id)


def _released_revision(db, drawing_id):
    return (
        db.query(DrawingRevision)
        .filter_by(drawing_id=drawing_id, release_state="Released")
        .order_by(DrawingRevision.sequence.desc())
        .first()
    )


def _drawing(db, user, project, drawing_id):
    row = db.get(Drawing, drawing_id)
    if not row or row.project_id != project.id:
        raise HTTPException(404, "Drawing not found")
    _require(db, user, project, "drawings.view", row.discipline_id)
    if not sees_unreleased(db, user, project) and not _released_revision(db, row.id):
        raise HTTPException(403, "Only released drawings are visible for this role")
    return row


def _same_company(user, drawing):
    if user.organization_admin:
        return
    if user.company_id != drawing.company_id:
        raise HTTPException(403, "This drawing belongs to another company")


def validate_drawing_policy(db, project, data):
    purposes = [(stage.purpose or "").strip() for stage in data.stages]
    if not purposes or purposes[0] != "internal" or len(purposes) > 2:
        raise HTTPException(422, "Drawing review needs an internal stage and an optional client stage")
    if any(item not in {"internal", "client"} for item in purposes):
        raise HTTPException(422, "Drawing stage purpose must be internal or client")
    if len(purposes) == 2 and purposes[1] != "client":
        raise HTTPException(422, "The second drawing stage must be the client review")
    for stage in data.stages:
        ids = [item.user_id for item in stage.reviewers] or list(stage.reviewer_ids)
        for user_id in ids:
            person = db.get(User, user_id)
            if not person or not person.active or person.organization_id != project.organization_id:
                raise HTTPException(422, "Reviewer must be an active user in this organization")
            member = db.query(Membership).filter_by(
                project_id=project.id, user_id=person.id, active=True
            ).first()
            if not member:
                raise HTTPException(422, f"{person.name} is not an active member of this project")


def _guard(db, user, instance, stage, action):
    if instance.record_type != DRAWING_REVIEW or action == "comment":
        return
    revision = db.get(DrawingRevision, instance.record_id)
    drawing = db.get(Drawing, revision.drawing_id) if revision else None
    if not revision or not drawing:
        raise HTTPException(404, "Drawing revision not found")
    if action in {"approve", "approve_with_comments"} and revision.version != instance.record_version:
        raise HTTPException(
            409,
            "This submitted file changed. The engineer must add a new revision and submit that revision.",
        )
    if user.id == revision.created_by and action in {"approve", "approve_with_comments"}:
        raise HTTPException(403, "The preparer cannot approve this revision")


def _stage_purpose(db, instance, stage):
    if stage is not None:
        return stage.purpose or ""
    from app.domains.coordination.service import current_stage

    current = current_stage(db, instance)
    return (current.purpose or "") if current else ""


def _policy_has_client(db, instance):
    from app.domains.coordination.models import PolicyStage

    return (
        db.query(PolicyStage).filter_by(policy_id=instance.policy_id, purpose="client").first()
        is not None
    )


def on_review_event(db, user, instance, project, action, stage=None):
    revision = db.get(DrawingRevision, instance.record_id)
    drawing = db.get(Drawing, revision.drawing_id)
    if action == "advance" and stage is not None:
        revision.review_state = "Client Review" if stage.purpose == "client" else "Internal Review"
        return
    if action == "return":
        revision.review_state = "Returned"
        if _stage_purpose(db, instance, stage) == "client":
            revision.client_outcome = "Revise and Resubmit"
        _event(db, user, drawing, revision, "returned")
        return
    if action == "reject":
        revision.review_state = "Rejected"
        if _stage_purpose(db, instance, stage) == "client":
            revision.client_outcome = "Rejected"
        _event(db, user, drawing, revision, "rejected")
        return
    if instance.state != "Approved":
        return
    decisions = (
        db.query(ReviewDecision)
        .filter_by(instance_id=instance.id, attempt=instance.attempt)
        .all()
    )
    with_comments = any(row.action == "approve_with_comments" for row in decisions)
    revision.review_state = "Accepted with Comments" if with_comments else "Accepted"
    if _policy_has_client(db, instance):
        revision.client_outcome = (
            "Approved with Comments" if with_comments else "Approved"
        )
    _event(
        db,
        user,
        drawing,
        revision,
        "accepted_with_comments" if with_comments else "accepted",
    )
    notify(
        db,
        db.get(User, instance.submitted_by),
        "Drawing accepted with comments" if with_comments else "Drawing accepted",
        f"{drawing.number} rev {revision.code} is not released for construction.",
        f"drawing-review:{instance.id}:accepted",
        project_id=project.id,
        link=f"/projects/{project.id}",
    )


def _bump(db, model, row_id, version, values):
    result = db.execute(
        update(model).where(model.id == row_id, model.version == version).values(
            **values, version=version + 1
        )
    )
    if result.rowcount != 1:
        raise HTTPException(409, "Record changed. Reload it and try again")


def create_drawing(db, user, project, data):
    _require(db, user, project, "drawings.edit", data.discipline_id)
    _check_refs(db, user, project, data)
    _apply_category(db, data)
    number = data.number.strip()
    key = number.casefold()
    if db.query(Drawing).filter_by(project_id=project.id, number_key=key).first():
        raise HTTPException(409, "Drawing number already exists in this project")
    row = Drawing(
        organization_id=user.organization_id,
        project_id=project.id,
        number=number,
        number_key=key,
        title=data.title.strip(),
        drawing_type=data.drawing_type.strip() or "General",
        category_id=data.category_id,
        discipline_id=data.discipline_id,
        company_id=data.company_id,
        responsible_user_id=data.responsible_user_id,
        reviewer_id=data.reviewer_id,
        area_id=data.area_id,
        planned_date=data.planned_date,
        planned_submission=data.planned_submission,
        planned_approval=data.planned_approval,
        priority=data.priority,
        work_reference=data.work_reference.strip(),
        remarks=data.remarks.strip(),
        provenance="manual",
        created_by=user.id,
    )
    _same_company(user, row)
    db.add(row)
    db.flush()
    return row


def _check_refs(db, user, project, data):
    from app.domains.identity_projects.router import company
    from app.domains.identity_projects.models import ProjectDiscipline

    enabled = {
        row.discipline_id
        for row in db.query(ProjectDiscipline).filter_by(project_id=project.id, active=True)
    }
    if data.discipline_id not in enabled:
        raise HTTPException(422, "Enable this discipline on the project first")
    if not db.get(Discipline, data.discipline_id):
        raise HTTPException(422, "Unknown discipline")
    company(db, user, data.company_id)
    if data.responsible_user_id:
        person = db.get(User, data.responsible_user_id)
        if not person or person.organization_id != user.organization_id or not person.active:
            raise HTTPException(422, "Responsible engineer must be an active user")
    if data.area_id:
        area = db.get(Area, data.area_id)
        if not area or area.project_id != project.id:
            raise HTTPException(422, "Location must belong to this project")
    if data.priority not in PRIORITIES:
        raise HTTPException(422, "Unknown priority")
    if getattr(data, "reviewer_id", None):
        person = db.get(User, data.reviewer_id)
        if not person or person.organization_id != user.organization_id or not person.active:
            raise HTTPException(422, "Reviewer must be an active user")


def _apply_category(db, data):
    if not getattr(data, "category_id", None):
        return
    category = db.get(DrawingCategory, data.category_id)
    if not category or not category.active or category.discipline_id != data.discipline_id:
        raise HTTPException(422, "Drawing category is not available for this discipline")
    data.drawing_type = category.name


def edit_drawing(db, user, project, drawing, data):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    _same_company(user, drawing)
    _check_refs(db, user, project, data)
    _apply_category(db, data)
    key = data.number.strip().casefold()
    if key != drawing.number_key and _latest(db, drawing.id):
        raise HTTPException(409, "Drawing number is fixed once a revision exists")
    taken = db.query(Drawing).filter_by(project_id=project.id, number_key=key).first()
    if taken and taken.id != drawing.id:
        raise HTTPException(409, "Drawing number already exists in this project")
    _bump(
        db,
        Drawing,
        drawing.id,
        data.version,
        {
            "number": data.number.strip(),
            "number_key": data.number.strip().casefold(),
            "title": data.title,
            "drawing_type": data.drawing_type,
            "discipline_id": data.discipline_id,
            "company_id": data.company_id,
            "responsible_user_id": data.responsible_user_id,
            "reviewer_id": data.reviewer_id,
            "area_id": data.area_id,
            "planned_date": data.planned_date,
            "planned_submission": data.planned_submission,
            "planned_approval": data.planned_approval,
            "priority": data.priority,
            "work_reference": data.work_reference,
            "remarks": data.remarks.strip(),
            "category_id": data.category_id,
            "drawing_type": data.drawing_type,
        },
    )


def start_preparation(db, user, project, drawing, version):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    _same_company(user, drawing)
    if drawing.preparation != "Planned":
        raise HTTPException(409, "Preparation has already started")
    if not drawing.responsible_user_id:
        raise HTTPException(422, "Assign a responsible engineer before preparation starts")
    if not drawing.planned_date:
        raise HTTPException(422, "A planned preparation date is required before preparation starts")
    _bump(db, Drawing, drawing.id, version, {"preparation": "In Preparation"})


def add_revision(db, user, project, drawing, code, upload, description="", purpose="Review"):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    _same_company(user, drawing)
    latest = _latest(db, drawing.id)
    if latest and latest.review_state in {"Not Submitted", "Internal Review", "Client Review"}:
        raise HTTPException(409, "Finish or correct the open revision before adding another")
    if latest and latest.review_state in {"Accepted", "Accepted with Comments"} and latest.release_state == "Not Released":
        raise HTTPException(409, "Release or withdraw the accepted revision before adding another")
    if not code.strip():
        raise HTTPException(422, "Revision code is required")
    stored = store_file(
        db,
        user,
        upload.filename or "drawing",
        upload.content_type or "",
        upload.data,
    )
    sequence = (latest.sequence + 1) if latest else 1
    revision = DrawingRevision(
        drawing_id=drawing.id,
        code=code.strip(),
        code_key=code.strip().casefold(),
        sequence=sequence,
        file_id=stored.id,
        description=(description or "").strip(),
        issue_purpose=purpose if purpose in {"Review", "Approval", "Construction", "As-Built"} else "Review",
        created_by=user.id,
    )
    db.add(revision)
    drawing.preparation = "Prepared"
    db.flush()
    _event(db, user, drawing, revision, "prepared")
    return revision


def submit_revision(db, user, project, drawing, revision):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    _same_company(user, drawing)
    if revision.drawing_id != drawing.id:
        raise HTTPException(404, "Revision not found")
    if revision.review_state != "Not Submitted":
        raise HTTPException(409, "This revision was already submitted")
    stored = db.get(StoredFile, revision.file_id)
    if not stored or stored.scan_status != "available":
        raise HTTPException(409, "Quarantined or missing files cannot be submitted")
    policy = _drawing_policy(db, project, drawing)
    if not policy:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    instance = submit_record(
        db, user, project, DRAWING_REVIEW, revision.id, revision.version, policy
    )
    revision.review_state = "Internal Review"
    _event(db, user, drawing, revision, "submitted")
    return instance


def _drawing_policy(db, project, drawing):
    from app.domains.coordination.models import ApprovalPolicy

    rows = (
        db.query(ApprovalPolicy)
        .filter_by(project_id=project.id, record_type=DRAWING_REVIEW, active=True)
        .all()
    )

    def fits(policy):
        if policy.company_id and policy.company_id != drawing.company_id:
            return False
        if policy.discipline_id and policy.discipline_id != drawing.discipline_id:
            return False
        return True

    matches = [row for row in rows if fits(row)]
    if not matches:
        return None
    return max(
        matches,
        key=lambda row: (
            int(row.company_id is not None) + int(row.discipline_id is not None),
            row.policy_version,
        ),
    )


def release_revision(db, user, project, drawing, revision, version, reason=""):
    _require(db, user, project, "drawings.release", drawing.discipline_id)
    if revision.drawing_id != drawing.id:
        raise HTTPException(404, "Revision not found")
    if user.id == revision.created_by:
        raise HTTPException(403, "The preparer cannot release this revision")
    if revision.review_state not in {"Accepted", "Accepted with Comments"}:
        raise HTTPException(409, "Construction release requires an accepted review")
    setting = db.get(EngineeringSetting, project.id)
    allowed = {"Approved"}
    if setting and setting.allow_release_with_comments:
        allowed.add("Approved with Comments")
    if revision.client_outcome not in allowed:
        raise HTTPException(
            409,
            "Construction release requires an approved client outcome. Approved with comments needs the project setting.",
        )
    if revision.release_state != "Not Released":
        raise HTTPException(409, "This revision is not waiting for release")
    stored = db.get(StoredFile, revision.file_id)
    if not stored or stored.scan_status != "available":
        raise HTTPException(409, "Quarantined or missing files cannot be released")
    blocking = (
        db.query(DrawingComment)
        .filter(
            DrawingComment.revision_id == revision.id,
            DrawingComment.severity == "Blocking",
            DrawingComment.status != "Closed",
        )
        .first()
    )
    if blocking:
        raise HTTPException(409, "Close blocking comments before release")
    _bump(
        db,
        DrawingRevision,
        revision.id,
        version,
        {
            "release_state": "Released",
            "released_by": user.id,
            "released_at": datetime.now(timezone.utc),
            "release_reference": (reason or "").strip(),
        },
    )
    db.refresh(revision)
    for previous in db.query(DrawingRevision).filter_by(drawing_id=drawing.id, release_state="Released"):
        if previous.id == revision.id:
            continue
        previous.release_state = "Superseded"
        _event(db, user, drawing, previous, "superseded", f"Superseded by {revision.code}")
        _publish(db, project, drawing, previous, "superseded")
        detail = f"Drawing {drawing.number} rev {previous.code} was superseded by {revision.code}"
        _procurement_impact(db, user, previous, detail)
        _construction_impact(db, user, previous, detail)
    _event(db, user, drawing, revision, "released")
    _publish(db, project, drawing, revision, "released")


def withdraw_revision(db, user, project, drawing, revision, version, reason):
    _require(db, user, project, "drawings.release", drawing.discipline_id)
    if revision.drawing_id != drawing.id:
        raise HTTPException(404, "Revision not found")
    if revision.release_state not in {"Released", "Not Released"}:
        raise HTTPException(409, "This revision cannot be withdrawn")
    if not reason.strip():
        raise HTTPException(422, "A reason is required")
    _bump(db, DrawingRevision, revision.id, version, {"release_state": "Withdrawn"})
    db.refresh(revision)
    _event(db, user, drawing, revision, "withdrawn", reason)
    _publish(db, project, drawing, revision, "withdrawn")
    detail = f"Drawing {drawing.number} rev {revision.code} was withdrawn: {reason}"
    _procurement_impact(db, user, revision, detail)
    _construction_impact(db, user, revision, detail)


def _procurement_impact(db, user, revision, detail):
    from app.domains.supply_chain.service import flag_revision

    flag_revision(db, user, revision.id, detail)


def _construction_impact(db, user, revision, detail):
    from app.domains.construction.impact import flag_revision as flag_packages

    flag_packages(db, user, revision, detail)


def _publish(db, project, drawing, revision, kind):
    recipients = {project.lead_pm_id} if project.lead_pm_id else set()
    if drawing.responsible_user_id:
        recipients.add(drawing.responsible_user_id)
    for person_id in recipients:
        person = db.get(User, person_id)
        if not person or not person.active:
            continue
        notify(
            db,
            person,
            f"Drawing {kind}",
            f"{drawing.number} rev {revision.code}",
            f"revision-event:{revision.id}:{kind}:{person.id}",
            project_id=project.id,
            link=f"/projects/{project.id}",
        )


def add_comment(db, user, project, drawing, revision, data):
    if revision.drawing_id != drawing.id:
        raise HTTPException(404, "Revision not found")
    if not (
        _can(db, user, project, "drawings.review", drawing.discipline_id)
        or _can(db, user, project, "drawings.client_review", drawing.discipline_id)
    ):
        raise HTTPException(403, "Only a reviewer can raise a comment")
    if data.severity not in SEVERITIES:
        raise HTTPException(422, "Unknown comment severity")
    owner = db.get(User, data.owner_id)
    if not owner or owner.organization_id != user.organization_id or not owner.active:
        raise HTTPException(422, "Comment owner must be an active user")
    stage = data.stage if data.stage in {"Internal", "Client"} else "Internal"
    number = db.query(DrawingComment).filter_by(revision_id=revision.id).count() + 1
    row = DrawingComment(
        revision_id=revision.id,
        page_ref=data.page_ref,
        body=data.body,
        severity=data.severity,
        stage=stage,
        comment_number=number,
        owner_id=data.owner_id,
        due_date=data.due_date,
        created_by=user.id,
    )
    db.add(row)
    db.flush()
    notify(
        db,
        owner,
        "Drawing comment assigned",
        data.body[:180],
        f"comment:{row.id}:assigned",
        project_id=project.id,
        link=f"/projects/{project.id}",
    )
    return row


def _comment_permission(comment):
    return "drawings.client_review" if comment.stage == "Client" else "drawings.review"


def can_respond_comment(db, user, project, drawing, comment):
    if comment.status != "Open":
        return False
    if user.organization_admin or user.id in {comment.owner_id, drawing.responsible_user_id}:
        return True
    return False


def can_close_comment(db, user, project, drawing, comment):
    if comment.status not in {"Open", "Responded"}:
        return False
    if user.organization_admin or user.id in {
        comment.owner_id,
        comment.created_by,
        drawing.responsible_user_id,
    }:
        return True
    return _can(db, user, project, _comment_permission(comment), drawing.discipline_id)


def respond_comment(db, user, project, drawing, comment, response, version):
    if not can_respond_comment(db, user, project, drawing, comment):
        if comment.status != "Open":
            raise HTTPException(409, "This comment is no longer open for a response")
        raise HTTPException(403, "Only the assigned engineer can respond")
    if not response.strip():
        raise HTTPException(422, "A response is required")
    _bump(db, DrawingComment, comment.id, version, {"response": response.strip(), "status": "Responded"})


def close_comment(db, user, project, drawing, comment, evidence, version):
    if not can_close_comment(db, user, project, drawing, comment):
        if comment.status not in {"Open", "Responded"}:
            raise HTTPException(409, "This comment is already closed")
        raise HTTPException(403, "You cannot close this comment")
    if not evidence.strip():
        raise HTTPException(422, "Closure evidence is required")
    _bump(
        db,
        DrawingComment,
        comment.id,
        version,
        {
            "closure_evidence": evidence,
            "status": "Closed",
            "closed_by": user.id,
            "closed_at": datetime.now(timezone.utc),
        },
    )


def drawing_recipients(db, user, project):
    _require(db, user, project, "drawings.view")
    people = []
    seen = set()
    memberships = db.query(Membership).filter_by(project_id=project.id, active=True)
    for membership in memberships:
        person = db.get(User, membership.user_id)
        if not person or not person.active or person.id in seen:
            continue
        seen.add(person.id)
        people.append({"id": person.id, "name": person.name})
    return people


def visible_transmittals(db, user, project):
    _require(db, user, project, "drawings.view")
    rows = db.query(Transmittal).filter_by(project_id=project.id).order_by(Transmittal.id).all()
    if user.organization_admin or _can(db, user, project, "transmittals.issue"):
        return rows
    allowed = {
        item.transmittal_id
        for item in db.query(TransmittalRecipient).filter_by(user_id=user.id)
    }
    return [row for row in rows if row.id in allowed]


def create_transmittal(db, user, project, data):
    _require(db, user, project, "transmittals.issue")
    if not data.revision_ids or not data.recipient_ids:
        raise HTTPException(422, "A transmittal needs revisions and recipients")
    row = Transmittal(
        project_id=project.id,
        number=data.number,
        number_key=data.number.casefold(),
        purpose=data.purpose,
        issue_date=data.issue_date,
        sender_id=user.id,
        sender_company_id=data.sender_company_id,
        recipient_company_id=data.recipient_company_id,
        expected_response=data.expected_response,
        method=(data.method or "").strip(),
        external_reference=(data.external_reference or "").strip(),
    )
    db.add(row)
    db.flush()
    for revision_id in dict.fromkeys(data.revision_ids):
        revision = db.get(DrawingRevision, revision_id)
        drawing = db.get(Drawing, revision.drawing_id) if revision else None
        if not drawing or drawing.project_id != project.id:
            raise HTTPException(422, "Revision is not on this project")
        db.add(TransmittalItem(transmittal_id=row.id, revision_id=revision_id))
    for user_id in dict.fromkeys(data.recipient_ids):
        person = db.get(User, user_id)
        if not person or person.organization_id != user.organization_id or not person.active:
            raise HTTPException(422, "Recipient must be an active user in this organization")
        db.add(TransmittalRecipient(transmittal_id=row.id, user_id=user_id))
    return row


def dispatch_transmittal(db, user, project, transmittal, version):
    _require(db, user, project, "transmittals.issue")
    if transmittal.project_id != project.id or transmittal.status != "Draft":
        raise HTTPException(409, "Only a draft transmittal can be dispatched")
    _bump(
        db,
        Transmittal,
        transmittal.id,
        version,
        {"status": "Issued", "dispatched_at": datetime.now(timezone.utc)},
    )
    for recipient in db.query(TransmittalRecipient).filter_by(transmittal_id=transmittal.id):
        person = db.get(User, recipient.user_id)
        notify(
            db,
            person,
            "Transmittal dispatched",
            transmittal.purpose,
            f"transmittal:{transmittal.id}:dispatched:{person.id}",
            project_id=project.id,
            link=f"/projects/{project.id}",
        )


def acknowledge_transmittal(db, user, project, transmittal):
    if transmittal.project_id != project.id or transmittal.status not in {"Issued", "Acknowledged"}:
        raise HTTPException(409, "This transmittal has not been dispatched")
    recipient = (
        db.query(TransmittalRecipient)
        .filter_by(transmittal_id=transmittal.id, user_id=user.id)
        .first()
    )
    if not recipient and not user.organization_admin:
        raise HTTPException(403, "You are not a recipient of this transmittal")
    if recipient and not recipient.acknowledged_at:
        recipient.acknowledged_at = datetime.now(timezone.utc)
    pending = (
        db.query(TransmittalRecipient)
        .filter_by(transmittal_id=transmittal.id)
        .filter(TransmittalRecipient.acknowledged_at.is_(None))
        .count()
    )
    if pending == 0 and transmittal.status != "Acknowledged":
        transmittal.status = "Acknowledged"
        transmittal.acknowledged_at = datetime.now(timezone.utc)


def import_rows(db, user, project, filename, payload, mode="create", location_mappings=None):
    _require(db, user, project, "drawings.edit")
    if mode not in {"create", "update"}:
        raise HTTPException(422, "Import mode must be create or update")
    if not payload:
        raise HTTPException(422, "The file is empty")
    rows, columns = _parse_table(filename, payload)
    missing = [name for name in ("number", "title", "discipline") if name not in columns]
    if missing:
        raise HTTPException(
            422,
            f"Missing column(s): {', '.join(missing)}. Found: {', '.join(columns) or 'no header row'}. "
            "Use the template: " + ", ".join(IMPORT_COLUMNS),
        )
    location_mappings = location_mappings or {}
    if not isinstance(location_mappings, dict):
        raise HTTPException(422, "Location mappings must be a JSON object")
    active_areas = list(db.query(Area).filter_by(project_id=project.id, active=True))
    area_by_id = {area.id: area for area in active_areas}
    validated_mappings = {}
    for source, target in location_mappings.items():
        if not isinstance(source, str) or not source.strip() or isinstance(target, bool):
            raise HTTPException(422, "Invalid location mapping")
        try:
            target_id = int(target)
        except (ValueError, TypeError):
            raise HTTPException(422, "Location mapping targets must be area IDs")
        if target_id not in area_by_id:
            raise HTTPException(422, "Mapped location is not an active area of this project")
        validated_mappings[source.strip().casefold()] = target_id
    created, skipped, updated, conflicts, lines = [], [], [], [], []
    seen_numbers = set()
    for index, raw in enumerate(rows, start=2):
        number = str(raw.get("number") or "").strip()
        title = str(raw.get("title") or "").strip()
        try:
            if number.casefold() in seen_numbers:
                raise HTTPException(409, "Duplicate drawing number within the uploaded file")
            seen_numbers.add(number.casefold())
            if raw.get("revision") or raw.get("approval_status") or raw.get("document_type") or raw.get("description"):
                raise HTTPException(422, "Revision, approval status, document type and description are not supported by the drawing register import. Use the dedicated controlled revision workflow; no approval can be granted by import.")
            outcome, row = _import_one(db, user, project, raw, mode, validated_mappings)
        except HTTPException as exc:
            conflicts.append({"row": index, "number": number, "reason": str(exc.detail)})
            lines.append({"row": index, "number": number, "title": title, "outcome": "conflict", "reason": str(exc.detail), "values": _display_values(raw)})
            continue
        if outcome == "created":
            created.append(row.number)
        elif outcome == "updated":
            updated.append(row.number)
        else:
            skipped.append(row.number)
        reason = "Already in the register with the same details" if outcome == "skipped" else ""
        lines.append({"row": index, "number": row.number, "title": row.title, "outcome": outcome, "reason": reason, "values": _display_values(raw)})
    return {
        "created": created,
        "skipped": skipped,
        "updated": updated,
        "conflicts": conflicts,
        "rows": lines,
        "columns": columns,
        "unsupported_columns": [c for c in columns if c not in IMPORT_COLUMNS],
        "total": len(lines),
        "valid": len(created) + len(updated),
        "invalid": len(conflicts),
        "duplicates": sum("duplicate" in str(c["reason"]).lower() or "already exists" in str(c["reason"]).lower() for c in conflicts) + len(skipped),
        "areas": [{"id": a.id, "code": a.code, "name": a.name} for a in active_areas],
    }


def _display_values(raw):
    return {key: (value.isoformat() if isinstance(value, (datetime, date)) else str(value) if value is not None else "") for key, value in raw.items()}


IMPORT_COLUMNS = (
    "number", "title", "type", "discipline", "company", "responsible", "location",
    "planned_date", "planned_submission", "planned_approval", "priority", "work_reference", "remarks",
)
_HEADER_ALIASES = {
    "number": "number", "drawing number": "number", "drawing no": "number", "dwg no": "number",
    "dwg number": "number", "drawing": "number", "no": "number",
    "title": "title", "drawing title": "title", "name": "title",
    "type": "type", "drawing type": "type",
    "discipline": "discipline",
    "company": "company", "originator": "company", "originator company": "company",
    "responsible": "responsible", "responsible engineer": "responsible", "engineer": "responsible",
    "owner": "responsible",
    "location": "location", "area": "location", "unit": "location", "block": "location",
    "unit area block": "location",
    "planned date": "planned_date", "planned preparation": "planned_date",
    "planned preparation date": "planned_date", "required date": "planned_date",
    "planned submission": "planned_submission", "planned client submission": "planned_submission",
    "planned approval": "planned_approval",
    "priority": "priority", "work reference": "work_reference", "remarks": "remarks",
    "project area": "location", "project area location": "location", "location area": "location",
    "revision": "revision", "drawing revision": "revision", "rev": "revision",
    "approval status": "approval_status", "document type": "document_type", "description": "description",
}


def _header(value):
    key = re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()
    return _HEADER_ALIASES.get(key, key.replace(" ", "_"))


def _parse_table(filename, payload):
    records = read_records(filename, payload)
    if not records:
        return [], []
    headers = [_header(cell) for cell in records[0]]
    parsed = []
    for record in records[1:]:
        if not any(str(cell or "").strip() for cell in record):
            continue
        parsed.append({headers[i]: record[i] for i in range(len(headers)) if i < len(record) and headers[i]})
    return parsed, [item for item in headers if item]


def read_records(filename, payload):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(422, "The Excel file could not be read. Save it as .xlsx or .csv and try again")
        records = list(book.active.iter_rows(values_only=True))
    elif name.endswith(".csv") or not name:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = payload.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;\t")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
        records = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise HTTPException(422, "Upload a .csv or .xlsx file")
    return records


def _import_date(value, label):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    for candidate in (text, text[:10]):
        for pattern in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(candidate, pattern).date()
            except ValueError:
                continue
    raise HTTPException(422, f"{label} '{text}' is not a date. Use YYYY-MM-DD or DD/MM/YYYY")


def _import_one(db, user, project, raw, mode="create", location_mappings=None):
    def cell(key):
        value = raw.get(key, "")
        return "" if value is None else str(value).strip()

    number = cell("number")
    title = cell("title")
    if not number or not title:
        raise HTTPException(422, "Number and title are required")
    discipline = next(
        (row for row in db.query(Discipline).all() if row.name.casefold() == cell("discipline").casefold()),
        None,
    )
    if not discipline:
        raise HTTPException(422, f"Discipline '{cell('discipline')}' was not found")
    if cell("company"):
        company = (
            db.query(Company)
            .filter_by(organization_id=user.organization_id, name_key=cell("company").casefold())
            .first()
        )
        if not company:
            raise HTTPException(422, f"Company '{cell('company')}' was not found")
    else:
        company = db.get(Company, user.company_id) if user.company_id else None
        if not company:
            raise HTTPException(422, "Company is required because your account has no company")
    responsible = None
    if cell("responsible"):
        wanted = cell("responsible").casefold()
        responsible = next(
            (
                person
                for person in db.query(User).filter_by(organization_id=user.organization_id, active=True)
                if person.email.casefold() == wanted or person.name.casefold() == wanted
            ),
            None,
        )
        if not responsible:
            raise HTTPException(422, f"Responsible engineer '{cell('responsible')}' was not found")
    area_id = None
    if cell("location"):
        wanted = cell("location").casefold()
        area = next(
            (
                row
                for row in db.query(Area).filter_by(project_id=project.id, active=True)
                if row.code_key == wanted or row.name.casefold() == wanted
            ),
            None,
        )
        if not area and location_mappings and wanted in location_mappings:
            area = db.query(Area).filter_by(id=location_mappings[wanted], project_id=project.id, active=True).first()
        if not area:
            raise HTTPException(422, f"Location '{cell('location')}' was not found on this project. Map it to an active project location and revalidate.")
        area_id = area.id
    planned = _import_date(raw.get("planned_date"), "Planned date")
    planned_submission = _import_date(raw.get("planned_submission"), "Planned submission")
    planned_approval = _import_date(raw.get("planned_approval"), "Planned approval")
    priority = (cell("priority") or "Normal").title()
    if priority not in {"Low", "Normal", "High", "Critical"}:
        raise HTTPException(422, "Priority must be Low, Normal, High or Critical")
    existing = (
        db.query(Drawing)
        .filter_by(project_id=project.id, number_key=number.casefold())
        .first()
    )
    if existing:
        if mode != "update":
            same = (
                existing.title == title
                and existing.discipline_id == discipline.id
                and existing.company_id == company.id
            )
            if same:
                return "skipped", existing
            raise HTTPException(409, "Drawing number already exists with different content")
        if _latest(db, existing.id):
            existing.title = title
            existing.priority = priority
            existing.planned_date = planned
            existing.remarks = cell("remarks")
            existing.version += 1
            return "updated", existing
        existing.title = title
        existing.discipline_id = discipline.id
        existing.company_id = company.id
        existing.responsible_user_id = responsible.id if responsible else existing.responsible_user_id
        existing.area_id = area_id or existing.area_id
        existing.planned_date = planned
        existing.priority = priority
        existing.remarks = cell("remarks")
        existing.version += 1
        return "updated", existing
    from app.domains.identity_projects.schemas import Strict

    class _In(Strict):
        number: str
        title: str
        drawing_type: str = "General"
        category_id: int | None = None
        reviewer_id: int | None = None
        discipline_id: int
        company_id: int
        responsible_user_id: int | None = None
        area_id: int | None = None
        planned_date: date | None = None
        planned_submission: date | None = None
        planned_approval: date | None = None
        priority: str = "Normal"
        work_reference: str = ""
        remarks: str = ""

    data = _In(
        number=number,
        title=title,
        drawing_type=cell("type") or "General",
        discipline_id=discipline.id,
        company_id=company.id,
        responsible_user_id=responsible.id if responsible else None,
        area_id=area_id,
        planned_date=planned,
        planned_submission=planned_submission,
        planned_approval=planned_approval,
        priority=priority,
        work_reference=cell("work_reference"),
        remarks=cell("remarks"),
    )
    row = create_drawing(db, user, project, data)
    row.provenance = "import"
    return "created", row


def save_wrench(db, project, data):
    if data.enabled:
        raise HTTPException(
            409,
            "Wrench synchronization stays disabled until a verified API specification and credentials are supplied",
        )
    row = db.get(WrenchSetting, project.id)
    if not row:
        row = WrenchSetting(project_id=project.id)
        db.add(row)
    row.enabled = False
    row.specification = data.specification
    row.credential_label = data.credential_label
    return row


def link_wrench(db, user, project, drawing, data):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    row = db.query(WrenchLink).filter_by(drawing_id=drawing.id).first()
    if not row:
        row = WrenchLink(
            drawing_id=drawing.id,
            external_document_id=data.external_document_id,
            external_revision_id=data.external_revision_id,
            provenance="manual",
        )
        db.add(row)
        db.flush()
        return row
    if (
        row.external_document_id != data.external_document_id
        or row.external_revision_id != data.external_revision_id
    ):
        row.conflict_status = "conflict"
        row.detail = "External ids differ from the stored mapping. The drawing was not changed."
        return row
    return row


def scan_overdue(db, today=None):
    today = today or datetime.now(timezone.utc).date()
    drawings = db.query(Drawing).all()
    for drawing in drawings:
        project = db.get(Project, drawing.project_id)
        latest = _latest(db, drawing.id)
        if (
            drawing.planned_date
            and drawing.planned_date < today
            and bucket_for(drawing, latest) in {"to_prepare", "in_preparation"}
        ):
            _escalate(db, project, drawing, "preparation", f"drawing:{drawing.id}:prep")
        if latest and latest.review_state in {"Internal Review", "Client Review"}:
            _escalate(db, project, drawing, "review", f"revision:{latest.id}:review")
    for comment in db.query(DrawingComment).filter_by(status="Open"):
        if comment.due_date and comment.due_date < today:
            revision = db.get(DrawingRevision, comment.revision_id)
            drawing = db.get(Drawing, revision.drawing_id)
            project = db.get(Project, drawing.project_id)
            _escalate(db, project, drawing, "comment", f"comment:{comment.id}:response", comment.owner_id)


def _escalate(db, project, drawing, kind, source, owner_id=None):
    owner_id = owner_id or drawing.responsible_user_id or drawing.created_by
    recipients = [owner_id]
    heads = []
    from app.domains.identity_projects.models import Membership, Role

    for member in db.query(Membership).filter_by(project_id=project.id, active=True):
        role = db.get(Role, member.role_id)
        if "drawings.review" not in (role.permissions or []):
            continue
        if member.scope == "SELECTED" and drawing.discipline_id not in (member.discipline_ids or []):
            continue
        heads.append(member.user_id)
    if project.lead_pm_id:
        heads.append(project.lead_pm_id)
    sent = set()
    for person_id in recipients:
        if not person_id or person_id in sent:
            continue
        sent.add(person_id)
        person = db.get(User, person_id)
        if person and person.active:
            notify(
                db,
                person,
                f"Overdue drawing {kind}",
                drawing.number,
                f"{source}:owner:{person.id}",
                project_id=project.id,
                link=f"/projects/{project.id}",
            )
    for person_id in heads:
        if not person_id or person_id in sent:
            continue
        sent.add(person_id)
        person = db.get(User, person_id)
        if not person or not person.active:
            continue
        notify(
            db,
            person,
            f"Drawing {kind} escalated",
            f"{drawing.number}. Escalation adds visibility and does not grant approval.",
            f"{source}:escalate:{person.id}",
            project_id=project.id,
            link=f"/projects/{project.id}",
        )


def revision_file(db, revision):
    stored = db.get(StoredFile, revision.file_id)
    if not stored or stored.scan_status == "quarantined":
        raise HTTPException(404, "File not found")
    return stored, read_file(stored)


def summary(db, drawings):
    counts = {name: 0 for name in BUCKETS}
    active = [(drawing, latest) for drawing, latest in drawings if drawing.record_status != "Cancelled"]
    for drawing, latest in active:
        counts[bucket_for(drawing, latest)] += 1
    measures = {name: 0 for name in MEASURES}
    today = date.today()
    for drawing, latest in active:
        for name in MEASURES:
            if name == "open_comments":
                continue
            if matches_measure(db, drawing, latest, name, today):
                measures[name] += 1
        measures["open_comments"] += _open_comment_count(db, drawing)
    approved = measures["approved"]
    total = measures["active"]
    return {
        "basis": "Each active drawing is counted once. Dashboard cards overlap and are not a single total. Client reviewed is Approved, Approved with comments, Revise and resubmit, and Rejected. Open comments count comment records.",
        "counts": counts,
        "measures": measures,
        "approval_percent": None if total == 0 else round(approved * 100 / total, 1),
        "total": len(active),
    }


MEASURES = (
    "active",
    "to_be_prepared",
    "in_preparation",
    "prepared",
    "pending_internal",
    "pending_client",
    "client_reviewed",
    "approved",
    "approved_with_comments",
    "revise_and_resubmit",
    "rejected",
    "released",
    "open_comments",
    "overdue_preparation",
    "overdue_client",
)
CLIENT_DONE = {"Approved", "Approved with Comments", "Revise and Resubmit", "Rejected"}


def matches_measure(db, drawing, latest, name, today=None):
    today = today or date.today()
    outcome = latest.client_outcome if latest else ""
    if name == "active":
        return drawing.record_status != "Cancelled"
    if name == "to_be_prepared":
        return drawing.preparation == "Planned"
    if name == "in_preparation":
        return drawing.preparation == "In Preparation"
    if name == "prepared":
        return drawing.preparation == "Prepared"
    if name == "pending_internal":
        return bool(latest and latest.review_state == "Internal Review")
    if name == "pending_client":
        return bool(latest and latest.review_state == "Client Review")
    if name == "approved":
        return outcome == "Approved"
    if name == "approved_with_comments":
        return outcome == "Approved with Comments"
    if name == "revise_and_resubmit":
        return outcome == "Revise and Resubmit"
    if name == "rejected":
        return outcome == "Rejected"
    if name == "client_reviewed":
        return outcome in CLIENT_DONE
    if name == "released":
        return _released_revision(db, drawing.id) is not None
    if name == "overdue_preparation":
        return bool(
            drawing.planned_date
            and drawing.planned_date < today
            and drawing.preparation != "Prepared"
        )
    if name == "overdue_client":
        return _overdue_client(db, latest, today)
    if name == "open_comments":
        return _open_comment_count(db, drawing) > 0
    return False


def _open_comment_count(db, drawing):
    revision_ids = [
        row.id for row in db.query(DrawingRevision.id).filter_by(drawing_id=drawing.id)
    ]
    if not revision_ids:
        return 0
    return (
        db.query(DrawingComment)
        .filter(DrawingComment.revision_id.in_(revision_ids), DrawingComment.status != "Closed")
        .count()
    )


def _overdue_client(db, latest, today):
    if not latest or latest.review_state != "Client Review":
        return False
    return (
        db.query(TransmittalItem)
        .join(Transmittal, Transmittal.id == TransmittalItem.transmittal_id)
        .filter(
            TransmittalItem.revision_id == latest.id,
            Transmittal.expected_response.isnot(None),
            Transmittal.expected_response < today,
            Transmittal.status.in_(["Issued", "Acknowledged"]),
        )
        .first()
        is not None
    )


def filtered_drawings(
    db,
    project,
    discipline_id=None,
    company_id=None,
    owner_id=None,
    planned_from=None,
    planned_to=None,
    category_id=None,
    area_id=None,
    priority="",
    source="",
    record_status="Active",
    date_kind="planned_preparation",
):
    query = db.query(Drawing).filter_by(project_id=project.id)
    if record_status and record_status != "All":
        query = query.filter_by(record_status=record_status)
    if discipline_id:
        query = query.filter_by(discipline_id=discipline_id)
    if company_id:
        query = query.filter_by(company_id=company_id)
    if owner_id:
        query = query.filter_by(responsible_user_id=owner_id)
    if category_id:
        query = query.filter_by(category_id=category_id)
    if area_id:
        query = query.filter_by(area_id=area_id)
    if priority:
        query = query.filter_by(priority=priority)
    if source:
        query = query.filter_by(provenance=source)
    column = {
        "planned_preparation": Drawing.planned_date,
        "planned_submission": Drawing.planned_submission,
        "planned_approval": Drawing.planned_approval,
    }.get(date_kind, Drawing.planned_date)
    if planned_from:
        query = query.filter(column >= planned_from)
    if planned_to:
        query = query.filter(column <= planned_to)
    rows = []
    for drawing in query.order_by(Drawing.number).all():
        rows.append((drawing, _latest(db, drawing.id)))
    return rows


def public_drawing(db, user, project, drawing, latest):
    discipline = db.get(Discipline, drawing.discipline_id)
    company = db.get(Company, drawing.company_id)
    responsible = db.get(User, drawing.responsible_user_id) if drawing.responsible_user_id else None
    return {
        "id": drawing.id,
        "number": drawing.number,
        "title": drawing.title,
        "drawing_type": drawing.drawing_type,
        "discipline_id": drawing.discipline_id,
        "discipline": discipline.name if discipline else "",
        "company_id": drawing.company_id,
        "company": company.name if company else "",
        "responsible_user_id": drawing.responsible_user_id,
        "responsible": responsible.name if responsible else "",
        "area_id": drawing.area_id,
        "planned_date": drawing.planned_date,
        "priority": drawing.priority,
        "preparation": drawing.preparation,
        "work_reference": drawing.work_reference,
        "provenance": drawing.provenance,
        "category_id": drawing.category_id,
        "reviewer_id": drawing.reviewer_id,
        "planned_submission": drawing.planned_submission,
        "planned_approval": drawing.planned_approval,
        "remarks": drawing.remarks,
        "record_status": drawing.record_status,
        "cancellation_reason": drawing.cancellation_reason,
        "version": drawing.version,
        "bucket": bucket_for(drawing, latest),
        "preparation_status": _preparation_label(drawing),
        "internal_status": _internal_label(latest),
        "client_status": _client_label(latest),
        "latest_revision": latest.code if latest else "",
        "review_state": latest.review_state if latest else "",
        "release_state": latest.release_state if latest else "",
    }


def _preparation_label(drawing):
    if drawing.preparation == "Planned":
        return "To Be Prepared"
    return drawing.preparation


def _internal_label(latest):
    if not latest or latest.review_state == "Not Submitted":
        return "Not Submitted"
    if latest.review_state == "Internal Review":
        return "Pending Review"
    if latest.review_state == "Rejected":
        return "Rejected"
    if latest.review_state == "Returned":
        return "Changes Required"
    if latest.review_state == "Accepted with Comments":
        return "Accepted with comments"
    if latest.review_state in {"Accepted", "Client Review"}:
        return "Accepted"
    return latest.review_state


def _client_label(latest):
    if not latest:
        return "Not Submitted"
    if latest.client_outcome:
        return latest.client_outcome
    if latest.review_state == "Client Review":
        return "Pending Client Review"
    return "Not Submitted"


def detail(db, user, project, drawing):
    latest = _latest(db, drawing.id)
    revisions = (
        db.query(DrawingRevision)
        .filter_by(drawing_id=drawing.id)
        .order_by(DrawingRevision.sequence)
        .all()
    )
    if not sees_unreleased(db, user, project):
        revisions = [
            row
            for row in revisions
            if row.release_state in {"Released", "Superseded", "Withdrawn"}
        ]
        latest = _released_revision(db, drawing.id)
    payload = public_drawing(db, user, project, drawing, latest)
    released = (
        db.query(DrawingRevision)
        .filter_by(drawing_id=drawing.id, release_state="Released")
        .order_by(DrawingRevision.sequence.desc())
        .first()
    )
    payload["latest_released"] = (
        {"id": released.id, "code": released.code} if released else None
    )
    payload["revisions"] = [public_revision(db, user, project, drawing, row) for row in revisions]
    payload["actions"] = drawing_actions(db, user, project, drawing, latest)
    payload["links"] = [
        {"id": link.id, "related_id": link.related_id}
        for link in db.query(DrawingLink).filter_by(drawing_id=drawing.id)
    ]
    setting = db.get(EngineeringSetting, project.id)
    payload["allow_release_with_comments"] = bool(
        setting and setting.allow_release_with_comments
    )
    from .dashboard import tracking

    payload["tracking"] = tracking(db, project, drawing, revisions, latest, released, setting)
    payload["assignments"] = assignments(db, user, project, drawing)
    return payload


def assignments(db, user, project, drawing):
    from app.domains.coordination.models import PolicyReviewer, PolicyStage
    from app.domains.identity_projects.models import Role

    discipline = db.get(Discipline, drawing.discipline_id)
    discipline_name = discipline.name if discipline else "this discipline"
    policy = _drawing_policy(db, project, drawing)
    route = None
    if policy:
        stages = []
        for stage in db.query(PolicyStage).filter_by(policy_id=policy.id).order_by(PolicyStage.position):
            names = [
                _user_name(db, row.user_id)
                for row in db.query(PolicyReviewer).filter_by(stage_id=stage.id)
            ]
            stages.append(
                {
                    "name": stage.name,
                    "purpose": stage.purpose or "internal",
                    "rule": "All must approve" if stage.rule == "all" else f"{stage.quorum_count} must approve",
                    "reviewers": [name for name in names if name],
                }
            )
        route = {
            "policy_id": policy.id,
            "applies_to": discipline_name if policy.discipline_id else "All disciplines",
            "stages": stages,
            "self_approval": bool(policy.allow_self_approval),
        }
    release, issuers = [], []
    for member in db.query(Membership).filter_by(project_id=project.id, active=True):
        role = db.get(Role, member.role_id)
        if not role:
            continue
        covers = member.scope == "ALL" or drawing.discipline_id in (member.discipline_ids or [])
        name = _user_name(db, member.user_id)
        if covers and "drawings.release" in (role.permissions or []) and name not in release:
            release.append(name)
        if "transmittals.issue" in (role.permissions or []) and name not in issuers:
            issuers.append(name)
    responsible = db.get(User, drawing.responsible_user_id) if drawing.responsible_user_id else None
    return {
        "discipline": discipline_name,
        "responsible": {"id": responsible.id, "name": responsible.name} if responsible else None,
        "can_assign": _can(db, user, project, "drawings.edit", drawing.discipline_id),
        "review_route": route,
        "review_route_missing": None
        if route
        else f"No drawing review policy covers {discipline_name}. An administrator sets one on the Review policy tab.",
        "release_authority": [name for name in release if name],
        "document_controllers": [name for name in issuers if name],
    }


def _public_comment(db, user, project, drawing, row):
    owner = db.get(User, row.owner_id)
    raiser = db.get(User, row.created_by)
    return {
        "id": row.id,
        "page_ref": row.page_ref,
        "body": row.body,
        "severity": row.severity,
        "stage": row.stage,
        "comment_number": row.comment_number,
        "owner_id": row.owner_id,
        "owner_name": owner.name if owner else "",
        "raised_by": raiser.name if raiser else "",
        "can_respond": can_respond_comment(db, user, project, drawing, row),
        "can_close": can_close_comment(db, user, project, drawing, row),
        "closed_by": row.closed_by,
        "closed_at": row.closed_at,
        "due_date": row.due_date,
        "response": row.response,
        "closure_evidence": row.closure_evidence,
        "status": row.status,
        "version": row.version,
    }


def _user_name(db, user_id):
    if not user_id:
        return ""
    person = db.get(User, user_id)
    return person.name if person else ""


def public_revision(db, user, project, drawing, revision):
    stored = db.get(StoredFile, revision.file_id)
    review = (
        db.query(ReviewInstance)
        .filter_by(record_type=DRAWING_REVIEW, record_id=revision.id)
        .order_by(ReviewInstance.id.desc())
        .first()
    )
    comments = db.query(DrawingComment).filter_by(revision_id=revision.id).order_by(DrawingComment.id).all()
    events = db.query(RevisionEvent).filter_by(revision_id=revision.id).order_by(RevisionEvent.id).all()
    return {
        "id": revision.id,
        "code": revision.code,
        "sequence": revision.sequence,
        "review_state": revision.review_state,
        "release_state": revision.release_state,
        "description": revision.description,
        "issue_purpose": revision.issue_purpose,
        "client_outcome": revision.client_outcome,
        "external_reviewer": revision.external_reviewer,
        "decision_reference": revision.decision_reference,
        "released_at": revision.released_at,
        "release_reference": revision.release_reference,
        "file_name": stored.original_name if stored else "",
        "sha256": stored.sha256 if stored else "",
        "scan_status": stored.scan_status if stored else "",
        "version": revision.version,
        "created_by": revision.created_by,
        "created_by_name": _user_name(db, revision.created_by),
        "prepared_at": revision.prepared_at,
        "released_by_name": _user_name(db, revision.released_by),
        "decision_recorded_by_name": _user_name(db, revision.decision_recorded_by),
        "review": review_payload(db, user, review) if review else None,
        "comments": [
            _public_comment(db, user, project, drawing, row) for row in comments
        ],
        "events": [
            {
                "id": row.id,
                "kind": row.kind,
                "detail": row.detail,
                "created_at": row.created_at,
                "actor": _user_name(db, row.actor_id),
            }
            for row in events
        ],
        "actions": revision_actions(db, user, project, drawing, revision),
    }


def drawing_actions(db, user, project, drawing, latest):
    actions = []
    if not _can(db, user, project, "drawings.edit", drawing.discipline_id):
        return actions
    if user.organization_admin or user.company_id == drawing.company_id:
        if drawing.preparation == "Planned":
            actions.append("start_preparation")
        if not latest or latest.review_state in {"Returned", "Rejected"} or latest.release_state in {
            "Released",
            "Superseded",
            "Withdrawn",
        }:
            actions.append("add_revision")
    return actions


def revision_actions(db, user, project, drawing, revision):
    actions = []
    if _can(db, user, project, "drawings.edit", drawing.discipline_id) and (
        user.organization_admin or user.company_id == drawing.company_id
    ):
        if revision.review_state == "Not Submitted":
            actions.append("submit")
    if (
        _can(db, user, project, "drawings.release", drawing.discipline_id)
        and user.id != revision.created_by
        and revision.review_state in {"Accepted", "Accepted with Comments"}
        and revision.release_state == "Not Released"
        and _release_allowed(db, project, revision)
    ):
        actions.append("release")
    if _can(db, user, project, "drawings.release", drawing.discipline_id) and revision.release_state in {
        "Released",
        "Not Released",
    }:
        actions.append("withdraw")
    if _can(db, user, project, "drawings.review", drawing.discipline_id) or _can(
        db, user, project, "drawings.client_review", drawing.discipline_id
    ):
        actions.append("comment")
    return actions


def submit_record(db, user, project, record_type, record_id, record_version, policy):
    open_review = (
        db.query(ReviewInstance)
        .filter_by(project_id=project.id, record_type=record_type, record_id=record_id)
        .filter(ReviewInstance.state.in_(["In Review", "Returned"]))
        .first()
    )
    if open_review:
        raise HTTPException(409, "A review is already open for this record")
    if not policy:
        raise HTTPException(422, "Mandatory reviewer configuration is missing")
    instance = ReviewInstance(
        organization_id=user.organization_id,
        project_id=project.id,
        record_type=record_type,
        record_id=record_id,
        record_version=record_version,
        policy_id=policy.id,
        state="In Review",
        submitted_by=user.id,
    )
    db.add(instance)
    db.flush()
    from app.domains.coordination.service import _open_stage, current_stage

    _open_stage(db, user, project, instance, current_stage(db, instance))
    return instance


def _release_allowed(db, project, revision):
    setting = db.get(EngineeringSetting, project.id)
    if revision.client_outcome == "Approved":
        return True
    return bool(
        revision.client_outcome == "Approved with Comments"
        and setting
        and setting.allow_release_with_comments
    )


CATEGORIES = {
    "Electrical": [
        "Single-line diagrams",
        "Cable layouts",
        "Cable schedules",
        "Earthing layouts",
        "Lighting layouts",
    ],
    "Mechanical": [
        "Equipment layouts",
        "General arrangement drawings",
        "Installation details",
        "Fabrication drawings",
    ],
    "Civil": [
        "Site layouts",
        "Foundations",
        "Structural drawings",
        "Reinforcement details",
    ],
    "Instrumentation": [
        "Instrument layouts",
        "Loop diagrams",
        "Hook-up drawings",
        "Wiring diagrams",
        "Instrument schedules",
    ],
    "Piping": ["P&IDs", "Piping layouts", "Isometric drawings", "Support drawings"],
}


def ensure_categories(db):
    for discipline_name, labels in CATEGORIES.items():
        discipline = db.query(Discipline).filter_by(name=discipline_name).first()
        if not discipline:
            continue
        for label in labels:
            key = label.casefold()
            if (
                db.query(DrawingCategory)
                .filter_by(discipline_id=discipline.id, name_key=key)
                .first()
            ):
                continue
            db.add(
                DrawingCategory(discipline_id=discipline.id, name=label, name_key=key, active=True)
            )
    db.flush()


def list_categories(db, user):
    if not user.organization_id:
        raise HTTPException(403, "Organization required")
    ensure_categories(db)
    rows = db.query(DrawingCategory).filter_by(active=True).order_by(DrawingCategory.name).all()
    return [
        {"id": row.id, "discipline_id": row.discipline_id, "name": row.name} for row in rows
    ]


def cancel_drawing(db, user, project, drawing, reason, version):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    _same_company(user, drawing)
    if drawing.record_status == "Cancelled":
        raise HTTPException(409, "This drawing is already cancelled")
    if not reason.strip():
        raise HTTPException(422, "A cancellation reason is required")
    _bump(
        db,
        Drawing,
        drawing.id,
        version,
        {"record_status": "Cancelled", "cancellation_reason": reason.strip()},
    )


def link_drawing(db, user, project, drawing, related_id):
    _require(db, user, project, "drawings.edit", drawing.discipline_id)
    related = db.get(Drawing, related_id)
    if (
        not related
        or related.project_id != project.id
        or related.organization_id != drawing.organization_id
    ):
        raise HTTPException(422, "Related drawing must belong to this project")
    if related.id == drawing.id:
        raise HTTPException(422, "A drawing cannot be linked to itself")
    if db.query(DrawingLink).filter_by(drawing_id=drawing.id, related_id=related.id).first():
        raise HTTPException(409, "These drawings are already linked")
    row = DrawingLink(project_id=project.id, drawing_id=drawing.id, related_id=related.id)
    db.add(row)
    db.flush()
    return row


def save_engineering_setting(db, project, enabled):
    row = db.get(EngineeringSetting, project.id)
    if not row:
        row = EngineeringSetting(project_id=project.id, allow_release_with_comments=enabled)
        db.add(row)
    else:
        row.allow_release_with_comments = enabled
    db.flush()
    return row


def record_client_decision(db, user, project, drawing, revision, data):
    _require(db, user, project, "transmittals.issue")
    if revision.drawing_id != drawing.id:
        raise HTTPException(404, "Revision not found")
    if revision.review_state != "Client Review":
        raise HTTPException(409, "An external decision can be recorded only during client review")
    outcomes = {
        "Approved": "Accepted",
        "Approved with Comments": "Accepted with Comments",
        "Revise and Resubmit": "Returned",
        "Rejected": "Rejected",
    }
    if data.outcome not in outcomes:
        raise HTTPException(422, "Unknown client outcome")
    if data.outcome in {"Revise and Resubmit", "Rejected"} and not data.reason.strip():
        raise HTTPException(422, "A comment is required for this outcome")
    if not data.external_reviewer.strip() or not data.reference.strip():
        raise HTTPException(422, "Record the external reviewer and a response reference")
    _bump(
        db,
        DrawingRevision,
        revision.id,
        data.version,
        {
            "review_state": outcomes[data.outcome],
            "client_outcome": data.outcome,
            "external_reviewer": data.external_reviewer.strip(),
            "decision_reference": data.reference.strip(),
            "decision_recorded_by": user.id,
        },
    )
    open_review = (
        db.query(ReviewInstance)
        .filter_by(record_type=DRAWING_REVIEW, record_id=revision.id)
        .filter(ReviewInstance.state == "In Review")
        .first()
    )
    if open_review:
        open_review.state = "Approved" if data.outcome.startswith("Approved") else (
            "Rejected" if data.outcome == "Rejected" else "Returned"
        )
    _event(
        db,
        user,
        drawing,
        revision,
        "client_decision",
        f"{data.outcome} by {data.external_reviewer.strip()}; recorded by user {user.id}; {data.reference.strip()}",
    )


def export_drawings(db, user, project, pairs):
    _require(db, user, project, "drawings.view")
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "number",
            "title",
            "discipline",
            "category",
            "company",
            "preparation",
            "internal_status",
            "client_status",
            "priority",
            "planned_date",
            "source",
            "status",
        ]
    )
    for drawing, latest in pairs:
        try:
            _require(db, user, project, "drawings.view", drawing.discipline_id)
        except HTTPException:
            continue
        if not sees_unreleased(db, user, project) and not _released_revision(db, drawing.id):
            continue
        category = db.get(DrawingCategory, drawing.category_id) if drawing.category_id else None
        discipline = db.get(Discipline, drawing.discipline_id)
        company = db.get(Company, drawing.company_id)
        writer.writerow(
            [
                drawing.number,
                drawing.title,
                discipline.name if discipline else "",
                category.name if category else drawing.drawing_type,
                company.name if company else "",
                _preparation_label(drawing),
                _internal_label(latest),
                _client_label(latest),
                drawing.priority,
                drawing.planned_date.isoformat() if drawing.planned_date else "",
                drawing.provenance,
                drawing.record_status,
            ]
        )
    return buffer.getvalue()


def my_tasks(db, user, project):
    _require(db, user, project, "drawings.view")
    today = date.today()
    preparation, reviews, comments, overdue = [], [], [], []
    for drawing, latest in filtered_drawings(db, project, record_status="Active"):
        try:
            _require(db, user, project, "drawings.view", drawing.discipline_id)
        except HTTPException:
            continue
        if drawing.responsible_user_id == user.id and drawing.preparation != "Prepared":
            preparation.append({"id": drawing.id, "number": drawing.number, "title": drawing.title})
            if drawing.planned_date and drawing.planned_date < today:
                overdue.append({"id": drawing.id, "number": drawing.number, "kind": "preparation"})
        if latest:
            from app.domains.coordination.models import ReviewInstance
            from app.domains.coordination.service import _reviewers, current_stage

            open_review = (
                db.query(ReviewInstance)
                .filter_by(record_type=DRAWING_REVIEW, record_id=latest.id, state="In Review")
                .order_by(ReviewInstance.id.desc())
                .first()
            )
            stage = current_stage(db, open_review) if open_review else None
            if stage and any(row.user_id == user.id for row in _reviewers(db, stage.id)):
                reviews.append(
                    {
                        "id": drawing.id,
                        "number": drawing.number,
                        "stage": stage.name or stage.purpose or "Review",
                    }
                )
        revision_ids = [
            row.id for row in db.query(DrawingRevision.id).filter_by(drawing_id=drawing.id)
        ]
        if not revision_ids:
            continue
        for comment in db.query(DrawingComment).filter(
            DrawingComment.revision_id.in_(revision_ids),
            DrawingComment.status != "Closed",
        ):
            if not (
                can_respond_comment(db, user, project, drawing, comment)
                or can_close_comment(db, user, project, drawing, comment)
            ):
                continue
            comments.append(
                {
                    "id": comment.id,
                    "drawing_id": drawing.id,
                    "number": drawing.number,
                    "status": comment.status,
                }
            )
    return {
        "preparation": preparation,
        "reviews": reviews,
        "comments": comments,
        "overdue": overdue,
    }


REVIEW_GUARDS.append(_guard)
