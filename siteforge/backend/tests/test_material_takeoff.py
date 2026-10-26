from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from .test_engineering import _ids, _member, _user, sign_in
from .test_phase2 import complete

HEADER = ["Item", "Description", "Specification", "Qty", "Unit"]
REV_A = [
    ["1", "XLPE cable 4C x 16", "1.1 kV armoured", "100", "m"],
    ["2", "Cable gland 25mm", "Brass double compression", "", "nos"],
    ["3", "Earth strip", "GI 25x3", "40", "zz"],
    ["4", "Cable tag", "SS", "12 nos", ""],
]
REV_B = [
    ["1", "XLPE cable 4C x 16", "1.1 kV armoured", "120", "m"],
    ["2", "Earth strip", "GI 50x6", "40", "m"],
    ["3", "Cable tag", "SS", "12", "nos"],
    ["4", "Cable tray 300mm", "HDG perforated", "30", "m"],
]
VERIFIED_A = [
    ("XLPE cable 4C x 16", Decimal("100"), "m"),
    ("Cable gland 25mm", None, "nos"),
    ("Earth strip", Decimal("40"), "zz"),
    ("Cable tag", Decimal("12"), "nos"),
]


def _escape(text):
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _pdf(rows=None, text_only=None):
    """A one-page PDF with a ruled table, a text line, or nothing readable (scanned)."""
    ops = ["0.5 w"]
    if rows:
        xs = [40, 80, 240, 390, 440, 490]
        top, height = 750, 18
        table = [HEADER, *rows]
        bottom = top - height * len(table)
        for index in range(len(table) + 1):
            y = top - index * height
            ops.append(f"{xs[0]} {y} m {xs[-1]} {y} l S")
        for x in xs:
            ops.append(f"{x} {top} m {x} {bottom} l S")
        for r, row in enumerate(table):
            for k, value in enumerate(row):
                if value:
                    ops.append(f"BT /F1 8 Tf {xs[k] + 3} {top - r * height - 13} Td ({_escape(value)}) Tj ET")
    elif text_only:
        ops.append(f"BT /F1 10 Tf 50 700 Td ({_escape(text_only)}) Tj ET")
    else:
        ops.append("100 100 400 500 re S")
    stream = "\n".join(ops).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def test_native_table_extraction_matches_the_engineer_list_and_never_invents_values():
    from app.domains.engineering_documents import extraction

    lines, provider, message = extraction.extract(_pdf(REV_A), "A.pdf", "application/pdf", "1")
    assert provider == "native-pdf"
    assert len(lines) == len(VERIFIED_A)
    for line, (description, quantity, unit) in zip(lines, VERIFIED_A):
        assert line["description"] == description
        assert line["quantity"] == quantity
        assert line["unit"] == unit
        assert line["page"] == 1 and line["table_ref"] == "T1" and line["bbox"]
    assert "Missing quantity" in lines[1]["flags"] and lines[1]["confidence"] < Decimal("0.7")
    assert any("not a recognised unit" in flag for flag in lines[2]["flags"])
    assert "Unit read from the quantity cell" in lines[3]["flags"]
    assert lines[0]["flags"] == [] and lines[0]["specification"] == "1.1 kV armoured"
    assert "4 line(s)" in message

    with pytest.raises(extraction.ExtractionProblem) as scanned:
        extraction.extract(_pdf(), "scan.pdf", "application/pdf", "")
    assert scanned.value.status == "Needs setup"
    assert "no text layer" in str(scanned.value) and "No OCR/table provider" in str(scanned.value)
    with pytest.raises(extraction.ExtractionProblem) as plain:
        extraction.extract(_pdf(text_only="General arrangement, see notes"), "ga.pdf", "application/pdf", "")
    assert plain.value.status == "Failed" and "No material table" in str(plain.value)
    with pytest.raises(extraction.ExtractionProblem) as cad:
        extraction.extract(b"AC1032", "plan.dwg", "application/acad", "")
    assert "not supported" in str(cad.value)
    with pytest.raises(extraction.ExtractionProblem) as pages:
        extraction.extract(_pdf(REV_A), "A.pdf", "application/pdf", "2")
    assert "has 1 page" in str(pages.value)


def test_scanned_table_uses_the_configured_provider(monkeypatch):
    import httpx

    from app.config import settings
    from app.domains.engineering_documents import extraction

    monkeypatch.setattr(settings, "ocr_provider", "test-ocr")
    monkeypatch.setattr(settings, "ocr_endpoint", "http://ocr.local/tables")
    monkeypatch.setattr(settings, "ocr_api_key", "server-secret")
    seen = {}

    class Reply:
        status_code = 200

        def json(self):
            return {
                "tables": [
                    {
                        "page": 1,
                        "rows": [["Description", "Qty", "UOM"], ["Pipe support", "8", "nos"], ["Grating", "", "m2"]],
                        "row_boxes": [[10, 10, 300, 20], [10, 20, 300, 30], [10, 30, 300, 40]],
                        "row_confidence": [0.99, 0.88, 0.7],
                    }
                ]
            }

    def fake_post(url, headers, data, files, timeout):
        seen.update(url=url, headers=headers, pages=data["pages"])
        return Reply()

    monkeypatch.setattr(httpx, "post", fake_post)
    lines, provider, _message = extraction.extract(_pdf(), "scan.pdf", "application/pdf", "")
    assert provider == "native-pdf+ocr:test-ocr"
    assert seen["headers"]["Authorization"] == "Bearer server-secret" and seen["pages"] == "1"
    assert [line["description"] for line in lines] == ["Pipe support", "Grating"]
    assert lines[0]["quantity"] == Decimal("8") and lines[0]["confidence"] == Decimal("0.880")
    assert lines[1]["quantity"] is None and "Missing quantity" in lines[1]["flags"]
    assert extraction.ocr_status()["configured"] is True


def _setup(c, maker):
    sign_in(c)
    roles, disciplines, me = _ids(c)
    electrical = disciplines["Electrical"]
    project = complete(c, code="MTO-1", discipline_ids=[electrical])
    pid = project["id"]
    company_id = me["company_id"]
    people = {
        "engineer": _user(c, "eng.mto@alpha.test", "Engineer MTO", company_id),
        "head": _user(c, "head.mto@alpha.test", "Head MTO", company_id),
        "approver": _user(c, "approver.mto@alpha.test", "Approver MTO", company_id),
        "site": _user(c, "site.mto@alpha.test", "Site MTO", company_id),
        "buyer": _user(c, "buyer.mto@alpha.test", "Buyer MTO", company_id),
        "viewer": _user(c, "viewer.mto@alpha.test", "Viewer MTO", company_id),
    }
    _member(c, pid, people["engineer"], roles["Engineering Engineer"], company_id, [electrical])
    _member(c, pid, people["head"], roles["Discipline Head"], company_id, [electrical])
    _member(c, pid, people["approver"], roles["Discipline Head"], company_id, [electrical])
    _member(c, pid, people["site"], roles["Site Engineer"], company_id, [])
    _member(c, pid, people["buyer"], roles["Procurement Officer"], company_id, [])
    _member(c, pid, people["viewer"], roles["Viewer"], company_id, [])
    policy = c.post(
        f"/api/projects/{pid}/approval-policies",
        json={
            "record_type": "procurement.requirement",
            "discipline_id": electrical,
            "stages": [{"name": "Discipline", "rule": "all", "reviewer_ids": [people["approver"]]}],
        },
    )
    assert policy.status_code == 201, policy.text
    items = {}
    for code, unit in (("CABLE-1", "m"), ("GLAND-1", "nos"), ("STRIP-1", "m"), ("TAG-1", "nos"), ("TRAY-1", "m")):
        item = c.post(
            "/api/procurement/items",
            json={
                "code": code,
                "description": code.title(),
                "item_type": "Material",
                "category": "Electrical",
                "discipline_ids": [electrical],
                "base_unit": unit,
            },
        )
        assert item.status_code == 201, item.text
        items[code] = item.json()["id"]
    drawing = c.post(
        f"/api/projects/{pid}/drawings",
        json={
            "number": "E-MTO",
            "title": "Cable schedule",
            "discipline_id": electrical,
            "company_id": company_id,
            "responsible_user_id": people["head"],
        },
    )
    assert drawing.status_code == 201, drawing.text
    did = drawing.json()["id"]
    upload = c.post(
        f"/api/projects/{pid}/drawings/{did}/revisions",
        data={"code": "A"},
        files={"upload": ("E-MTO-A.pdf", _pdf(REV_A), "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    rev_a = next(row["id"] for row in upload.json()["revisions"] if row["code"] == "A")
    return pid, did, rev_a, electrical, company_id, people, items, me


def _edit(c, pid, takeoff, line, **values):
    response = c.put(
        f"/api/projects/{pid}/takeoffs/{takeoff['id']}/lines/{line['id']}",
        json={**values, "version": takeoff["version"]},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _prepare_and_verify(c, pid, takeoff, matches, fixes=None):
    sign_in(c, "eng.mto@alpha.test")
    for line in list(takeoff["lines"]):
        fresh = next(row for row in takeoff["lines"] if row["id"] == line["id"])
        values = {"item_id": matches[fresh["description"]]}
        values.update((fixes or {}).get(fresh["description"], {}))
        takeoff = _edit(c, pid, takeoff, fresh, **values)
    sent = c.post(f"/api/projects/{pid}/takeoffs/{takeoff['id']}/submit", json={"version": takeoff["version"]})
    assert sent.status_code == 200, sent.text
    own = c.post(f"/api/projects/{pid}/takeoffs/{takeoff['id']}/verify", json={"version": sent.json()["version"]})
    assert own.status_code == 403
    sign_in(c, "head.mto@alpha.test")
    verified = c.post(f"/api/projects/{pid}/takeoffs/{takeoff['id']}/verify", json={"version": sent.json()["version"]})
    assert verified.status_code == 200, verified.text
    return verified.json()


def test_drawing_mto_becomes_one_requirement_then_stock_and_requisition(env, tmp_path):
    from app.config import settings
    from app.domains.coordination.models import Notification
    from app.domains.supply_chain.models import PurchaseOrder, StockBalance

    settings.file_root = str(tmp_path)
    c, maker = env
    pid, did, rev_a, electrical, company_id, people, items, me = _setup(c, maker)
    base = f"/api/projects/{pid}/drawings/{did}/revisions/{rev_a}/takeoffs"

    sign_in(c, "viewer.mto@alpha.test")
    assert c.post(f"{base}/extract", json={"pages": "1"}).status_code == 403
    assert c.get(f"/api/projects/{pid}/takeoffs").status_code == 403
    sign_in(c, "admin@beta.test")
    assert c.post(f"{base}/extract", json={"pages": "1"}).status_code == 404

    sign_in(c, "site.mto@alpha.test")
    hidden = c.post(f"{base}/extract", json={"pages": "1"})
    assert hidden.status_code == 403 and "released" in hidden.json()["detail"]

    sign_in(c, "eng.mto@alpha.test")
    started = c.post(f"{base}/extract", json={"pages": "1"})
    assert started.status_code == 202, started.text
    tid = started.json()["id"]
    takeoff = c.get(f"/api/projects/{pid}/takeoffs/{tid}").json()
    assert takeoff["status"] == "Draft" and takeoff["job"]["status"] == "Succeeded"
    assert takeoff["source"] == "Extraction" and takeoff["revision_code"] == "A"
    assert [line["quantity"] for line in takeoff["lines"]] == ["100", None, "40", "12"]
    gland = takeoff["lines"][1]
    assert "Missing quantity" in gland["flags"] and gland["original"]["quantity"] is None
    again = c.post(f"{base}/extract", json={"pages": "1"})
    assert again.status_code == 200 and again.json()["id"] == tid
    assert len(c.get(f"/api/projects/{pid}/takeoffs/{tid}").json()["jobs"]) == 1
    with maker() as db:
        ready = db.query(Notification).filter_by(user_id=people["head"], title="MTO extraction ready for review").count()
        assert ready == 1

    original = c.get(f"/api/projects/{pid}/takeoffs/{tid}/export?view=original&format=csv")
    assert original.status_code == 200
    text = original.content.decode("utf-8-sig")
    assert "Unverified extraction" in text and "does not verify or approve" in text
    assert c.get(f"/api/projects/{pid}/takeoffs/{tid}").json()["status"] == "Draft"

    takeoff = _edit(c, pid, takeoff, gland, quantity="5")
    takeoff = _edit(c, pid, takeoff, takeoff["lines"][2], unit="m")
    extra = c.post(
        f"/api/projects/{pid}/takeoffs/{tid}/lines",
        json={"description": "Spare lugs", "version": takeoff["version"]},
    )
    assert extra.status_code == 201, extra.text
    takeoff = extra.json()
    assert takeoff["lines"][-1]["quantity"] is None
    assert "Missing quantity" in takeoff["lines"][-1]["flags"]
    stored_gland = next(line for line in takeoff["lines"] if line["id"] == gland["id"])
    assert stored_gland["quantity"] == "5" and stored_gland["original"]["quantity"] is None
    assert any(item["field"] == "quantity" and item["new"] == "5" for item in takeoff["history"])

    matches = {
        "XLPE cable 4C x 16": items["CABLE-1"],
        "Cable gland 25mm": items["GLAND-1"],
        "Earth strip": items["STRIP-1"],
        "Cable tag": items["TAG-1"],
        "Spare lugs": items["TAG-1"],
    }
    sign_in(c, "eng.mto@alpha.test")
    for line in takeoff["lines"]:
        takeoff = _edit(c, pid, takeoff, line, item_id=matches[line["description"]])
    sent = c.post(f"/api/projects/{pid}/takeoffs/{tid}/submit", json={"version": takeoff["version"]})
    assert sent.status_code == 200, sent.text
    sign_in(c, "head.mto@alpha.test")
    blocked = c.post(f"/api/projects/{pid}/takeoffs/{tid}/verify", json={"version": sent.json()["version"]})
    assert blocked.status_code == 422 and "no quantity or unit" in blocked.json()["detail"]
    lugs = sent.json()["lines"][-1]
    removed = c.post(
        f"/api/projects/{pid}/takeoffs/{tid}/lines/{lugs['id']}/delete",
        json={"reason": "Not on the drawing", "version": sent.json()["version"]},
    )
    assert removed.status_code == 200, removed.text
    verified = c.post(f"/api/projects/{pid}/takeoffs/{tid}/verify", json={"version": removed.json()["version"]})
    assert verified.status_code == 200, verified.text
    takeoff = verified.json()
    assert takeoff["status"] == "Verified" and takeoff["verified_by_name"] == "Head MTO"
    assert len(takeoff["deleted_lines"]) == 1
    locked = c.put(
        f"/api/projects/{pid}/takeoffs/{tid}/lines/{takeoff['lines'][0]['id']}",
        json={"quantity": "999", "version": takeoff["version"]},
    )
    assert locked.status_code == 409
    reviewed = c.get(f"/api/projects/{pid}/takeoffs/{tid}/export?view=reviewed&format=xlsx")
    sheet = load_workbook(BytesIO(reviewed.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[1][4] == "Verification status" and rows[2][4] == "Verified" and rows[2][2] == "A"

    body = {
        "number": "MR-MTO-1",
        "requested_by": people["head"],
        "required_date": "2026-12-01",
        "purpose": "Cable schedule E-MTO rev A",
        "version": takeoff["version"],
    }
    made = c.post(f"/api/projects/{pid}/takeoffs/{tid}/requirement", json=body)
    assert made.status_code == 201, made.text
    requirement = made.json()["requirement"]
    repeat = c.post(f"/api/projects/{pid}/takeoffs/{tid}/requirement", json={**body, "number": "MR-MTO-2"})
    assert repeat.status_code == 200 and repeat.json()["requirement"]["id"] == requirement["id"]
    assert requirement["source"] == "Drawing MTO" and requirement["takeoff_number"] == takeoff["number"]
    assert requirement["status"] == "Draft" and requirement["drawing_revision_code"] == "A"
    assert [line["quantity"] for line in requirement["lines"]] == ["100.000", "5.000", "40.000", "12.000"]
    assert all(line["takeoff_line_id"] and line["drawing_revision_code"] == "A" for line in requirement["lines"])
    submitted = c.post(
        f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/submit",
        json={"version": requirement["version"]},
    )
    assert submitted.status_code == 201, submitted.text
    sign_in(c, "approver.mto@alpha.test")
    review = submitted.json()["review"]
    approved = c.post(
        f"/api/reviews/{review['id']}/decisions",
        json={"action": "approve", "reason": "", "version": review["version"]},
    )
    assert approved.status_code == 200, approved.text

    sign_in(c)
    warehouse = c.post(
        f"/api/projects/{pid}/procurement/warehouses",
        json={"code": "WH-MTO", "name": "Site store", "officer_id": me["id"]},
    )
    assert warehouse.status_code == 201, warehouse.text
    with maker() as db:
        from app.domains.identity_projects.models import Project

        db.add(
            StockBalance(
                organization_id=db.get(Project, pid).organization_id,
                project_id=pid,
                warehouse_id=warehouse.json()["id"],
                item_id=items["CABLE-1"],
                unit="m",
                on_hand=Decimal("30"),
                reserved=Decimal("0"),
            )
        )
        db.commit()
    sign_in(c, "buyer.mto@alpha.test")
    check = c.get(f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/fulfilment").json()
    cable = check["lines"][0]
    assert check["blocked"] == "" and check["can_allocate"] and check["can_requisition"]
    assert cable["suggested_allocations"] == [{"warehouse_id": warehouse.json()["id"], "quantity": "30.000"}]
    assert cable["takeoff_number"] == takeoff["number"] and cable["drawing_revision_code"] == "A"
    allocated = c.post(
        f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/allocations",
        json={"lines": [{"line_id": cable["line_id"], "warehouse_id": warehouse.json()["id"], "quantity": "30"}]},
    )
    assert allocated.status_code == 201, allocated.text
    assert allocated.json()["lines"][0]["remaining_to_purchase"] == "70.000"
    twice = c.post(
        f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/allocations",
        json={"lines": [{"line_id": cable["line_id"], "warehouse_id": warehouse.json()["id"], "quantity": "1"}]},
    )
    assert twice.status_code == 422 and "free" in twice.json()["detail"]
    needs = c.get(f"/api/projects/{pid}/procurement/purchase-needs").json()
    assert [row["number"] for row in needs["items"]] == ["MR-MTO-1"]
    assert needs["items"][0]["lines"][0]["remaining_to_purchase"] == "70.000"
    pr_body = {
        "number": "PR-MTO-1",
        "request_date": "2026-10-12",
        "required_delivery": "2026-11-20",
        "justification": "Uncovered cable schedule quantities",
    }
    pr = c.post(f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/requisition", json=pr_body)
    assert pr.status_code == 201, pr.text
    assert [line["quantity"] for line in pr.json()["lines"]] == ["70.000", "5.000", "40.000", "12.000"]
    first = pr.json()["lines"][0]
    assert first["requirement_number"] == "MR-MTO-1" and first["takeoff_number"] == takeoff["number"]
    assert first["drawing_revision_code"] == "A" and pr.json()["status"] == "Draft"
    duplicate = c.post(
        f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/requisition",
        json={**pr_body, "number": "PR-MTO-2"},
    )
    assert duplicate.status_code == 409
    assert c.get(f"/api/projects/{pid}/procurement/purchase-needs").json()["items"] == []
    with maker() as db:
        assert db.query(PurchaseOrder).filter_by(project_id=pid).count() == 0

    with maker() as db:
        from app.domains.engineering_documents.models import DrawingRevision

        db.get(DrawingRevision, rev_a).review_state = "Returned"
        db.commit()
    sign_in(c)
    upload = c.post(
        f"/api/projects/{pid}/drawings/{did}/revisions",
        data={"code": "B"},
        files={"upload": ("E-MTO-B.pdf", _pdf(REV_B), "application/pdf")},
    )
    assert upload.status_code == 201, upload.text
    rev_b = next(row["id"] for row in upload.json()["revisions"] if row["code"] == "B")
    sign_in(c, "eng.mto@alpha.test")
    started_b = c.post(f"/api/projects/{pid}/drawings/{did}/revisions/{rev_b}/takeoffs/extract", json={"pages": ""})
    assert started_b.status_code == 202, started_b.text
    takeoff_b = c.get(f"/api/projects/{pid}/takeoffs/{started_b.json()['id']}").json()
    assert takeoff_b["change_summary"]["previous"] == takeoff["number"]
    verified_b = _prepare_and_verify(
        c,
        pid,
        takeoff_b,
        {
            "XLPE cable 4C x 16": items["CABLE-1"],
            "Earth strip": items["STRIP-1"],
            "Cable tag": items["TAG-1"],
            "Cable tray 300mm": items["TRAY-1"],
        },
    )
    summary = verified_b["change_summary"]
    assert verified_b["change_state"] == "Review required"
    assert [row["new"] for row in summary["quantity_changed"]] == ["120 m"]
    assert [row["new"] for row in summary["specification_changed"]] == ["GI 50x6"]
    assert [row["label"] for row in summary["added"]] == ["TRAY-1 · Tray-1"]
    assert [row["label"] for row in summary["removed"]] == ["GLAND-1 · Gland-1"]
    assert summary["requisitions"][0]["number"] == "PR-MTO-1" and summary["allocations"] == 1
    assert c.get(f"/api/projects/{pid}/takeoffs/{tid}").json()["status"] == "Superseded"
    with maker() as db:
        alerted = db.query(Notification).filter_by(
            user_id=people["buyer"], title="Drawing revision affects a requirement"
        ).count()
        assert alerted == 1
    sign_in(c, "buyer.mto@alpha.test")
    held = c.get(f"/api/projects/{pid}/procurement/requirements/{requirement['id']}").json()
    assert held["impact_state"] == "Review required"
    waiting = c.get(f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/fulfilment").json()
    assert waiting["blocked"] and not waiting["can_allocate"]

    sign_in(c, "head.mto@alpha.test")
    early = c.post(f"/api/projects/{pid}/takeoffs/{verified_b['id']}/apply", json={"version": verified_b["version"]})
    assert early.status_code == 409 and "impact decision" in early.json()["detail"]
    sign_in(c, "approver.mto@alpha.test")
    decided = c.post(
        f"/api/projects/{pid}/procurement/requirements/{requirement['id']}/impact",
        json={"decision": "Apply rev B quantities", "affected_quantity": "20", "version": held["version"]},
    )
    assert decided.status_code == 200, decided.text
    sign_in(c, "head.mto@alpha.test")
    on_pr = c.post(f"/api/projects/{pid}/takeoffs/{verified_b['id']}/apply", json={"version": verified_b["version"]})
    assert on_pr.status_code == 422 and "requisition" in on_pr.json()["detail"]
    sign_in(c, "buyer.mto@alpha.test")
    cancelled = c.post(
        f"/api/projects/{pid}/procurement/requisitions/{pr.json()['id']}/cancel",
        json={"reason": "Rev B changes the schedule", "version": pr.json()["version"]},
    )
    assert cancelled.status_code == 200, cancelled.text
    sign_in(c, "head.mto@alpha.test")
    current_b = c.get(f"/api/projects/{pid}/takeoffs/{verified_b['id']}").json()
    applied = c.post(f"/api/projects/{pid}/takeoffs/{verified_b['id']}/apply", json={"version": current_b["version"]})
    assert applied.status_code == 200, applied.text
    assert applied.json()["change_state"] == "Applied"
    moved = c.get(f"/api/projects/{pid}/procurement/requirements/{requirement['id']}").json()
    assert moved["status"] == "Draft" and moved["takeoff_id"] == verified_b["id"]
    assert moved["drawing_revision_code"] == "B"
    assert sorted((line["item_code"], line["quantity"]) for line in moved["lines"]) == [
        ("CABLE-1", "120.000"),
        ("GLAND-1", "0.000"),
        ("STRIP-1", "40.000"),
        ("TAG-1", "12.000"),
        ("TRAY-1", "30.000"),
    ]
    assert next(line for line in moved["lines"] if line["item_code"] == "CABLE-1")["stock_allocated_quantity"] == "30.000"


def test_manual_and_imported_lists_use_the_same_requirement_flow(env, tmp_path):
    from app.config import settings

    settings.file_root = str(tmp_path)
    c, maker = env
    pid, did, rev_a, electrical, company_id, people, items, me = _setup(c, maker)

    sign_in(c, "eng.mto@alpha.test")
    csv_body = (
        "Description,Specification,Qty,UOM,Item code,Page\n"
        "XLPE cable 4C x 16,1.1 kV,100,m,CABLE-1,1\n"
        "Cable gland,Brass,,nos,GLAND-1,1\n"
        "Earth strip,GI,forty,m,STRIP-1,1\n"
        "Cable tag,SS,12,,TAG-9,1\n"
    ).encode()
    base = f"/api/projects/{pid}/drawings/{did}/revisions/{rev_a}/takeoffs"
    preview = c.post(f"{base}/import?mode=preview", files={"upload": ("mto.csv", csv_body, "text/csv")})
    assert preview.status_code == 200, preview.text
    rows = preview.json()["rows"]
    assert preview.json()["takeoff"] is None
    assert rows[1]["quantity"] is None and "Missing quantity" in rows[1]["flags"]
    assert rows[2]["quantity"] is None and any("not a number" in flag for flag in rows[2]["flags"])
    assert "Missing unit" in rows[3]["flags"] and any("TAG-9" in flag for flag in rows[3]["flags"])
    saved = c.post(f"{base}/import?mode=confirm", files={"upload": ("mto.csv", csv_body, "text/csv")})
    assert saved.status_code == 201, saved.text
    imported = saved.json()["takeoff"]
    assert imported["source"] == "Import" and imported["status"] == "Draft" and imported["line_count"] == 4
    again = c.post(f"{base}/import?mode=confirm", files={"upload": ("mto.csv", csv_body, "text/csv")})
    assert again.json()["takeoff"]["id"] == imported["id"]

    manual = c.post(f"{base}/manual")
    assert manual.status_code == 201
    assert c.post(f"{base}/manual").json()["id"] == manual.json()["id"]
    sign_in(c, "head.mto@alpha.test")
    returned = c.post(
        f"/api/projects/{pid}/takeoffs/{imported['id']}/return",
        json={"reason": "x", "version": imported["version"]},
    )
    assert returned.status_code == 409
    sign_in(c, "eng.mto@alpha.test")
    sent = c.post(f"/api/projects/{pid}/takeoffs/{imported['id']}/submit", json={"version": imported["version"]})
    assert sent.status_code == 200
    sign_in(c, "head.mto@alpha.test")
    returned = c.post(
        f"/api/projects/{pid}/takeoffs/{imported['id']}/return",
        json={"reason": "Gland and strip quantities are missing", "version": sent.json()["version"]},
    )
    assert returned.status_code == 200 and returned.json()["status"] == "Returned"
    sign_in(c, "eng.mto@alpha.test")
    tasks = c.get("/api/tasks").json()
    titles = [row["title"] for row in (tasks["items"] if isinstance(tasks, dict) else tasks)]
    assert any(title.startswith(f"Correct {imported['number']}") for title in titles)

    sign_in(c, "site.mto@alpha.test")
    created = c.post(
        f"/api/projects/{pid}/procurement/requirements",
        json={
            "number": "MR-MAN-1",
            "company_id": company_id,
            "discipline_id": electrical,
            "requested_by": people["site"],
            "required_date": "2026-12-01",
            "purpose": "Site spares",
            "lines": [{"item_id": items["GLAND-1"], "quantity": "10", "unit": "nos", "required_date": "2026-12-01"}],
        },
    )
    assert created.status_code == 201, created.text
    manual_requirement = created.json()
    assert manual_requirement["source"] == "Manual" and manual_requirement["takeoff_id"] is None
    check = c.get(f"/api/projects/{pid}/procurement/requirements/{manual_requirement['id']}/fulfilment").json()
    assert "after the requirement is approved" in check["blocked"]
    submitted = c.post(
        f"/api/projects/{pid}/procurement/requirements/{manual_requirement['id']}/submit",
        json={"version": manual_requirement["version"]},
    )
    assert submitted.status_code == 201
    sign_in(c, "approver.mto@alpha.test")
    review = submitted.json()["review"]
    assert c.post(
        f"/api/reviews/{review['id']}/decisions",
        json={"action": "approve", "reason": "", "version": review["version"]},
    ).status_code == 200
    sign_in(c, "buyer.mto@alpha.test")
    needs = c.get(f"/api/projects/{pid}/procurement/purchase-needs").json()
    assert [(row["number"], row["source"]) for row in needs["items"]] == [("MR-MAN-1", "Manual")]
    assert needs["items"][0]["lines"][0]["remaining_to_purchase"] == "10.000"
