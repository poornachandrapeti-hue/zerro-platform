import csv
import io
import json
from types import SimpleNamespace
from openpyxl import load_workbook
from app.domains.supply_chain.catalogue_exchange import (
    EXPORT_COLUMNS,
    IMPORT_COLUMNS,
    exchange_row,
    export_catalogue_csv,
    export_catalogue_xlsx,
    is_approved_item,
)


def sample(status="Approved", active=True):
    return SimpleNamespace(
        code="MEC-001",
        description="Pump",
        item_type="Equipment",
        category="Pumps",
        subcategory="Centrifugal",
        base_unit="Nos",
        purchase_unit="Nos",
        issue_unit="Nos",
        approval_status=status,
        active=active,
        specification="50 m3/h",
        specification_attributes={"flow": 50, "head": "30 m"},
        size_rating="DN50",
        material_grade="SS316",
        applicable_standard="API 610",
        approved_make="ACME",
        model_number="P-1",
        manufacturer_part_number="MPN-1",
        hsn_code="8413",
        inspection_required=True,
        tracking_required=False,
        mtc_required=False,
        calibration_certificate_required=False,
        shelf_life_applicable=False,
        shelf_life_days=None,
        version=2,
        discipline_ids=[1],
    )


def test_only_approved_active_item_is_eligible():
    assert is_approved_item(sample())
    for status in ("Draft", "Pending Approval", "Rejected", None):
        assert not is_approved_item(sample(status=status))
    assert not is_approved_item(sample(active=False))
    assert not is_approved_item(None)


def test_export_has_import_compatible_columns_and_json():
    row = exchange_row(sample(), ["Mechanical"])
    raw = export_catalogue_csv([row])
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert list(rows[0]) == list(EXPORT_COLUMNS)
    assert list(IMPORT_COLUMNS) == list(EXPORT_COLUMNS[: len(IMPORT_COLUMNS)])
    assert json.loads(rows[0]["specification_attributes"]) == {"flow": 50, "head": "30 m"}
    assert rows[0]["code"] == "MEC-001"
    assert rows[0]["type"] == "Equipment"
    assert rows[0]["unit"] == "Nos"
    assert rows[0]["disciplines"] == "Mechanical"
    assert rows[0]["inspection"] == "yes"
    assert rows[0]["tracking"] == "no"
    assert rows[0]["approval_status"] == "Approved"


def test_export_empty_still_has_headers():
    raw = export_catalogue_csv([]).decode("utf-8-sig")
    assert raw.strip().splitlines()[0] == ",".join(EXPORT_COLUMNS)


def test_csv_escaping_and_unicode():
    item = sample()
    item.description = 'Pump, "Heavy" – 15 kW'
    row = exchange_row(item, ["Mechanical"])
    rows = list(csv.DictReader(io.StringIO(export_catalogue_csv([row]).decode("utf-8-sig"))))
    assert rows[0]["description"] == item.description


def test_export_does_not_mutate_input():
    item = sample()
    before = dict(item.__dict__)
    export_catalogue_csv([exchange_row(item, ["Mechanical"])])
    assert item.__dict__ == before


def test_export_xlsx_round_trip_headers_match_import():
    row = exchange_row(sample(), ["Mechanical"])
    book = load_workbook(io.BytesIO(export_catalogue_xlsx([row])))
    headers = [cell.value for cell in next(book.active.iter_rows(min_row=1, max_row=1))]
    assert headers[: len(IMPORT_COLUMNS)] == list(IMPORT_COLUMNS)
    values = [cell.value for cell in next(book.active.iter_rows(min_row=2, max_row=2))]
    assert values[0] == "MEC-001"
    assert values[headers.index("disciplines")] == "Mechanical"
